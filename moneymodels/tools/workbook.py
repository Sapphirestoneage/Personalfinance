#!/usr/bin/env python3
"""Build the Money Models workbook (an .xlsx) from the same tables the app reads.

    python3 moneymodels/tools/workbook.py            -> moneymodels/Money-Models-Workbook.xlsx

One sheet per planet with every level (lesson, question, boxes, checklists),
a Dashboard where every figure is a live formula over named cells, a Plays
sheet that says which play is open, and a Start Here sheet with the legend,
the example switch and the planets-by-bands grid. Money is in dollars here
(a spreadsheet displays; the app keeps cents), percents are fractions.
"""
import json, os, re, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, RadarChart, Reference, Series

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
def load(n): return json.load(open(os.path.join(APP, 'data', n + '.json')))
LV, RC, PL, DEMO = load('levels'), load('recipes'), load('plays'), load('demo')
OUT = os.path.join(APP, 'Money-Models-Workbook.xlsx')

FONT = 'Arial'
def f(size=10, bold=False, italic=False, color='000000'): return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
BLUE, GREEN, GREY, WHITE = '0000FF', '008000', '6B6B6B', 'FFFFFF'
INK = '1F2A44'
FILL_INPUT = PatternFill('solid', fgColor='FFFF00')
FILL_HEAD = PatternFill('solid', fgColor=INK)
FILL_BAND = {1: 'E8EEF9', 2: 'E6F4EA', 3: 'FDF1E1', 4: 'EFEAFB', 5: 'FBF3D9'}
FILL_SOFT = PatternFill('solid', fgColor='F3F4F6')
THIN = Side(style='thin', color='D0D5DD')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center')
MONEY = '$#,##0;($#,##0);-'
MONEY2 = '$#,##0.00;($#,##0.00);-'
PCT = '0.0%'
RATIO = '0.0"x"'
NUM = '#,##0'

wb = Workbook()
names = {}          # fact key -> "'Sheet'!$F$12"
example_of = {}     # fact key -> example value (sheet units)

def q(sheet): return "'" + sheet + "'"
def define(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

def sheet_units(field, v):
    if v is None: return None
    if field['kind'] == 'cents': return v / 100
    if field['kind'] == 'percent': return v / 100
    return v

# ---------------------------------------------------------------- Start Here
start = wb.active; start.title = 'Start Here'

# ---------------------------------------------------------------- planet sheets
PLANET_SHEETS = {}
BANDS = {b['band']: b for b in LV['bands']}
level_done_cells = {}   # level id -> "'Sheet'!$H$5"

def planet_sheet(planet, idx):
    title = f"{idx} {planet['label']}"
    ws = wb.create_sheet(title)
    PLANET_SHEETS[planet['id']] = title
    widths = {'A': 7, 'B': 10, 'C': 52, 'D': 24, 'E': 20, 'F': 16, 'G': 44, 'H': 10}
    for c, w in widths.items(): ws.column_dimensions[c].width = w
    ws['A1'] = f"Planet {idx} of 6: {planet['label']}"; ws['A1'].font = f(16, True, color=WHITE)
    ws['A2'] = planet['blurb']; ws['A2'].font = f(10, italic=True, color=WHITE)
    for r in (1, 2):
        for c in range(1, 9): ws.cell(r, c).fill = FILL_HEAD
    ws.merge_cells('A1:H1'); ws.merge_cells('A2:H2')
    heads = ['Level', 'Band', 'Title, lesson, questions and boxes', 'Your answer', 'Example', 'In use', 'Help', 'Done?']
    for i, h in enumerate(heads, 1):
        c = ws.cell(4, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX; c.alignment = CENTER
    ws['A3'] = 'Type in the yellow cells (column D). Column F shows the number in use: yours, or the example while the switch on Start Here says Yes.'; ws['A3'].font = f(9, italic=True, color=GREY)
    ws.freeze_panes = 'A5'
    row = 5
    dv_yes = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True); ws.add_data_validation(dv_yes)
    for lv in [l for l in LV['levels'] if l['planet'] == planet['id']]:
        band = BANDS[lv['band']]
        fill = PatternFill('solid', fgColor=FILL_BAND[lv['band']])
        head_row = row
        ws.cell(row, 1, lv['level']).alignment = CENTER
        ws.cell(row, 2, band['name'])
        ws.cell(row, 3, lv['title']).font = f(11, True)
        ws.cell(row, 7, f"About {lv['minutes']} minute{'s' if lv['minutes'] != 1 else ''}. {band['meaning']}.").font = f(9, italic=True, color=GREY)
        for c in range(1, 9):
            ws.cell(row, c).fill = fill; ws.cell(row, c).border = BOX
            if c in (1, 2): ws.cell(row, c).font = f(10, True, color=INK)
        row += 1
        lesson = '\n\n'.join(lv['lesson'])
        ws.cell(row, 3, lesson).font = f(10, color='333333'); ws.cell(row, 3).alignment = WRAP
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        lines = sum(max(1, math.ceil(len(p) / 140)) for p in lv['lesson']) + (len(lv['lesson']) - 1)
        ws.row_dimensions[row].height = max(15, 14 * lines + 4)
        row += 1
        done_formula = None
        if lv['kind'] == 'quiz':
            qz = lv['quiz']
            ws.cell(row, 3, qz['q']).font = f(10, True); ws.cell(row, 3).alignment = WRAP
            opts = qz['options']
            dv = DataValidation(type='list', formula1='"' + ','.join(o.replace(',', ';').replace('"', "'") for o in opts) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            d = ws.cell(row, 4); d.fill = FILL_INPUT; d.font = f(10, color=BLUE); d.alignment = WRAP; dv.add(d)
            ws.cell(row, 5, opts[qz['answer']].replace(',', ';').replace('"', "'")).font = f(9, color=GREY); ws.cell(row, 5).alignment = WRAP
            ws.cell(row, 6, f'=IF(useExample="Yes",E{row},D{row})').font = f(10, color=INK)
            right = opts[qz['answer']].replace(',', ';').replace('"', "'")
            ws.cell(row, 7, f'=IF(F{row}="","",IF(F{row}="{right}","Right. {qz["why"].replace(chr(34), chr(39))}","Not quite. {qz["why"].replace(chr(34), chr(39))}"))').alignment = WRAP
            ws.row_dimensions[row].height = max(30, 14 * math.ceil(len(qz['why']) / 44) + 4)
            done_formula = f'=IF(F{row}="","",IF(F{row}="{right}","Done","Not yet"))'
            row += 1
            ws.cell(row, 3, 'Choices: ' + ' / '.join(opts)).font = f(9, italic=True, color=GREY); ws.cell(row, 3).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
            ws.row_dimensions[row].height = max(15, 14 * math.ceil((len(' / '.join(opts)) + 9) / 140) + 2)
            row += 1
        elif lv['kind'] == 'checklist':
            first = row
            for it in lv['items']:
                ws.cell(row, 3, it['text']).alignment = WRAP
                d = ws.cell(row, 4); d.fill = FILL_INPUT; d.font = f(10, color=BLUE); d.alignment = CENTER; dv_yes.add(d)
                ws.cell(row, 5, 'Yes').font = f(9, color=GREY); ws.cell(row, 5).alignment = CENTER
                ws.cell(row, 6, f'=IF(useExample="Yes",E{row},D{row})').alignment = CENTER
                ws.cell(row, 7, 'Yes when it is true').font = f(9, italic=True, color=GREY)
                row += 1
            last = row - 1
            n = len(lv['items'])
            done_formula = f'=IF(COUNTIF(F{first}:F{last},"Yes")={n},"Done",IF(COUNTIF(F{first}:F{last},"Yes")>0,"Part",""))'
        else:
            cells = []
            for fld in lv['fields']:
                ws.cell(row, 3, fld['label']).alignment = WRAP
                d = ws.cell(row, 4); d.fill = FILL_INPUT; d.font = f(10, color=BLUE); d.alignment = WRAP
                ex = sheet_units(fld, DEMO['facts'].get(fld['key']))
                e = ws.cell(row, 5, ex); e.font = f(9, color=GREY); e.alignment = WRAP
                if fld['kind'] == 'choice':
                    labels = [o['label'].replace(',', ';') for o in fld['options']]
                    dv = DataValidation(type='list', formula1='"' + ','.join(labels) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(d)
                    e.value = next(o['label'].replace(',', ';') for o in fld['options'] if o['id'] == DEMO['facts'][fld['key']])
                    help_text = 'Pick one: ' + ' / '.join(labels)
                else:
                    help_text = fld.get('help', '')
                    fmt = {'cents': MONEY2, 'percent': PCT, 'count': NUM, 'days': NUM}.get(fld['kind'])
                    if fmt: d.number_format = fmt; e.number_format = fmt
                    if fld['kind'] == 'cents': help_text = 'Dollars. ' + help_text
                    if fld['kind'] == 'percent': help_text = 'Type a percent, e.g. 25%. ' + help_text
                    if fld['kind'] == 'days': help_text = 'Days. ' + help_text
                    if fld['kind'] in ('cents', 'percent', 'count', 'days'):
                        dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='-1000000000', allow_blank=True, error='Type a number, or leave it blank if you do not know yet.'); ws.add_data_validation(dv); dv.add(d)
                u = ws.cell(row, 6, f'=IF(useExample="Yes",E{row},IF(D{row}="","",D{row}))'); u.font = f(10, color=INK); u.alignment = WRAP
                if fmt if fld['kind'] != 'choice' else None: u.number_format = fmt
                ws.cell(row, 7, help_text).font = f(9, italic=True, color=GREY); ws.cell(row, 7).alignment = WRAP
                ws.row_dimensions[row].height = max(15, 14 * max(math.ceil(len(help_text) / 44), math.ceil(len(fld['label']) / 52), math.ceil(len(str(ex or '')) / 20)) + 2)
                key = fld['key']
                names[key] = f"{q(title)}!$F${row}"
                example_of[key] = ex
                define(key, names[key])
                cells.append(f'F{row}')
                row += 1
            done_formula = '=IF(COUNTBLANK(' + ','.join(cells) + ')=0,"Done",IF(COUNTBLANK(' + ','.join(cells) + f')<{len(cells)},"Part",""))' if len(cells) == 1 else \
                '=IF(AND(' + ','.join(c + '<>""' for c in cells) + '),"Done",IF(OR(' + ','.join(c + '<>""' for c in cells) + '),"Part",""))'
            if len(cells) == 1: done_formula = f'=IF({cells[0]}<>"","Done","")'
        h = ws.cell(head_row, 8, done_formula); h.font = f(10, True, color=INK); h.alignment = CENTER
        level_done_cells[lv['id']] = f"{q(title)}!$H${head_row}"
        for r in range(head_row + 1, row):
            for c in range(3, 9): ws.cell(r, c).border = BOX
        row += 1
    ws.conditional_formatting.add(f'H5:H{row}', CellIsRule(operator='equal', formula=['"Done"'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
    ws.conditional_formatting.add(f'H5:H{row}', CellIsRule(operator='equal', formula=['"Part"'], fill=PatternFill('solid', fgColor='FFEB9C')))
    ws.conditional_formatting.add(f'H5:H{row}', CellIsRule(operator='equal', formula=['"Not yet"'], fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    return ws

for i, p in enumerate(LV['planets'], 1): planet_sheet(p, i)

# ---------------------------------------------------------------- Dashboard
dash = wb.create_sheet('Dashboard')
for c, w in {'A': 40, 'B': 16, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 48}.items(): dash.column_dimensions[c].width = w
dash['A1'] = 'Dashboard: every figure the model can compute'; dash['A1'].font = f(16, True, color=WHITE)
dash['A2'] = 'Nothing here is typed. Every cell is a formula over the answers on the planet sheets. A figure reads "not yet" until every number it needs is in.'; dash['A2'].font = f(10, italic=True, color=WHITE)
for r in (1, 2):
    for c in range(1, 8): dash.cell(r, c).fill = FILL_HEAD
dash.merge_cells('A1:G1'); dash.merge_cells('A2:G2')
dash.sheet_view.showGridLines = False
dash.freeze_panes = 'A4'
K = {}   # recipe id -> cell address on Dashboard (e.g. 'B6')
chart_row = [0]
def place(chart, near):
    a = max(near, chart_row[0]); dash.add_chart(chart, f'I{a}'); chart_row[0] = a + 20; return a + 20
row = [4]
def head(text):
    r = row[0]; c = dash.cell(r, 1, text); c.font = f(12, True, color=WHITE)
    for i in range(1, 8): dash.cell(r, i).fill = PatternFill('solid', fgColor='3B4A6B')
    row[0] += 1
def needs_ok(needs):
    parts = []
    for n in needs:
        if n in K: parts.append(f'ISNUMBER({n})')
        else: parts.append(f'{n}<>""')
    return 'AND(' + ','.join(parts) + ')' if len(parts) > 1 else parts[0]
def figure(rid, label, expr, needs, fmt=None, note=None, status=None):
    r = row[0]
    dash.cell(r, 1, label).alignment = WRAP
    c = dash.cell(r, 2, f'=IF({needs_ok(needs)},{expr},"not yet")'); c.font = f(10, True, color=INK); c.alignment = Alignment(horizontal='right')
    if fmt: c.number_format = fmt
    if status: dash.cell(r, 3, status.replace('{v}', f'B{r}')).font = f(10, color=INK)
    dash.cell(r, 7, note or next((x['says'] for x in RC['recipes'] if x['id'] == rid), '')).font = f(9, italic=True, color=GREY); dash.cell(r, 7).alignment = WRAP
    dash.row_dimensions[r].height = max(15, 13 * math.ceil(len(dash.cell(r, 7).value or '') / 48) + 2)
    for i in range(1, 8): dash.cell(r, i).border = BOX
    K[rid] = f'B{r}'; define(rid, f"{q('Dashboard')}!$B${r}")
    row[0] += 1
    return f'B{r}'

head('Foundations: the business today')
figure('gpPerSale', 'Gross profit per sale, today', 'price-deliveryCost', ['price', 'deliveryCost'], MONEY)
figure('grossMargin', 'Gross margin, today', 'IF(price=0,"no margin",(price-deliveryCost)/price)', ['price', 'deliveryCost'], PCT, status='=IF(ISNUMBER({v}),IF({v}>=0.5,"good",IF({v}>=0.3,"watch","low")),"")')
figure('cac', 'Cost of a customer (CAC)', 'IF(newCustomers=0,"no customers",adSpend/newCustomers)', ['adSpend', 'newCustomers'], MONEY)
figure('firstSaleGap', 'First sale minus the cost of the customer', 'gpPerSale-cac', ['gpPerSale', 'cac'], MONEY, status='=IF(ISNUMBER({v}),IF({v}>=0,"pays for itself","the rest of the model covers this"),"")')
figure('cpl', 'Cost per lead', 'IF(leads=0,"no leads",adSpend/leads)', ['adSpend', 'leads'], MONEY)
figure('impliedCustomers', 'Customers the funnel implies', 'ROUND(leads*attrConversion,0)', ['leads', 'attrConversion'], NUM)
figure('baselineRatio', '30-day ratio, today (cash collected / CAC)', 'IF(cac=0,"free customers",cash30Today/cac)', ['cash30Today', 'cac'], RATIO, status='=IF(ISNUMBER({v}),IF({v}>=2,"self-funding",IF({v}>=1,"break even","losing")),"")')
figure('baselinePayback', 'Days to get cash back, today', 'daysToCash', ['daysToCash'], NUM)
row[0] += 1
head('The sequence: each offer, per customer')
figure('attrGp', 'Attraction offer gross profit', 'attrPrice-attrCost', ['attrPrice', 'attrCost'], MONEY)
figure('upGp', 'Upsell gross profit (per taker)', 'upPrice-upCost', ['upPrice', 'upCost'], MONEY)
figure('upExpected', 'Upsell profit per customer', 'upGp*upTake', ['upGp', 'upTake'], MONEY)
figure('downGp', 'Downsell gross profit (per taker)', 'downPrice-downCost', ['downPrice', 'downCost'], MONEY)
figure('downExpected', 'Downsell profit per customer', 'downGp*(1-upTake)*downTake', ['downGp', 'upTake', 'downTake'], MONEY)
figure('contGp', 'Continuity gross profit per month', 'contPrice-contCost', ['contPrice', 'contCost'], MONEY)
figure('contMonth1', 'Continuity, first month, per customer', 'contGp*contJoin', ['contGp', 'contJoin'], MONEY, note='Monthly gross profit times the share who join.')
figure('expectedMonths', 'Months a member stays, on average', 'IF(churn=0,"nobody leaves",1/churn)', ['churn'], '0.0')
figure('contLtgp', 'Continuity lifetime profit per customer', 'contGp*contJoin*expectedMonths', ['contGp', 'contJoin', 'expectedMonths'], MONEY)
row[0] += 1
head('The headline: the 30-day rule and lifetime value')
figure('gp30', '30-day cash per customer (planned, gross profit)', 'attrGp+upExpected+downExpected+contMonth1', ['attrGp', 'upExpected', 'downExpected', 'contMonth1'], MONEY)
figure('collected30', '30-day cash collected per customer (before delivery)', 'attrPrice+upPrice*upTake+downPrice*(1-upTake)*downTake+contPrice*contJoin', ['attrPrice', 'upPrice', 'upTake', 'downPrice', 'downTake', 'contPrice', 'contJoin'], MONEY)
figure('ratio30', '30-day ratio (planned)', 'IF(cac=0,"free customers",gp30/cac)', ['gp30', 'cac'], RATIO, status='=IF(ISNUMBER({v}),IF({v}>=2,"SELF-FUNDING: each customer pays for the next",IF({v}>=1,"break even inside the month","not paying back inside 30 days")),"")')
figure('maxCac', 'The most you could pay for a customer', 'gp30/2', ['gp30'], MONEY)
figure('ltgp', 'Lifetime gross profit per customer (LTGP)', 'attrGp+upExpected+downExpected+contLtgp', ['attrGp', 'upExpected', 'downExpected', 'contLtgp'], MONEY)
figure('ltgpCac', 'LTGP to CAC', 'IF(cac=0,"free customers",ltgp/cac)', ['ltgp', 'cac'], RATIO, status='=IF(ISNUMBER({v}),IF({v}>=3,"good: 3x or more",IF({v}>=1,"thin","losing")),"")')
figure('cashPerLead', '30-day cash per lead', 'gp30*attrConversion', ['gp30', 'attrConversion'], MONEY)
figure('breakEvenCustomers', 'Customers a month to cover fixed costs', 'IF(gp30<=0,"no profit in 30 days",ROUNDUP(fixedCosts/gp30,0))', ['fixedCosts', 'gp30'], NUM)
figure('affordableCustomers', 'Customers you can afford per month', 'IF(cac=0,"no ceiling",ROUNDDOWN(growthCash/cac,0))', ['growthCash', 'cac'], NUM)
figure('capacityHeadroom', 'Room to grow before capacity', 'capacity-newCustomers', ['capacity', 'newCustomers'], NUM)
figure('gpPerHour', 'Gross profit per hour of your time', 'IF(hoursPerCustomer=0,"no hours",gp30/hoursPerCustomer)', ['gp30', 'hoursPerCustomer'], MONEY)
row[0] += 1
head('Prove: what actually happened')
figure('effectiveUpTake', 'Real upsell take across all customers', 'showRate*upOfferedShare*upTake', ['showRate', 'upOfferedShare', 'upTake'], PCT)
figure('attrGpNet', 'Attraction gross profit after refunds', 'attrGp-attrPrice*refundRate', ['attrGp', 'attrPrice', 'refundRate'], MONEY)
figure('downCash30', 'Downsell cash inside 30 days (first payment only)', 'downPrice*planFirstShare*(1-upTake)*downTake', ['downPrice', 'planFirstShare', 'upTake', 'downTake'], MONEY)
figure('cash30Verified', '30-day cash, checked against timing', 'attrGp+IF(upWhen<=30,upExpected+(downPrice*planFirstShare-downCost)*(1-upTake)*downTake,0)+IF(contFirstDays<=30,contMonth1,0)', ['attrGp', 'upWhen', 'upExpected', 'downPrice', 'planFirstShare', 'downCost', 'upTake', 'downTake', 'contFirstDays', 'contMonth1'], MONEY)
figure('planLeak', 'Cash lost to failed payment plans, per customer', 'downPrice*(1-planFirstShare)*planDefault*(1-upTake)*downTake', ['downPrice', 'planFirstShare', 'planDefault', 'upTake', 'downTake'], MONEY)
figure('retain3Predicted', 'Month-three retention, predicted by churn', '(1-churn)^3', ['churn'], PCT)
figure('retain3Measured', 'Month-three retention, measured', 'retain3', ['retain3'], PCT)
figure('measuredRatio', '30-day ratio, measured (cash collected / CAC)', 'IF(cac=0,"free customers",measuredCash30/cac)', ['measuredCash30', 'cac'], RATIO, status='=IF(ISNUMBER({v}),IF({v}>=2,"self-funding",IF({v}>=1,"break even","losing")),"")')
figure('customerGrowth', 'Customer growth since the model', 'IF(newCustomers=0,"no customers before",measuredNewCustomers/newCustomers-1)', ['measuredNewCustomers', 'newCustomers'], PCT)
row[0] += 1
head('Optimize')
figure('runRevenue', 'Cash a full run brings in', 'attrCap*collected30', ['attrCap', 'collected30'], MONEY)
figure('up2Expected', 'Second upsell profit per customer', 'up2Price*up2Take*upTake', ['up2Price', 'up2Take', 'upTake'], MONEY)
figure('anchorGap', 'The upsell as a share of the anchor', 'IF(anchorPrice=0,"no anchor",upPrice/anchorPrice)', ['anchorPrice', 'upPrice'], PCT)
figure('annualLift', 'Cash the yearly plan moves into month one', '(contAnnualPrice-contPrice)*contAnnualTake*contJoin', ['contAnnualPrice', 'contPrice', 'contAnnualTake', 'contJoin'], MONEY)
figure('gp30Stacked', '30-day cash with the yearly plan', 'gp30+annualLift', ['gp30', 'annualLift'], MONEY)
figure('ltgpFull', 'LTGP with the second upsell and the yearly plan', 'ltgp+up2Expected+annualLift', ['ltgp', 'up2Expected', 'annualLift'], MONEY)
figure('entryDownExpected', 'Entry downsell cash per lead', 'entryDownPrice*entryDownTake*(1-attrConversion)', ['entryDownPrice', 'entryDownTake', 'attrConversion'], MONEY)
figure('cashPerLeadFull', '30-day cash per lead with the entry downsell', 'cashPerLead+entryDownExpected', ['cashPerLead', 'entryDownExpected'], MONEY)
figure('breakEvenTake', 'Take rate that keeps a price test even', 'IF(upGp+upPrice*priceRaise<=0,"no profit",upTake*upGp/(upGp+upPrice*priceRaise))', ['upGp', 'upPrice', 'priceRaise', 'upTake'], PCT)

# ---- Tables that feed the charts
row[0] += 1
head('Chart data: the 30-day cash, offer by offer')
r0 = row[0]
for i, h in enumerate(['Offer', 'Gross profit', 'Cost of a customer (1x)', 'Self-funding line (2x)'], 1):
    c = dash.cell(r0, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
rows30 = [('Attraction', 'attrGp'), ('Upsell', 'upExpected'), ('Downsell', 'downExpected'), ('Continuity, month 1', 'contMonth1'), ('Total, 30 days', 'gp30')]
for i, (lab, rid) in enumerate(rows30, 1):
    r = r0 + i
    dash.cell(r, 1, lab); dash.cell(r, 2, f'=IF(ISNUMBER({rid}),{rid},0)').number_format = MONEY
    dash.cell(r, 3, '=IF(ISNUMBER(cac),cac,0)').number_format = MONEY; dash.cell(r, 4, '=IF(ISNUMBER(cac),2*cac,0)').number_format = MONEY
    for cc in range(1, 5): dash.cell(r, cc).border = BOX
dash.cell(r0 + 6, 1, 'Text in a figure counts as 0 in the charts, so an unfinished chart reads low, never wrong.').font = f(9, italic=True, color=GREY)
ch = BarChart(); ch.type = 'col'; ch.title = 'The 30-day cash, offer by offer'; ch.y_axis.title = 'Gross profit per customer ($)'
ch.add_data(Reference(dash, min_col=2, min_row=r0, max_row=r0 + 5), titles_from_data=True)
ch.set_categories(Reference(dash, min_col=1, min_row=r0 + 1, max_row=r0 + 5))
ln = LineChart(); ln.add_data(Reference(dash, min_col=3, max_col=4, min_row=r0, max_row=r0 + 5), titles_from_data=True)
ch += ln; ch.height = 9; ch.width = 18; ch.legend.position = 'b'
cb = place(ch, r0 - 1)
row[0] = max(r0 + 8, cb)

head('Chart data: before, planned, measured (cash collected in 30 days)')
r1 = row[0]
for i, h in enumerate(['', 'Cash collected', 'Cost of a customer'], 1):
    c = dash.cell(r1, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
for i, (lab, ex) in enumerate([('Before (today)', 'IF(cash30Today="",0,cash30Today)'), ('Planned', 'IF(ISNUMBER(collected30),collected30,0)'), ('Measured', 'IF(measuredCash30="",0,measuredCash30)')], 1):
    dash.cell(r1 + i, 1, lab); dash.cell(r1 + i, 2, '=' + ex).number_format = MONEY; dash.cell(r1 + i, 3, '=IF(ISNUMBER(cac),cac,0)').number_format = MONEY
    for cc in range(1, 4): dash.cell(r1 + i, cc).border = BOX
ch = BarChart(); ch.type = 'col'; ch.title = 'Before, planned, measured'; ch.y_axis.title = '$ per customer, first 30 days'
ch.add_data(Reference(dash, min_col=2, min_row=r1, max_row=r1 + 3), titles_from_data=True); ch.set_categories(Reference(dash, min_col=1, min_row=r1 + 1, max_row=r1 + 3))
ln = LineChart(); ln.add_data(Reference(dash, min_col=3, min_row=r1, max_row=r1 + 3), titles_from_data=True); ch += ln
ch.height = 9; ch.width = 18; ch.legend.position = 'b'; cb = place(ch, r1 - 1)
row[0] = max(r1 + 6, cb)

head('Chart data: the funnel, from 100 leads')
r2 = row[0]
for i, h in enumerate(['Stage', 'People'], 1):
    c = dash.cell(r2, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
fun = [('Leads', '100'), ('Take the attraction offer', 'IF(attrConversion="",0,100*attrConversion)'), ('Take the upsell', f'IF(upTake="",0,B{r2 + 2}*upTake)'), ('Take the downsell', f'IF(downTake="",0,(B{r2 + 2}-B{r2 + 3})*downTake)'), ('Join continuity', f'IF(contJoin="",0,B{r2 + 2}*contJoin)')]
for i, (lab, ex) in enumerate(fun, 1):
    dash.cell(r2 + i, 1, lab); dash.cell(r2 + i, 2, '=' + ex).number_format = '0.0'
    for cc in range(1, 3): dash.cell(r2 + i, cc).border = BOX
ch = BarChart(); ch.type = 'bar'; ch.title = 'The sequence as a funnel (from 100 leads)'
ch.add_data(Reference(dash, min_col=2, min_row=r2, max_row=r2 + 5), titles_from_data=True); ch.set_categories(Reference(dash, min_col=1, min_row=r2 + 1, max_row=r2 + 5))
ch.height = 9; ch.width = 18; ch.legend = None; ch.x_axis.scaling.orientation = 'maxMin'; cb = place(ch, r2 - 1)
row[0] = max(r2 + 8, cb)

head('Chart data: twelve months if you reinvest the 30-day cash')
r3 = row[0]
for i, h in enumerate(['Month', 'Budget for customers', 'Customers, reinvesting', 'Customers, flat budget', 'Capacity'], 1):
    c = dash.cell(r3, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
ready = 'AND(ISNUMBER(cac),cac>0,ISNUMBER(gp30),growthCash<>"",reinvestShare<>"",capacity<>"")'
for m in range(1, 13):
    r = r3 + m
    dash.cell(r, 1, m)
    dash.cell(r, 2, f'=IF({ready},' + ('growthCash' if m == 1 else f'growthCash+reinvestShare*gp30*C{r - 1}') + ',0)').number_format = MONEY
    dash.cell(r, 3, f'=IF({ready},MIN(capacity,B{r}/cac),0)').number_format = '0.0'
    dash.cell(r, 4, f'=IF({ready},MIN(capacity,growthCash/cac),0)').number_format = '0.0'
    dash.cell(r, 5, '=IF(capacity="",0,capacity)').number_format = NUM
    for cc in range(1, 6): dash.cell(r, cc).border = BOX
ch = LineChart(); ch.title = 'Customers per month: reinvesting against a flat budget'; ch.y_axis.title = 'New customers'; ch.x_axis.title = 'Month'
ch.add_data(Reference(dash, min_col=3, max_col=5, min_row=r3, max_row=r3 + 12), titles_from_data=True); ch.set_categories(Reference(dash, min_col=1, min_row=r3 + 1, max_row=r3 + 12))
ch.height = 9; ch.width = 18; ch.legend.position = 'b'; cb = place(ch, r3 - 1)
row[0] = max(r3 + 15, cb)

head('Chart data: members left, month by month (100 join)')
r4 = row[0]
for i, h in enumerate(['Month', 'Members left'], 1):
    c = dash.cell(r4, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
for m in range(0, 13):
    r = r4 + 1 + m
    dash.cell(r, 1, m); dash.cell(r, 2, f'=IF(churn="",0,100*(1-churn)^{m})').number_format = '0.0'
    for cc in range(1, 3): dash.cell(r, cc).border = BOX
ch = LineChart(); ch.title = 'Members left, month by month'; ch.y_axis.title = 'Of 100 who joined'; ch.x_axis.title = 'Month'
ch.add_data(Reference(dash, min_col=2, min_row=r4, max_row=r4 + 13), titles_from_data=True); ch.set_categories(Reference(dash, min_col=1, min_row=r4 + 1, max_row=r4 + 13))
ch.height = 9; ch.width = 18; ch.legend = None; cb = place(ch, r4 - 1)
row[0] = max(r4 + 16, cb)

head('Chart data: cash coming back, day by day (gross profit, cumulative)')
r5 = row[0]
for i, h in enumerate(['Day', 'Cumulative gross profit', 'Cost of a customer', 'Paid back on day'], 1):
    c = dash.cell(r5, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
tl_ready = 'AND(ISNUMBER(attrGp),ISNUMBER(upExpected),ISNUMBER(downExpected),ISNUMBER(contMonth1),upWhen<>"",contFirstDays<>"",churn<>"")'
days = list(range(0, 91, 3))
for i, d in enumerate(days):
    r = r5 + 1 + i
    dash.cell(r, 1, d)
    cont = '+'.join(f'IF({d}>=contFirstDays+{30 * k},contMonth1*(1-churn)^{k},0)' for k in range(0, 4))
    dash.cell(r, 2, f'=IF({tl_ready},attrGp+IF({d}>=upWhen,upExpected+downExpected,0)+{cont},0)').number_format = MONEY
    dash.cell(r, 3, '=IF(ISNUMBER(cac),cac,0)').number_format = MONEY
    dash.cell(r, 4, f'=IF(AND({tl_ready},ISNUMBER(cac),B{r}>=cac),A{r},"")')
    for cc in range(1, 5): dash.cell(r, cc).border = BOX
last = r5 + len(days)
ch = LineChart(); ch.title = 'Cash coming back against the cost of a customer'; ch.y_axis.title = '$ per customer'; ch.x_axis.title = 'Day'
ch.add_data(Reference(dash, min_col=2, max_col=3, min_row=r5, max_row=last), titles_from_data=True); ch.set_categories(Reference(dash, min_col=1, min_row=r5 + 1, max_row=last))
ch.height = 9; ch.width = 18; ch.legend.position = 'b'; cb = place(ch, r5 - 1)
row[0] = max(last + 2, cb)
figure('paybackDays', 'Payback days (planned, to the nearest 3 days)', f'IF(COUNT(D{r5 + 1}:D{last})=0,"not inside 90 days",MIN(D{r5 + 1}:D{last}))', ['gp30', 'cac', 'upWhen', 'contFirstDays', 'churn'], NUM, note='The first day in the table above where the cumulative gross profit passes the cost of a customer.')

row[0] += 1
head('Chart data: the model\'s health, six ways (share of target, capped at 100%)')
r6 = row[0]
for i, h in enumerate(['Measure', 'Your figure', 'Target', 'Share of target'], 1):
    c = dash.cell(r6, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
axes = [('Margin', 'grossMargin', 0.5, PCT), ('30-day ratio', 'ratio30', 2, RATIO), ('LTGP to CAC', 'ltgpCac', 3, RATIO), ('Upsell take', 'upTake', 0.3, PCT), ('Monthly retention', 'IF(churn="","not yet",1-churn)', 0.95, PCT), ('Lead conversion', 'attrConversion', 0.2, PCT)]
for i, (lab, ex, tgt, fmt) in enumerate(axes, 1):
    r = r6 + i
    dash.cell(r, 1, lab)
    e = ex if ex.startswith('IF(') else f'IF(ISNUMBER({ex}),{ex},IF({ex}="","not yet",{ex}))' if ex in K else f'IF({ex}="","not yet",{ex})'
    dash.cell(r, 2, '=' + e).number_format = fmt
    dash.cell(r, 3, tgt).number_format = fmt
    dash.cell(r, 4, f'=IF(ISNUMBER(B{r}),MIN(1,MAX(0,B{r}/C{r})),0)').number_format = PCT
    for cc in range(1, 5): dash.cell(r, cc).border = BOX
ch = RadarChart(); ch.type = 'filled'; ch.title = "The model's health, six ways"
ch.add_data(Reference(dash, min_col=4, min_row=r6, max_row=r6 + 6), titles_from_data=True); ch.set_categories(Reference(dash, min_col=1, min_row=r6 + 1, max_row=r6 + 6))
ch.height = 9; ch.width = 18; ch.legend = None; ch.y_axis.scaling.max = 1; ch.y_axis.scaling.min = 0; cb = place(ch, r6 - 1)
row[0] = max(r6 + 9, cb)

dash.conditional_formatting.add(f'C5:C{row[0]}', FormulaRule(formula=['OR(LEFT(C5,4)="SELF",LEFT(C5,4)="good",LEFT(C5,4)="self",LEFT(C5,4)="pays")'], font=Font(name=FONT, bold=True, color='006100')))
dash.conditional_formatting.add(f'C5:C{row[0]}', FormulaRule(formula=['OR(LEFT(C5,5)="break",LEFT(C5,4)="thin",LEFT(C5,5)="watch")'], font=Font(name=FONT, bold=True, color='9C5700')))
dash.conditional_formatting.add(f'C5:C{row[0]}', FormulaRule(formula=['OR(LEFT(C5,3)="not",LEFT(C5,3)="low",LEFT(C5,6)="losing")'], font=Font(name=FONT, bold=True, color='9C0006')))
dash.page_setup.orientation = 'landscape'; dash.sheet_properties.pageSetUpPr.fitToPage = True; dash.page_setup.fitToWidth = 1; dash.page_setup.fitToHeight = 0

# ---------------------------------------------------------------- Your Model
model = wb.create_sheet('Your Model')
for c, w in {'A': 22, 'B': 60, 'C': 18, 'D': 30}.items(): model.column_dimensions[c].width = w
model['A1'] = 'Your money model: the four offers in order'; model['A1'].font = f(16, True, color=WHITE)
model['A2'] = 'What you wrote on the Build bands, with the price and take rate beside each. Read it top to bottom: that is the sequence a customer walks.'; model['A2'].font = f(10, italic=True, color=WHITE)
for r in (1, 2):
    for c in range(1, 5): model.cell(r, c).fill = FILL_HEAD
model.merge_cells('A1:D1'); model.merge_cells('A2:D2'); model.sheet_view.showGridLines = False
for i, h in enumerate(['Offer', 'What you say', 'Price', 'Who says yes'], 1):
    c = model.cell(4, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
steps = [('1. Attraction', 'attrOffer', 'attrPrice', 'IF(attrConversion="","not yet",TEXT(attrConversion,"0%")&" of leads")'),
         ('   then the string', 'attrString', None, None),
         ('2. Upsell', 'upScript', 'upPrice', 'IF(upTake="","not yet",TEXT(upTake,"0%")&" of customers")'),
         ('   on a no', 'upNoReply', None, None),
         ('3. Downsell', 'downScript', 'downPrice', 'IF(downTake="","not yet",TEXT(downTake,"0%")&" of the nos")'),
         ('   what is smaller', 'downCut', None, None),
         ('4. Continuity', 'contScript', 'contPrice', 'IF(contJoin="","not yet",TEXT(contJoin,"0%")&" join"&IF(churn="","",", "&TEXT(churn,"0%")&" leave each month"))'),
         ('   the reason to stay', 'contHook', None, None)]
for i, (lab, key, price, take) in enumerate(steps):
    r = 5 + i
    model.cell(r, 1, lab).font = f(10, not lab.startswith(' '))
    c = model.cell(r, 2, f'=IF({key}="","not written yet",{key})'); c.font = f(10, color=GREEN); c.alignment = WRAP
    if price: model.cell(r, 3, f'=IF({price}="","not yet",{price})').number_format = MONEY; model.cell(r, 3).font = f(10, color=GREEN)
    if take: model.cell(r, 4, '=' + take).font = f(10, color=GREEN)
    model.row_dimensions[r].height = 42
    for cc in range(1, 5): model.cell(r, cc).border = BOX
r = 14
model.cell(r, 1, 'The numbers that judge it').font = f(12, True); r += 1
for lab, rid, fmt in [('30-day ratio', 'ratio30', RATIO), ('Payback days', 'paybackDays', NUM), ('LTGP to CAC', 'ltgpCac', RATIO), ('Most you could pay for a customer', 'maxCac', MONEY), ('The promise', 'offerLine', None), ('The constraint', 'constraint', None), ('Where you start', 'firstPlanet', None)]:
    model.cell(r, 1, lab)
    if rid == 'offerLine': c = model.cell(r, 2, '=IF(OR(customer="",promise=""),"not yet","For "&customer&": "&promise)')
    elif rid in ('constraint', 'firstPlanet'): c = model.cell(r, 2, f'=IF({rid}="","not yet",{rid})')
    else: c = model.cell(r, 2, f'=IF(ISNUMBER({rid}),{rid},{rid})'); c.number_format = fmt
    c.font = f(10, True, color=GREEN); c.alignment = Alignment(horizontal='left', wrap_text=True)
    for cc in range(1, 3): model.cell(r, cc).border = BOX
    r += 1

# ---------------------------------------------------------------- Plays
plays = wb.create_sheet('Plays')
for c, w in {'A': 14, 'B': 30, 'C': 60, 'D': 18, 'E': 44}.items(): plays.column_dimensions[c].width = w
plays['A1'] = 'The plays: nineteen ways to build the four offers'; plays['A1'].font = f(16, True, color=WHITE)
plays['A2'] = 'A play reads Open when your numbers say it fits, Waiting until then. The five steps under each are the way to run it; mark them Yes as you go.'; plays['A2'].font = f(10, italic=True, color=WHITE)
for r in (1, 2):
    for c in range(1, 6): plays.cell(r, c).fill = FILL_HEAD
plays.merge_cells('A1:E1'); plays.merge_cells('A2:E2'); plays.sheet_view.showGridLines = False
for i, h in enumerate(['Planet', 'Play', 'What it is, and the steps', 'Open? / Done?', 'Opens when'], 1):
    c = plays.cell(4, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX
plays.freeze_panes = 'A5'
dv_yes = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True); plays.add_data_validation(dv_yes)
planet_label = {p['id']: p['label'] for p in LV['planets']}
def rule_expr(w):
    of, op = w['of'], w['op']
    if op == 'known': return f'{of}<>""' if of not in K else f'ISNUMBER({of})'
    v = w['value']
    if of in ('grossMargin', 'upTake', 'contJoin', 'churn'): v = v / 100
    if of in ('cpl', 'price', 'upPrice', 'contPrice', 'attrPrice', 'upExpected', 'upCost', 'contGp'): v = v / 100
    return f'AND(ISNUMBER({of}),{of}{op}{v})'
r = 5
for p in PL['plays']:
    plays.cell(r, 1, planet_label[p['planet']]).font = f(10, True)
    plays.cell(r, 2, p['label']).font = f(11, True)
    plays.cell(r, 3, p['what']).alignment = WRAP
    cond = 'AND(' + ','.join(rule_expr(w) for w in p['when']) + ')' if len(p['when']) > 1 else rule_expr(p['when'][0])
    s = plays.cell(r, 4, f'=IF({cond},"Open","Waiting")'); s.font = f(10, True, color=INK); s.alignment = CENTER
    plays.cell(r, 5, '; '.join(w['says'] for w in p['when'])).font = f(9, italic=True, color=GREY); plays.cell(r, 5).alignment = WRAP
    plays.row_dimensions[r].height = max(30, 14 * max(math.ceil(len(p['what']) / 60), math.ceil(len(plays.cell(r, 5).value) / 44)) + 2)
    for cc in range(1, 6): plays.cell(r, cc).fill = PatternFill('solid', fgColor=FILL_BAND[3]); plays.cell(r, cc).border = BOX
    r += 1
    for i, st in enumerate(p['steps'], 1):
        plays.cell(r, 3, f'{i}. {st}').alignment = WRAP
        d = plays.cell(r, 4); d.fill = FILL_INPUT; d.font = f(10, color=BLUE); d.alignment = CENTER; dv_yes.add(d)
        plays.cell(r, 5, 'Yes when done').font = f(9, italic=True, color=GREY)
        for cc in range(3, 6): plays.cell(r, cc).border = BOX
        r += 1
    r += 1
plays.conditional_formatting.add(f'D5:D{r}', CellIsRule(operator='equal', formula=['"Open"'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
plays.conditional_formatting.add(f'D5:D{r}', CellIsRule(operator='equal', formula=['"Waiting"'], fill=PatternFill('solid', fgColor='FFEB9C')))

# ---------------------------------------------------------------- Start Here (filled last, needs the sheet names)
ws = start
for c, w in {'A': 26, 'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 40}.items(): ws.column_dimensions[c].width = w
ws['A1'] = 'Money Models: a workbook for $100M Money Models'; ws['A1'].font = f(18, True, color=WHITE)
ws['A2'] = 'Get paid back before the bill is due. A money model is the order in which you ask a customer to say yes; this workbook walks you through building one.'; ws['A2'].font = f(10, italic=True, color=WHITE)
for r in (1, 2):
    for c in range(1, 9): ws.cell(r, c).fill = FILL_HEAD
ws.merge_cells('A1:H1'); ws.merge_cells('A2:H2'); ws.sheet_view.showGridLines = False
ws['A4'] = 'How to use it'; ws['A4'].font = f(12, True)
how = ['1. Six planet sheets, one per part of the model. Each has fifteen levels in five bands: Learn (read, answer one question), Numbers (type your figures), Build (the exercises), Prove (what happened when you ran it), Optimize (sharpen it).',
       '2. Type only in the yellow cells. Blue text is yours; black text is a formula; leave formulas alone.',
       '3. Every figure on the Dashboard is live. It reads "not yet" until every number it needs is in, so nothing is ever a fake zero.',
       '4. Plays says which of the nineteen plays your numbers support. Your Model reads the four offers back to you in order.',
       '5. Do the planets in order, or start where the constraint points: Foundations level 8 tells you.']
for i, t in enumerate(how):
    c = ws.cell(5 + i, 1, t); c.alignment = WRAP; ws.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=8); ws.row_dimensions[5 + i].height = 30
ws['A11'] = 'Use the example numbers?'; ws['A11'].font = f(11, True)
ws['B11'] = 'Yes'; ws['B11'].fill = FILL_INPUT; ws['B11'].font = f(11, True, color=BLUE); ws['B11'].alignment = CENTER; ws['B11'].border = BOX
dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=False); ws.add_data_validation(dv); dv.add(ws['B11'])
define('useExample', "'Start Here'!$B$11")
ws['C11'] = 'Yes: every sheet runs on a made-up six-week fitness challenge, so you can see the whole thing working. No: your own answers, and nothing else.'; ws['C11'].font = f(9, italic=True, color=GREY); ws['C11'].alignment = WRAP
ws.merge_cells('C11:H11'); ws.row_dimensions[11].height = 30
ws['A13'] = 'Legend'; ws['A13'].font = f(12, True)
ws['A14'] = 'Type here'; ws['A14'].fill = FILL_INPUT; ws['A14'].font = f(10, color=BLUE); ws['A14'].border = BOX; ws['B14'] = 'A yellow cell with blue text is yours to fill. Blank means "not known yet"; a typed 0 is a real zero.'; ws.merge_cells('B14:H14')
ws['A15'] = 'Computed'; ws['A15'].font = f(10, True, color=INK); ws['A15'].border = BOX; ws['B15'] = 'Black text is a formula. Green text is a figure read from another sheet.'; ws.merge_cells('B15:H15')
ws['A16'] = 'not yet'; ws['A16'].font = f(10, color=GREY); ws['A16'].border = BOX; ws['B16'] = 'A figure that is still missing an input. The Dashboard says which; the planet sheet says where to type it.'; ws.merge_cells('B16:H16')
ws['A17'] = 'Example'; ws['A17'].font = f(9, color=GREY); ws['A17'].border = BOX; ws['B17'] = 'The grey column on every planet sheet: invented numbers for a made-up business, never real ones.'; ws.merge_cells('B17:H17')
ws['A19'] = 'Where you are: levels done, by planet and band (3 of 3 clears a band)'; ws['A19'].font = f(12, True)
for i, h in enumerate(['Planet'] + [f"{b['band']} {b['name']}" for b in LV['bands']] + ['Done', 'Sheet'], 1):
    c = ws.cell(20, i, h); c.font = f(10, True); c.fill = FILL_SOFT; c.border = BOX; c.alignment = CENTER
for i, p in enumerate(LV['planets']):
    r = 21 + i; sh = PLANET_SHEETS[p['id']]
    ws.cell(r, 1, p['label']).font = f(10, True)
    for b in LV['bands']:
        c = ws.cell(r, 1 + b['band'], f'=COUNTIFS({q(sh)}!$H:$H,"Done",{q(sh)}!$B:$B,"{b["name"]}")'); c.alignment = CENTER; c.border = BOX
    c = ws.cell(r, 7, f'=COUNTIF({q(sh)}!$H:$H,"Done")&" of 15"'); c.alignment = CENTER; c.border = BOX
    ws.cell(r, 8, f'=HYPERLINK("#\'{sh}\'!A1","Open {sh}")').font = f(10, color=BLUE, )
    ws.cell(r, 1).border = BOX
ws.conditional_formatting.add('B21:F26', CellIsRule(operator='equal', formula=['3'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
ws.conditional_formatting.add('B21:F26', CellIsRule(operator='between', formula=['1', '2'], fill=PatternFill('solid', fgColor='FFEB9C')))
ws['A28'] = 'The sun: 30-day ratio'; ws['A28'].font = f(12, True)
ws['B28'] = '=IF(ISNUMBER(ratio30),ratio30,"not yet")'; ws['B28'].number_format = RATIO; ws['B28'].font = f(14, True, color=GREEN); ws['B28'].alignment = CENTER
ws['C28'] = '=IF(ISNUMBER(ratio30),IF(ratio30>=2,"Every customer pays for the next one.",IF(ratio30>=1,"Break even inside the month. Twice the cost is the self-funding line.","Not yet paying back inside 30 days.")),"Needs the Numbers band on Foundations, Attraction, Upsell, Downsell and Continuity.")'; ws.merge_cells('C28:H28'); ws['C28'].alignment = WRAP; ws.row_dimensions[28].height = 30
ws['A29'] = 'LTGP to CAC'; ws['A29'].font = f(12, True)
ws['B29'] = '=IF(ISNUMBER(ltgpCac),ltgpCac,"not yet")'; ws['B29'].number_format = RATIO; ws['B29'].font = f(14, True, color=GREEN); ws['B29'].alignment = CENTER
ws['C29'] = 'How many times over a customer pays for themselves. The book wants at least 3.'; ws.merge_cells('C29:H29')
ws['A31'] = 'About this workbook'; ws['A31'].font = f(12, True)
about = ['Every lesson is a paraphrase in this workbook\'s own words, never the book\'s text; the play names are plain-word labels for the book\'s offer types. Read the book for the reasoning and the stories.',
         'Thresholds (pay a customer back inside 30 days, twice over to self-fund, lifetime gross profit at least three times the cost of a customer) are the book\'s rules of thumb, not measured findings. Your own numbers decide.',
         'Not financial or business advice. The example numbers are invented. Built from the same tables as the Money Models app in the Personalfinance repository (moneymodels/).']
for i, t in enumerate(about):
    c = ws.cell(32 + i, 1, t); c.alignment = WRAP; c.font = f(9, color=GREY); ws.merge_cells(start_row=32 + i, start_column=1, end_row=32 + i, end_column=8); ws.row_dimensions[32 + i].height = 28

for sh in wb.worksheets:
    sh.page_setup.orientation = 'landscape'; sh.sheet_properties.pageSetUpPr.fitToPage = True; sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
    sh.print_options.horizontalCentered = True
    for rr in sh.iter_rows():
        for c in rr:
            if c.font is None or c.font.name != FONT:
                c.font = c.font.copy(name=FONT)
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print('wrote', OUT, 'names', len(wb.defined_names))
