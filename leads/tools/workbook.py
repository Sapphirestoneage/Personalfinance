#!/usr/bin/env python3
"""leads/tools/workbook.py: the Leads Ladder as one spreadsheet (LD-008).

Reads leads/data/book.json and writes leads/Leads-Ladder.xlsx: the same
book, exercises, checklists and Machine formulas as the app, for someone
who would rather work in a sheet. Example numbers only; every figure is
invented and says so.

    python3 leads/tools/workbook.py            writes the file
"""
import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
BOOK = json.load(open(os.path.join(ROOT, 'data', 'book.json')))
OUT = os.environ.get('OUT') or os.path.join(ROOT, 'Leads-Ladder.xlsx')

FONT = 'Arial'
def f(bold=False, size=10, color='000000', italic=False): return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)
NAVY = PatternFill('solid', fgColor='1F2A44'); GREY = PatternFill('solid', fgColor='E9ECF1'); LIGHT = PatternFill('solid', fgColor='F5F6F8')
YELLOW = PatternFill('solid', fgColor='FFFFCC'); GOLD = PatternFill('solid', fgColor='FCE9B6'); GREEN = PatternFill('solid', fgColor='D8F0DF'); AMBER = PatternFill('solid', fgColor='FBE5C3'); RED = PatternFill('solid', fgColor='F7D2D2')
BLUE = '0000FF'; LINK = '008000'
thin = Side(style='thin', color='C9CDD4'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top'); CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
SHEET = {'magnet': 'Magnet', 'warm': 'Warm', 'content': 'Content', 'cold': 'Cold', 'paid': 'Ads', 'referrals': 'Referrals', 'employees': 'Employees', 'agencies': 'Agencies', 'affiliates': 'Affiliates'}
MACHINE_INPUTS = [  # key, label, kind, example (display units), hint
    ('reachPerDay', 'Primary actions a day', 'count', 100, 'Messages, reach-outs, or ad impressions bought. The rule of 100 says a hundred.'),
    ('daysPerMonth', 'Days a month you do them', 'count', 22, ''),
    ('replyRate', 'Reach-outs that become engaged leads', 'rate', 0.05, 'Replied, clicked, or took the magnet and gave a way to reach them.'),
    ('bookRate', 'Engaged leads that book a call or a visit', 'rate', 0.40, ''),
    ('showRate', 'Bookings that show up', 'rate', 0.70, ''),
    ('closeRate', 'Shows that buy', 'rate', 0.30, ''),
    ('firstPurchaseCents', 'What a new customer pays up front ($)', 'money', 300, ''),
    ('monthlyCents', 'What they pay each month after that ($)', 'money', 200, 'Zero if nothing, and zero is an answer.'),
    ('monthsKept', 'Months a customer stays, on average', 'count', 8, ''),
    ('marginRate', 'Gross margin: what is left after delivering', 'rate', 0.70, ''),
    ('adSpendCents', 'Ad spend a month ($)', 'money', 3000, 'Zero if you run no ads.'),
    ('laborCents', 'Pay for the people doing outreach and content, a month ($)', 'money', 0, 'Zero if it is only you. Zero is an answer.'),
    ('customersNow', 'Paying customers today', 'count', 12, ''),
    ('churnRate', 'Share of customers who leave in a month', 'rate', 0.08, ''),
    ('referralRate', 'Share of customers who bring one in a month', 'rate', 0.10, ''),
]
MCELL = {}   # key -> "'Machine'!$B$n"

wb = Workbook()
wb.remove(wb.active)

def title_row(ws, row, text, sub=None, span=6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text); c.font = f(True, 14, 'FFFFFF'); c.fill = NAVY; c.alignment = Alignment(vertical='center')
    for col in range(1, span + 1): ws.cell(row=row, column=col).fill = NAVY
    ws.row_dimensions[row].height = 26
    if sub:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=span)
        s = ws.cell(row=row + 1, column=1, value=sub); s.font = f(False, 10, '444444', True); s.alignment = WRAP
        ws.row_dimensions[row + 1].height = 30
        return row + 2
    return row + 1

def head(ws, row, labels, fill=GREY):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l); c.font = f(True, 10); c.fill = fill; c.border = BOX; c.alignment = CENTER
    ws.row_dimensions[row].height = 30

def inp(c, kind=None):
    c.fill = YELLOW; c.font = f(False, 10, BLUE); c.border = BOX
    if kind == 'rate': c.number_format = '0.0%'
    elif kind == 'money': c.number_format = '$#,##0;($#,##0);-'
    elif kind == 'count': c.number_format = '#,##0.0;(#,##0.0);-' if False else '#,##0'

def legend(ws, row, span=6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value='How to use this sheet: type only in the yellow cells (blue text). Everything in black is worked out for you. A blank yellow cell means "not yet", not zero. Choose Y in a checklist cell to tick it.')
    c.font = f(False, 9, '444444', True); c.alignment = WRAP; ws.row_dimensions[row].height = 30
    return row + 1

yes = DataValidation(type='list', formula1='"Y"', allow_blank=True, showDropDown=False)

# ---------------------------------------------------------------- Machine
ws = wb.create_sheet('Machine')
ws.column_dimensions['A'].width = 52; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 16; ws.column_dimensions['D'].width = 16; ws.column_dimensions['E'].width = 16; ws.column_dimensions['F'].width = 60
r = title_row(ws, 1, 'The Machine: the numbers behind the book', 'Reach, the rates, what a customer costs and is worth, the book\'s two checks, the levers, referrals against churn. The Measure levels on the planet sheets read from here.')
r = legend(ws, r)
r += 1
head(ws, r, ['Your numbers', 'Value', '', '', '', 'What it means']); r += 1
ws.cell(row=r, column=1, value='These are the example numbers of an invented piano teacher. Replace every one with yours.').font = f(False, 9, 'A05A00', True); r += 1
for key, label, kind, ex, hint in MACHINE_INPUTS:
    ws.cell(row=r, column=1, value=label).font = f()
    c = ws.cell(row=r, column=2, value=ex); inp(c, kind)
    MCELL[key] = "'Machine'!$B$%d" % r
    ws.cell(row=r, column=6, value=hint).font = f(False, 9, '555555', True)
    r += 1
K = {k: 'B%d' % (int(MCELL[k].split('$')[-1])) for k in MCELL}
r += 1
head(ws, r, ['A month at these numbers', 'Value', '', '', '', 'How it is worked out']); r += 1
allrates = 'COUNT(%s,%s,%s,%s,%s,%s)=6' % (K['reachPerDay'], K['daysPerMonth'], K['replyRate'], K['bookRate'], K['showRate'], K['closeRate'])
def calc(label, formula, fmt, how, name=None):
    global r
    ws.cell(row=r, column=1, value=label).font = f(True)
    c = ws.cell(row=r, column=2, value=formula); c.font = f(); c.number_format = fmt; c.border = BOX
    ws.cell(row=r, column=6, value=how).font = f(False, 9, '555555', True)
    ref = 'B%d' % r
    if name: K[name] = ref
    r += 1
    return ref
calc('Reached a month', '=IF(COUNT(%s,%s)=2,%s*%s,"")' % (K['reachPerDay'], K['daysPerMonth'], K['reachPerDay'], K['daysPerMonth']), '#,##0', 'actions a day times days a month', 'reach')
calc('Engaged leads a month', '=IF(OR(%s="",%s=""),"",%s*%s)' % (K['reach'], K['replyRate'], K['reach'], K['replyRate']), '#,##0.0', 'reached times the reply rate', 'engaged')
calc('Booked', '=IF(OR(%s="",%s=""),"",%s*%s)' % (K['engaged'], K['bookRate'], K['engaged'], K['bookRate']), '#,##0.0', 'engaged times the booking rate', 'booked')
calc('Showed', '=IF(OR(%s="",%s=""),"",%s*%s)' % (K['booked'], K['showRate'], K['booked'], K['showRate']), '#,##0.0', 'booked times the show rate', 'shows')
calc('Customers a month', '=IF(OR(%s="",%s=""),"",%s*%s)' % (K['shows'], K['closeRate'], K['shows'], K['closeRate']), '#,##0.0', 'showed times the close rate', 'customers')
calc('Engaged leads a week', '=IF(%s="","",%s*12/52)' % (K['engaged'], K['engaged']), '#,##0.0', 'engaged a month, as a week (against the goal on Start Here)', 'perWeek')
r += 1
head(ws, r, ['What a lead and a customer cost', 'Value', '', '', '', 'How it is worked out']); r += 1
calc('Spend on getting leads, a month', '=IF(COUNT(%s,%s)=2,%s+%s,"")' % (K['adSpendCents'], K['laborCents'], K['adSpendCents'], K['laborCents']), '$#,##0;($#,##0);-', 'ads plus pay', 'spend')
calc('Cost per engaged lead', '=IF(OR(%s="",%s=""),"",IF(%s=0,"none made",%s/%s))' % (K['spend'], K['engaged'], K['engaged'], K['spend'], K['engaged']), '$#,##0.00;($#,##0.00);-', 'spend over engaged leads', 'cpel')
calc('Cost per customer', '=IF(OR(%s="",%s=""),"",IF(%s=0,"none made",%s/%s))' % (K['spend'], K['customers'], K['customers'], K['spend'], K['customers']), '$#,##0;($#,##0);-', 'spend over customers', 'cac')
r += 1
head(ws, r, ['What a customer is worth', 'Value', '', '', '', 'How it is worked out']); r += 1
vin = 'COUNT(%s,%s,%s,%s)=4' % (K['firstPurchaseCents'], K['monthlyCents'], K['monthsKept'], K['marginRate'])
calc('Lifetime gross profit', '=IF(%s,(%s+%s*%s)*%s,"")' % (vin, K['firstPurchaseCents'], K['monthlyCents'], K['monthsKept'], K['marginRate']), '$#,##0;($#,##0);-', '(up front + monthly times months kept) times margin', 'ltgp')
calc('Gross profit in the first thirty days', '=IF(%s,(%s+%s)*%s,"")' % (vin, K['firstPurchaseCents'], K['monthlyCents'], K['marginRate']), '$#,##0;($#,##0);-', '(up front + one month) times margin', 'first30')
r += 1
head(ws, r, ["The book's two checks", 'Value', 'Rule', '', '', 'Verdict']); r += 1
cacnum = 'ISNUMBER(%s)' % K['cac']
ws.cell(row=r, column=1, value='Lifetime gross profit over cost per customer').font = f(True)
c = ws.cell(row=r, column=2, value='=IF(OR(%s="",%s=""),"",IF(%s=0,"free",IF(%s,%s/%s,"")))' % (K['ltgp'], K['spend'], K['spend'], cacnum, K['ltgp'], K['cac'])); c.number_format = '0.0"x"'; c.font = f(True, 11); c.border = BOX; K['ratio'] = 'B%d' % r
ws.cell(row=r, column=3, value=3).number_format = '0"x or more"'; ws.cell(row=r, column=3).font = f(False, 10, BLUE); K['floor'] = 'C%d' % r
ws.cell(row=r, column=3).comment = Comment("The book's rule of thumb: lifetime gross profit at least three times what a customer cost to get. From $100M Leads, the paid ads chapter.", 'Leads Ladder')
ws.cell(row=r, column=6, value='=IF(%s="","needs the numbers above",IF(%s="free","Your leads cost nothing you counted, so every customer is profit.",IF(%s>=%s,"Passes: a customer is worth more than three times what they cost.","Under three to one: fix the offer, the margin, or the funnel before spending more.")))' % (K['ratio'], K['ratio'], K['ratio'], K['floor'])).font = f(); ws.cell(row=r, column=6).alignment = WRAP
ws.conditional_formatting.add(K['ratio'], FormulaRule(formula=['AND(ISNUMBER(%s),%s>=%s)' % (K['ratio'], K['ratio'], K['floor'])], fill=GREEN))
ws.conditional_formatting.add(K['ratio'], FormulaRule(formula=['AND(ISNUMBER(%s),%s<%s,%s>=1)' % (K['ratio'], K['ratio'], K['floor'], K['ratio'])], fill=AMBER))
ws.conditional_formatting.add(K['ratio'], FormulaRule(formula=['AND(ISNUMBER(%s),%s<1)' % (K['ratio'], K['ratio'])], fill=RED))
ratio_row = r; r += 1
ws.cell(row=r, column=1, value='First thirty days over cost per customer').font = f(True)
c = ws.cell(row=r, column=2, value='=IF(OR(%s="",%s=""),"",IF(%s=0,"free",IF(%s,%s/%s,"")))' % (K['first30'], K['spend'], K['spend'], cacnum, K['first30'], K['cac'])); c.number_format = '0.0"x"'; c.font = f(True, 11); c.border = BOX; K['payback'] = 'B%d' % r
ws.cell(row=r, column=3, value=2).number_format = '0"x or more"'; ws.cell(row=r, column=3).font = f(False, 10, BLUE); K['paymult'] = 'C%d' % r
ws.cell(row=r, column=3).comment = Comment("The book's second rule: the first thirty days should bring in twice what a customer cost, so the customers pay for the ads.", 'Leads Ladder')
ws.cell(row=r, column=6, value='=IF(%s="","needs the numbers above",IF(%s="free","Nothing to pay back.",IF(%s>=%s,"Passes: the first month pays for the next one. The book says this is when you can scale.","The first month does not pay back twice. Growth will eat cash; raise the up-front price or lower the cost.")))' % (K['payback'], K['payback'], K['payback'], K['paymult'])).font = f(); ws.cell(row=r, column=6).alignment = WRAP
ws.conditional_formatting.add(K['payback'], FormulaRule(formula=['AND(ISNUMBER(%s),%s>=%s)' % (K['payback'], K['payback'], K['paymult'])], fill=GREEN))
ws.conditional_formatting.add(K['payback'], FormulaRule(formula=['AND(ISNUMBER(%s),%s<%s)' % (K['payback'], K['payback'], K['paymult'])], fill=AMBER))
r += 2
head(ws, r, ['The levers: one small change', 'Customers a month', 'Extra customers', 'Change', '', 'Why this matters']); r += 1
lever_top = r
base = '%s*%s*%s*%s*%s*%s' % (K['reachPerDay'], K['daysPerMonth'], K['replyRate'], K['bookRate'], K['showRate'], K['closeRate'])
levers = [('As entered', base, ''), ]
for rk, word in [('replyRate', 'Reply rate'), ('bookRate', 'Booking rate'), ('showRate', 'Show rate'), ('closeRate', 'Close rate')]:
    levers.append((word + ' up one point', base.replace(K[rk], 'MIN(1,%s+0.01)' % K[rk]), '+1 point'))
levers.append(('Ten more actions a day', base.replace(K['reachPerDay'], '(%s+10)' % K['reachPerDay']), '+10 a day'))
for i, (label, formula, change) in enumerate(levers):
    ws.cell(row=r, column=1, value=label).font = f(i > 0)
    c = ws.cell(row=r, column=2, value='=IF(%s,%s,"")' % (allrates, formula)); c.number_format = '#,##0.00'; c.border = BOX; c.font = f()
    c = ws.cell(row=r, column=3, value='' if i == 0 else '=IF(B%d="","",B%d-$B$%d)' % (r, r, lever_top)); c.number_format = '#,##0.00'; c.border = BOX; c.font = f()
    ws.cell(row=r, column=4, value=change).font = f(False, 9, '555555')
    if i == 0: ws.cell(row=r, column=6, value='Every rate multiplies the same chain, so a tenth more of any rate is the same tenth more customers. One point is different: a point on a low rate is worth far more than a point on a high one. The biggest number in "Extra customers" is your lowest rate, and that is where "better" beats "more".').font = f(False, 9, '555555', True); ws.cell(row=r, column=6).alignment = WRAP
    r += 1
ws.merge_cells(start_row=lever_top, start_column=6, end_row=r - 1, end_column=6)
lever_end = r - 1
ch = BarChart(); ch.type = 'bar'; ch.title = 'Extra customers a month from one small change'; ch.style = 10
ch.add_data(Reference(ws, min_col=3, min_row=lever_top + 1, max_row=lever_end), titles_from_data=False)
ch.set_categories(Reference(ws, min_col=1, min_row=lever_top + 1, max_row=lever_end)); ch.legend = None; ch.height = 7; ch.width = 18
ws.add_chart(ch, 'H%d' % lever_top)
r += 1
head(ws, r, ['Referrals against churn: a year of customers', 'With referrals as entered', 'If nobody referred', '', '', 'What it means']); r += 1
g_top = r
gin = 'COUNT(%s,%s,%s)=3' % (K['customersNow'], K['churnRate'], K['referralRate'])
newc = 'IF(%s="",0,%s)' % (K['customers'], K['customers'])
for m in range(13):
    ws.cell(row=r, column=1, value='Now' if m == 0 else 'Month %d' % m).font = f()
    if m == 0:
        ws.cell(row=r, column=2, value='=IF(%s,%s,"")' % (gin, K['customersNow'])); ws.cell(row=r, column=3, value='=IF(%s,%s,"")' % (gin, K['customersNow']))
    else:
        ws.cell(row=r, column=2, value='=IF(B%d="","",B%d*(1-%s+%s)+%s)' % (r - 1, r - 1, K['churnRate'], K['referralRate'], newc))
        ws.cell(row=r, column=3, value='=IF(C%d="","",C%d*(1-%s)+%s)' % (r - 1, r - 1, K['churnRate'], newc))
    for col in (2, 3): ws.cell(row=r, column=col).number_format = '#,##0.0'; ws.cell(row=r, column=col).border = BOX; ws.cell(row=r, column=col).font = f()
    if m == 0:
        ws.cell(row=r, column=6, value='=IF(NOT(%s),"needs customers today, churn and the referral rate",IF(%s>%s,"Referrals outnumber the customers who leave: the business grows on its own, before any lead you buy.",IF(%s=%s,"Referrals exactly replace the customers who leave. Every new customer is growth, none of it free.","More customers leave each month than referrals bring. Growth has to be bought until the referral rate passes the churn.")))' % (gin, K['referralRate'], K['churnRate'], K['referralRate'], K['churnRate'])).font = f(); ws.cell(row=r, column=6).alignment = WRAP
    r += 1
ws.merge_cells(start_row=g_top, start_column=6, end_row=r - 1, end_column=6)
ws.cell(row=r, column=1, value='The gap after a year').font = f(True)
c = ws.cell(row=r, column=2, value='=IF(B%d="","",B%d-C%d)' % (r - 1, r - 1, r - 1)); c.number_format = '#,##0.0'; c.font = f(True); c.border = BOX
ws.cell(row=r, column=6, value='customers, what asking for referrals is worth').font = f(False, 9, '555555', True)
lc = LineChart(); lc.title = 'Customers over twelve months'; lc.style = 12
lc.add_data(Reference(ws, min_col=2, min_row=g_top - 1, max_col=3, max_row=r - 1), titles_from_data=True)
lc.set_categories(Reference(ws, min_col=1, min_row=g_top, max_row=r - 1)); lc.height = 8; lc.width = 18
ws.add_chart(lc, 'H%d' % (lever_top + 15))
fc = BarChart(); fc.type = 'bar'; fc.title = 'From reached to customers, a month'; fc.style = 10
frow = int(K['reach'][1:])
fc.add_data(Reference(ws, min_col=2, min_row=frow, max_row=frow + 4), titles_from_data=False)
fc.set_categories(Reference(ws, min_col=1, min_row=frow, max_row=frow + 4)); fc.legend = None; fc.height = 7; fc.width = 18
ws.add_chart(fc, 'H%d' % frow)
ws.freeze_panes = 'A5'

# ---------------------------------------------------------------- Planets
BANDS = {b['n']: b for b in BOOK['bands']}
PROG = {}   # planet id -> sheet name (for Start Here formulas)
for p in BOOK['planets']:
    name = SHEET[p['id']]; PROG[p['id']] = name
    ws = wb.create_sheet(name)
    ws.column_dimensions['A'].width = 58; ws.column_dimensions['B'].width = 40; ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 3; ws.column_dimensions['E'].width = 7; ws.column_dimensions['F'].width = 7; ws.column_dimensions['G'].width = 7; ws.column_dimensions['H'].width = 13
    kind = 'the sun' if p.get('sun') else 'a lead getter' if p.get('getter') else 'core four: %s, %s' % (p['who'].lower(), p['how'].lower())
    r = title_row(ws, 1, '%s (%s)' % (p['name'], kind), p['chapter'] + '. ' + p['blurb'], span=8)
    r = legend(ws, r, span=8)
    ws.add_data_validation(yes)
    # progress block
    head(ws, r, ['Band', 'Levels done', 'Of', '', 'Band', 'Done items', 'Of', 'Status']); r += 1
    prog_top = r
    for n in range(1, 6):
        ws.cell(row=r, column=1, value='%d. %s: %s' % (n, BANDS[n]['name'], BANDS[n]['means'])).font = f()
        c = ws.cell(row=r, column=2, value='=COUNTIFS($E$1:$E$600,%d,$H$1:$H$600,"Done")' % n); c.font = f(True); c.border = BOX
        c = ws.cell(row=r, column=3, value='=COUNTIF($E$1:$E$600,%d)' % n); c.font = f(); c.border = BOX
        r += 1
    ws.cell(row=r, column=1, value='This planet').font = f(True)
    c = ws.cell(row=r, column=2, value='=SUM(B%d:B%d)' % (prog_top, r - 1)); c.font = f(True); c.border = BOX
    c = ws.cell(row=r, column=3, value='=SUM(C%d:C%d)' % (prog_top, r - 1)); c.font = f(); c.border = BOX
    ws.cell(row=r, column=5, value='The columns E to H on the right are the workings for each level: its band, how many of its items are done, how many there are, and its status. Leave them be.').font = f(False, 8, '777777', True)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8); ws.cell(row=r, column=5).alignment = WRAP; ws.row_dimensions[r].height = 40
    r += 2
    levels = [l for l in BOOK['levels'] if l['planet'] == p['id']]
    for n in range(1, 6):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = ws.cell(row=r, column=1, value='Band %d: %s. %s. Payoff: %s.' % (n, BANDS[n]['name'], BANDS[n]['means'], BANDS[n]['payoff'].lower())); c.font = f(True, 11, 'FFFFFF'); c.fill = NAVY
        for col in range(1, 9): ws.cell(row=r, column=col).fill = NAVY
        r += 1
        for lv in [l for l in levels if l['band'] == n]:
            # title row
            c = ws.cell(row=r, column=1, value=lv['title']); c.font = f(True, 11); c.fill = GREY
            ws.cell(row=r, column=2, value='about %d min' % lv['minutes']).font = f(False, 9, '555555'); ws.cell(row=r, column=2).fill = GREY
            ws.cell(row=r, column=3, value=lv['chapter']).font = f(False, 9, '555555'); ws.cell(row=r, column=3).fill = GREY
            ws.cell(row=r, column=5, value=n).font = f(False, 9, '999999')
            title_r = r; r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            c = ws.cell(row=r, column=1, value=lv['says']); c.font = f(); c.alignment = WRAP
            ws.row_dimensions[r].height = max(30, 15 * (len(lv['says']) // 120 + 1)); r += 1
            conds = []
            # fields
            for fld in lv.get('fields', []):
                if fld['type'] == 'multi':
                    ws.cell(row=r, column=1, value=fld['label'] + ' (Y beside each that applies)').font = f(True); r += 1
                    top = r
                    for o in fld['options']:
                        ws.cell(row=r, column=1, value='    ' + o[1]).font = f()
                        c = ws.cell(row=r, column=2); inp(c); c.alignment = Alignment(horizontal='center'); yes.add(c)
                        r += 1
                    conds.append('(COUNTIF(B%d:B%d,"Y")>0)' % (top, r - 1))
                    continue
                ws.cell(row=r, column=1, value=fld['label']).font = f(); ws.cell(row=r, column=1).alignment = WRAP
                c = ws.cell(row=r, column=2); c.alignment = WRAP
                if fld['type'] == 'choice':
                    opts = [o[1] for o in fld['options']]
                    dv = DataValidation(type='list', formula1='"%s"' % ','.join(o.replace(',', ';') for o in opts), allow_blank=True); ws.add_data_validation(dv); dv.add(c)
                    inp(c); ws.cell(row=r, column=3, value='choose one: ' + ' / '.join(opts)).font = f(False, 9, '555555', True)
                elif fld['type'] == 'number':
                    inp(c, 'money' if fld['unit'] == 'dollars' else 'count')
                    ws.cell(row=r, column=3, value='for example: %s' % ('$%s' % fld['example'] if fld['unit'] == 'dollars' else fld['example'])).font = f(False, 9, '555555', True)
                else:
                    inp(c)
                    if fld.get('example') is not None: ws.cell(row=r, column=3, value='for example: ' + str(fld['example'])).font = f(False, 9, '555555', True)
                    if fld['type'] == 'long': ws.row_dimensions[r].height = 45
                ws.cell(row=r, column=3).alignment = WRAP
                if fld.get('hint'): ws.cell(row=r, column=3, value=(ws.cell(row=r, column=3).value or '') + ' ' + fld['hint'])
                if not fld.get('optional'): conds.append('(B%d<>"")' % r)
                fld['_cell'] = 'B%d' % r
                r += 1
            for item in lv.get('checklist', []):
                ws.cell(row=r, column=1, value='[ ] ' + item).font = f(); ws.cell(row=r, column=1).alignment = WRAP
                c = ws.cell(row=r, column=2); inp(c); c.alignment = Alignment(horizontal='center'); yes.add(c)
                ws.cell(row=r, column=3, value='Y when true').font = f(False, 9, '555555', True)
                conds.append('(B%d="Y")' % r); r += 1
            if lv.get('confirm'):
                ws.cell(row=r, column=1, value='Read it? Put Y here when you have.').font = f()
                c = ws.cell(row=r, column=2); inp(c); c.alignment = Alignment(horizontal='center'); yes.add(c)
                conds.append('(B%d="Y")' % r); r += 1
            for key in lv.get('needs', []):
                lab = [m for m in MACHINE_INPUTS if m[0] == key][0]
                ws.cell(row=r, column=1, value='From the Machine sheet: ' + lab[1]).font = f()
                c = ws.cell(row=r, column=2, value='=IF(%s="","not yet",%s)' % (MCELL[key], MCELL[key])); c.font = f(False, 10, LINK); c.border = BOX
                c.number_format = '0.0%' if lab[2] == 'rate' else '$#,##0' if lab[2] == 'money' else '#,##0'
                ws.cell(row=r, column=3, value='type it on the Machine sheet, it shows here').font = f(False, 9, '555555', True)
                conds.append('(%s<>"")' % MCELL[key]); r += 1
            # reading
            rd = lv.get('reading')
            if rd:
                cells = {fl['key']: fl['_cell'] for fl in lv.get('fields', []) if '_cell' in fl}
                ws.cell(row=r, column=1, value=rd['label']).font = f(True, 10, '1F2A44')
                c = ws.cell(row=r, column=2); c.font = f(True); c.border = BOX; c.alignment = WRAP
                if rd['kind'] == 'share':
                    c.value = '=IF(OR(%s="",%s=""),"not yet",IF(%s=0,"none",%s/%s))' % (cells[rd['of']], cells[rd['over']], cells[rd['over']], cells[rd['of']], cells[rd['over']]); c.number_format = '0%'
                elif rd['kind'] == 'sum':
                    cs = [cells[k] for k in rd['keys']]
                    c.value = '=IF(COUNT(%s)<%d,"not yet",SUM(%s))' % (','.join(cs), len(cs), ','.join(cs)); c.number_format = '#,##0" people"'
                elif rd['kind'] == 'sentence':
                    import re
                    keys = re.findall(r'\{(\w+)\}', rd['template'])
                    parts = re.split(r'\{\w+\}', rd['template'])
                    expr = '"%s"' % parts[0]
                    for i, k in enumerate(keys): expr += '&%s&"%s"' % (cells[k], parts[i + 1])
                    c.value = '=IF(COUNTA(%s)<%d,"fill every blank and it reads back here",%s)' % (','.join(cells[k] for k in keys), len(keys), expr)
                    ws.row_dimensions[r].height = 45
                elif rd['kind'] == 'costPer':
                    cs = [cells[k] for k in rd['keys']]
                    c.value = '=IF(OR(COUNT(%s)<%d,%s=""),"not yet",IF(%s=0,"none",SUM(%s)/%s))' % (','.join(cs), len(cs), cells[rd['over']], cells[rd['over']], ','.join(cs), cells[rd['over']]); c.number_format = '$#,##0.00'
                elif rd['kind'] == 'funnel':
                    ks = rd['keys']
                    c.value = 'each step as a share of the one before'
                    c.font = f(False, 9, '555555', True); c.border = Border()
                    r += 1
                    for i in range(1, len(ks)):
                        a, b = cells[ks[i - 1]], cells[ks[i]]
                        la = [fl['label'] for fl in lv['fields'] if fl['key'] == ks[i]][0]
                        ws.cell(row=r, column=1, value='    ' + la + ', as a share').font = f()
                        cc = ws.cell(row=r, column=2, value='=IF(OR(%s="",%s=""),"not yet",IF(%s=0,"none",%s/%s))' % (a, b, a, b, a)); cc.number_format = '0%'; cc.font = f(True); cc.border = BOX
                        r += 1
                    r -= 1
                r += 1
            ws.cell(row=r, column=1, value='You have leveled this when: ' + lv['checkpoint']).font = f(False, 9, '7A5A00', True); ws.cell(row=r, column=1).alignment = WRAP
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); ws.row_dimensions[r].height = 28
            # workings on the title row
            ws.cell(row=title_r, column=6, value='=0+%s' % '+'.join(conds)).font = f(False, 9, '999999')
            ws.cell(row=title_r, column=7, value=len(conds)).font = f(False, 9, '999999')
            c = ws.cell(row=title_r, column=8, value='=IF(F%d=G%d,"Done",IF(F%d>0,"In progress","Not yet"))' % (title_r, title_r, title_r)); c.font = f(True, 10); c.border = BOX; c.alignment = Alignment(horizontal='center')
            r += 2
    ws.conditional_formatting.add('H1:H%d' % r, CellIsRule(operator='equal', formula=['"Done"'], fill=GREEN))
    ws.conditional_formatting.add('H1:H%d' % r, CellIsRule(operator='equal', formula=['"In progress"'], fill=AMBER))
    ws.freeze_panes = 'A%d' % (prog_top - 1)

# ---------------------------------------------------------------- Log
ws = wb.create_sheet('Daily Log')
for col, w in zip('ABCDEFG', [14, 20, 12, 3, 44, 14, 40]): ws.column_dimensions[col].width = w
r = title_row(ws, 1, 'The rule of 100: your daily log', 'A hundred primary actions a day on your chosen planet: messages, minutes of content, reach-outs, or dollars. Log every day. A day you skip is a blank row, not a zero.', span=7)
r = legend(ws, r, span=7)
head(ws, r, ['Day', 'Planet', 'Count', '', 'The last thirty days', 'Value', '']); r += 1
log_top = r
today = datetime.date.today()
example = [100, 110, 95, None, 120, 100, 100, 130, 60, 100, 105, 100, 100, 115]
dvp = DataValidation(type='list', formula1='"Warm outreach,Free content,Cold outreach,Paid ads"', allow_blank=True); ws.add_data_validation(dvp)
ws.cell(row=r - 1, column=5).value = 'The last thirty days (from today)'
for i, cnt in enumerate(example):
    d = today - datetime.timedelta(days=len(example) - 1 - i)
    c = ws.cell(row=r, column=1, value=d); inp(c); c.number_format = 'yyyy-mm-dd'
    c = ws.cell(row=r, column=2, value='Warm outreach' if cnt is not None else None); inp(c); dvp.add(c)
    c = ws.cell(row=r, column=3, value=cnt); inp(c, 'count')
    r += 1
for _ in range(int(os.environ.get('LOG_ROWS', '200'))):
    c = ws.cell(row=r, column=1); inp(c); c.number_format = 'yyyy-mm-dd'
    c = ws.cell(row=r, column=2); inp(c); dvp.add(c)
    c = ws.cell(row=r, column=3); inp(c, 'count')
    r += 1
log_end = r - 1
ws.cell(row=log_top, column=7, value='These fourteen rows are example days. Replace them with yours; blank rows below are ready.').font = f(False, 9, 'A05A00', True); ws.cell(row=log_top, column=7).alignment = WRAP
stats = [('Days logged', '=COUNTIFS($A$%d:$A$%d,">="&TODAY()-29,$C$%d:$C$%d,"<>")' % (log_top, log_end, log_top, log_end)),
         ('Days at 100 or more', '=COUNTIFS($A$%d:$A$%d,">="&TODAY()-29,$C$%d:$C$%d,">=100")' % (log_top, log_end, log_top, log_end)),
         ('Actions in the last thirty days', '=SUMIFS($C$%d:$C$%d,$A$%d:$A$%d,">="&TODAY()-29)' % (log_top, log_end, log_top, log_end)),
         ('Average on a logged day', '=IF(F%d=0,"not yet",F%d/F%d)' % (log_top, log_top + 2, log_top))]
for i, (lab, fo) in enumerate(stats):
    ws.cell(row=log_top + i, column=5, value=lab).font = f(True)
    c = ws.cell(row=log_top + i, column=6, value=fo); c.font = f(); c.border = BOX; c.number_format = '#,##0'
bc = BarChart(); bc.title = 'Primary actions a day'; bc.style = 10
bc.add_data(Reference(ws, min_col=3, min_row=log_top, max_row=log_top + 29), titles_from_data=False)
bc.set_categories(Reference(ws, min_col=1, min_row=log_top, max_row=log_top + 29)); bc.legend = None; bc.height = 8; bc.width = 20
ws.add_chart(bc, 'E%d' % (log_top + 6))
ws.freeze_panes = 'A%d' % log_top

# ---------------------------------------------------------------- Start Here
ws = wb.create_sheet('Start Here', 0)
for col, w in zip('ABCDEFGHI', [46, 40, 3, 30, 12, 12, 12, 12, 12]): ws.column_dimensions[col].width = w
r = title_row(ws, 1, 'The Leads Ladder: %s by %s, as a workbook' % (BOOK['book']['title'], BOOK['book']['author']), 'The book restated as exercises. The lead magnet is the sun and everything orbits it; warm outreach, free content, cold outreach and paid ads are the core four; referrals, employees, agencies and affiliates are the lead getters. Every planet has the same five bands: Learn, Sketch, Do, Measure, Scale.', span=9)
r = legend(ws, r, span=9)
r += 1
head(ws, r, ['Start here: seven answers', 'Your answer', '', 'Why it is asked', '', '', '', '', '']); r += 1
ws.cell(row=r, column=1, value='These answers are the example (an invented piano teacher). Replace them with yours.').font = f(False, 9, 'A05A00', True); r += 1
START = {}
ex = {'sells': 'Piano lessons for adults who gave up as kids', 'serves': 'Busy adults in their 30s and 40s', 'customers': 'One to ten', 'budget': 'Some, and I could afford to lose it while I learn', 'team': 'Just me', 'goal': 20, 'first': 'Warm outreach: I know people'}
for fld in BOOK['start']['fields']:
    ws.cell(row=r, column=1, value=fld['label']).font = f(); ws.cell(row=r, column=1).alignment = WRAP
    c = ws.cell(row=r, column=2, value=ex[fld['key']]); c.alignment = WRAP
    if fld['type'] == 'choice':
        dv = DataValidation(type='list', formula1='"%s"' % ','.join(o[1].replace(',', ';') for o in fld['options']), allow_blank=True); ws.add_data_validation(dv); dv.add(c)
        c.value = c.value.replace(',', ';')
        inp(c)
    elif fld['type'] == 'number': inp(c, 'count')
    else: inp(c)
    ws.cell(row=r, column=4, value=fld.get('hint', 'for example: %s' % fld.get('example', ''))).font = f(False, 9, '555555', True); ws.cell(row=r, column=4).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
    ws.row_dimensions[r].height = 30
    START[fld['key']] = '$B$%d' % r
    r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
c = ws.cell(row=r, column=1, value=BOOK['start']['rule']); c.font = f(True, 10, '7A5A00'); c.alignment = WRAP; c.fill = GOLD; ws.row_dimensions[r].height = 32
r += 2
head(ws, r, ['Where you stand', 'Sheet', '', 'Applies?', 'Learn', 'Sketch', 'Do', 'Measure', 'Scale']); r += 1
grid_top = r
applies = {'always': 'TRUE', 'budget': '%s="Some; and I could afford to lose it while I learn"' % START['budget'], 'customers': 'OR(%s="One to ten",%s="More than ten")' % (START['customers'], START['customers']), 'team': 'OR(%s="A team",%s="Me; and I am hiring")' % (START['team'], START['team'])}
for pid in BOOK['order']:
    p = [x for x in BOOK['planets'] if x['id'] == pid][0]; sh = SHEET[pid]
    ws.cell(row=r, column=1, value=p['name'] + (' (the sun)' if p.get('sun') else '')).font = f(True)
    ws.cell(row=r, column=2, value=sh).font = f(False, 10, LINK); ws.cell(row=r, column=2).hyperlink = "#'%s'!A1" % sh
    c = ws.cell(row=r, column=4, value='=IF(%s,"yes","not yet")' % applies[p['appliesWhen']]); c.font = f(); c.border = BOX; c.alignment = Alignment(horizontal='center')
    if p.get('dimMessage'): c.comment = Comment(p['dimMessage'], 'Leads Ladder')
    for b in range(1, 6):
        c = ws.cell(row=r, column=4 + b, value='=COUNTIFS(\'%s\'!$E$1:$E$600,%d,\'%s\'!$H$1:$H$600,"Done")&" of "&COUNTIF(\'%s\'!$E$1:$E$600,%d)' % (sh, b, sh, sh, b)); c.font = f(); c.border = BOX; c.alignment = Alignment(horizontal='center')
    r += 1
grid_end = r - 1
ws.cell(row=r, column=1, value='Ring cleared? (every planet that applies has finished the band)').font = f(True); ws.cell(row=r, column=1).alignment = WRAP
for b in range(1, 6):
    parts = []
    for pid in BOOK['order']:
        p = [x for x in BOOK['planets'] if x['id'] == pid][0]; sh = SHEET[pid]
        parts.append('(%s)*(COUNTIFS(\'%s\'!$E$1:$E$600,%d,\'%s\'!$H$1:$H$600,"Done")<COUNTIF(\'%s\'!$E$1:$E$600,%d))' % (applies[p['appliesWhen']], sh, b, sh, sh, b))
    c = ws.cell(row=r, column=4 + b, value='=IF(SUM(%s)=0,"cleared","open")' % ','.join(parts)); c.font = f(True); c.border = BOX; c.alignment = Alignment(horizontal='center')
    ws.conditional_formatting.add(c.coordinate, CellIsRule(operator='equal', formula=['"cleared"'], fill=GOLD))
ring_r = r; r += 2
# levels done per planet, for the chart
head(ws, r, ['Levels done, for the picture', 'Done', '', 'Of', '', '', '', '', '']); r += 1
ch_top = r
for pid in BOOK['order']:
    sh = SHEET[pid]; p = [x for x in BOOK['planets'] if x['id'] == pid][0]
    ws.cell(row=r, column=1, value=p['name']).font = f()
    c = ws.cell(row=r, column=2, value='=COUNTIF(\'%s\'!$H$1:$H$600,"Done")' % sh); c.font = f(); c.border = BOX
    c = ws.cell(row=r, column=4, value='=COUNT(\'%s\'!$E$1:$E$600)' % sh); c.font = f(); c.border = BOX
    r += 1
ws.cell(row=r, column=1, value='The whole book').font = f(True)
c = ws.cell(row=r, column=2, value='=SUM(B%d:B%d)' % (ch_top, r - 1)); c.font = f(True); c.border = BOX
c = ws.cell(row=r, column=4, value='=SUM(D%d:D%d)' % (ch_top, r - 1)); c.font = f(True); c.border = BOX
ws.cell(row=r, column=5, value='=IF(D%d=0,"",B%d/D%d)' % (r, r, r)).number_format = '0%'; ws.cell(row=r, column=5).font = f(True)
total_r = r; r += 1
pc = BarChart(); pc.type = 'bar'; pc.title = 'Levels done by planet'; pc.style = 10
pc.add_data(Reference(ws, min_col=2, min_row=ch_top, max_row=ch_top + 8), titles_from_data=False)
pc.set_categories(Reference(ws, min_col=1, min_row=ch_top, max_row=ch_top + 8)); pc.legend = None; pc.height = 8; pc.width = 16
ws.add_chart(pc, 'F%d' % ch_top)
r += 1
head(ws, r, ['The Machine, in brief', 'Value', '', 'Goal', '', '', '', '', '']); r += 1
ws.cell(row=r, column=1, value='Engaged leads a week at your numbers').font = f()
c = ws.cell(row=r, column=2, value="=IF('Machine'!%s=\"\",\"not yet\",'Machine'!%s)" % (K['perWeek'], K['perWeek'])); c.font = f(False, 10, LINK); c.number_format = '#,##0.0'; c.border = BOX
c = ws.cell(row=r, column=4, value='=IF(%s="","set the goal above",IF(B%d="not yet","fill the Machine",IF(B%d>=%s,"hit","short by "&TEXT(%s-B%d,"0.0"))))' % (START['goal'], r, r, START['goal'], START['goal'], r)); c.font = f(); c.border = BOX
r += 1
ws.cell(row=r, column=1, value='Customers a month').font = f()
c = ws.cell(row=r, column=2, value="=IF('Machine'!%s=\"\",\"not yet\",'Machine'!%s)" % (K['customers'], K['customers'])); c.font = f(False, 10, LINK); c.number_format = '#,##0.0'; c.border = BOX; r += 1
ws.cell(row=r, column=1, value='Lifetime gross profit over cost per customer').font = f()
c = ws.cell(row=r, column=2, value="=IF('Machine'!%s=\"\",\"not yet\",'Machine'!%s)" % (K['ratio'], K['ratio'])); c.font = f(False, 10, LINK); c.number_format = '0.0"x"'; c.border = BOX
ws.cell(row=r, column=4, value='the book says three or more').font = f(False, 9, '555555', True); r += 1
ws.cell(row=r, column=1, value='First thirty days over cost per customer').font = f()
c = ws.cell(row=r, column=2, value="=IF('Machine'!%s=\"\",\"not yet\",'Machine'!%s)" % (K['payback'], K['payback'])); c.font = f(False, 10, LINK); c.number_format = '0.0"x"'; c.border = BOX
ws.cell(row=r, column=4, value='the book says two or more').font = f(False, 9, '555555', True); r += 2
head(ws, r, ['The book, chapter by chapter', 'Sheet', '', 'What it is', '', '', '', '', '']); r += 1
rows = [('Chapters 1 and 3', 'Start Here', 'The core four, and which one first'), ('Chapter 12', 'Start Here', 'The rule of 100, and the one-year promise')]
for pid in BOOK['order']:
    p = [x for x in BOOK['planets'] if x['id'] == pid][0]
    rows.append((p['chapter'].split(',')[0], SHEET[pid], p['name'] + (' (the sun)' if p.get('sun') else ' (a lead getter)' if p.get('getter') else ' (core four)')))
rows.append(('The numbers', 'Machine', 'KPIs, levers, the two checks, referrals against churn')); rows.append(('The daily hundred', 'Daily Log', 'The rule of 100, day by day')); rows.append(('The words', 'Words', 'The glossary'))
for ch_, sh, what in rows:
    ws.cell(row=r, column=1, value=ch_).font = f()
    c = ws.cell(row=r, column=2, value=sh); c.font = f(False, 10, LINK); c.hyperlink = "#'%s'!A1" % sh
    ws.cell(row=r, column=4, value=what).font = f(); r += 1
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
c = ws.cell(row=r, column=1, value=BOOK['source'] + ' ' + BOOK['confidenceNote'] + ' Every example number in this workbook is invented.'); c.font = f(False, 8, '777777', True); c.alignment = WRAP; ws.row_dimensions[r].height = 40
ws.freeze_panes = 'A5'

# ---------------------------------------------------------------- Words
ws = wb.create_sheet('Words')
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 100
r = title_row(ws, 1, 'The words', 'Every term the book leans on, in one sentence.', span=2)
head(ws, r, ['Term', 'Meaning']); r += 1
for w in BOOK['words']:
    ws.cell(row=r, column=1, value=w[0]).font = f(True); ws.cell(row=r, column=2, value=w[1]).font = f(); ws.cell(row=r, column=2).alignment = WRAP; r += 1

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.page_setup.orientation = 'landscape'; s.page_setup.fitToWidth = 1; s.sheet_properties.pageSetUpPr.fitToPage = True; s.page_setup.fitToHeight = 0
wb.save(OUT)
print('wrote', os.path.relpath(OUT, os.path.join(ROOT, '..')), 'with', len(wb.worksheets), 'sheets')
