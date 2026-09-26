#!/usr/bin/env python3
"""marketing/tools/sheet.py, the Marketing Scoreboard as one spreadsheet (MD-010).

The same three logs and the same readings as the app, as a workbook a person
can open in Google Sheets or Excel and check cell by cell: every number on the
Scoreboard is a formula over the Posts, People and Touches tabs. Nothing is
hidden and nothing runs in the background.

    python3 marketing/tools/sheet.py            writes marketing/Marketing-Scoreboard.xlsx
"""
import datetime as dt
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(__file__), '..', 'Marketing-Scoreboard.xlsx')
TODAY = dt.date.today()
POST_ROWS, PEOPLE_ROWS, TOUCH_ROWS = int(os.environ.get('SHEET_POSTS', 250)), int(os.environ.get('SHEET_PEOPLE', 150)), int(os.environ.get('SHEET_TOUCHES', 500))

FONT = 'Arial'
F = lambda **k: Font(name=FONT, **k)
HEAD_FILL = PatternFill('solid', start_color='1F3A5F')
INPUT_FILL = PatternFill('solid', start_color='FFF8DC')
CALC_FILL = PatternFill('solid', start_color='EEF2F7')
SECTION_FILL = PatternFill('solid', start_color='DCE6F2')
GOOD = PatternFill('solid', start_color='D9EAD3'); WATCH = PatternFill('solid', start_color='FFF2CC'); BAD = PatternFill('solid', start_color='F4CCCC')
THIN = Side(style='thin', color='C9D1DB')
BORDER = Border(bottom=THIN)

CHANNELS = ['Instagram', 'TikTok', 'YouTube', 'LinkedIn', 'X', 'Threads', 'Facebook', 'Newsletter', 'Podcast', 'Blog', 'Other']
FORMATS = ['Short video', 'Long video', 'Carousel', 'Image', 'Text post', 'Story', 'Live', 'Email', 'Article', 'Episode']
ASKS = ['No ask', 'Follow', 'Comment a word', 'DM me', 'Click the link', 'Book a call', 'Buy', 'Share it']
STAGES = [('Stranger', 'never', 'Comment on something of theirs first. No pitch.'),
          ('Follower', 30, 'Reply to their content, or a DM about something they posted.'),
          ('In conversation', 7, 'A DM or a message that asks a question about them.'),
          ('Lead', 3, 'An email or a voice note. Offer the next step plainly.'),
          ('Call booked', 1, 'Confirm the call the day before. Send what to bring.'),
          ('Client', 30, 'A check-in call or email. Ask what changed.'),
          ('Advocate', 90, 'Thank them. Ask who else should hear about this.'),
          ('Not now', 'never', 'Leave it. A kind note in six months if they said so.')]
LANES = ['Warm (they know me)', 'Cold (they do not)', 'From content', 'Paid', 'Referral']
KINDS = ['Comment', 'DM', 'Email', 'Text', 'Call', 'Voice note', 'Met in person', 'Gift or card']
OUTCOMES = ['No reply yet', 'They replied', 'Call booked', 'They bought', 'Not now']
WHO = ['Me', 'Them']
TARGETS = [('Posts a week', 5), ('Reach-outs a week', 25), ('Conversations a week', 10), ('Leads a week', 3), ('Calls booked a week', 2)]

wb = Workbook()

def head(ws, row, labels, widths=None):
    for i, lbl in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lbl)
        c.font = F(bold=True, color='FFFFFF'); c.fill = HEAD_FILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = F(bold=True, size=16)
    if sub: ws['A2'] = sub; ws['A2'].font = F(italic=True, color='555555')

def section(ws, row, text, span=6):
    c = ws.cell(row=row, column=1, value=text); c.font = F(bold=True, size=12); c.fill = SECTION_FILL
    for col in range(2, span + 1): ws.cell(row=row, column=col).fill = SECTION_FILL

def label(ws, row, col, text, bold=False):
    c = ws.cell(row=row, column=col, value=text); c.font = F(bold=bold); return c

def calc(ws, row, col, formula, fmt=None):
    c = ws.cell(row=row, column=col, value=formula); c.font = F(); c.fill = CALC_FILL
    if fmt: c.number_format = fmt
    return c

# ---------------------------------------------------------------- Lists
L = wb.active; L.title = 'Lists'
title(L, 'Lists and targets', 'The dropdowns read these. Change a target here and the Scoreboard follows. Add a channel or a format by typing it below the last one.')
head(L, 4, ['Channels', 'Formats', 'Asks', 'Targets', 'Target', '', 'Stage', 'Days between touches', 'How to reach out at this stage', 'How we met (lanes)', 'Touch kinds', 'What happened', 'Who started'],
     [16, 14, 16, 22, 9, 3, 16, 12, 54, 22, 14, 16, 10])
for i, v in enumerate(CHANNELS): L.cell(row=5 + i, column=1, value=v).font = F()
for i, v in enumerate(FORMATS): L.cell(row=5 + i, column=2, value=v).font = F()
for i, v in enumerate(ASKS): L.cell(row=5 + i, column=3, value=v).font = F()
for i, (k, v) in enumerate(TARGETS):
    L.cell(row=5 + i, column=4, value=k).font = F()
    c = L.cell(row=5 + i, column=5, value=v); c.font = F(color='0000FF'); c.fill = INPUT_FILL
for i, (s, d, h) in enumerate(STAGES):
    L.cell(row=5 + i, column=7, value=s).font = F()
    c = L.cell(row=5 + i, column=8, value=d); c.font = F(color='0000FF'); c.fill = INPUT_FILL
    L.cell(row=5 + i, column=9, value=h).font = F()
for i, v in enumerate(LANES): L.cell(row=5 + i, column=10, value=v).font = F()
for i, v in enumerate(KINDS): L.cell(row=5 + i, column=11, value=v).font = F()
for i, v in enumerate(OUTCOMES): L.cell(row=5 + i, column=12, value=v).font = F()
for i, v in enumerate(WHO): L.cell(row=5 + i, column=13, value=v).font = F()
L['E5'].comment = Comment('Blue cells on a cream background are yours to change. Everything else is a formula.', 'Scoreboard')
L['H5'].comment = Comment('Days between touches for people at this stage. "never" means the sheet will not nudge you about them.', 'Scoreboard')
L.freeze_panes = 'A5'

# ---------------------------------------------------------------- Posts
P = wb.create_sheet('Posts')
PCOLS = ['Date', 'Channel', 'Format', 'Topic', 'Hook (the first line)', 'Ask', 'Link', 'Minutes to make', 'Reach', 'Likes', 'Comments', 'Shares', 'Saves', 'Link clicks', 'DMs it started', 'New followers', 'Leads', 'Results typed on', 'Notes', 'Engagements', 'Engagement rate', 'Week starting']
head(P, 1, PCOLS, [11, 12, 12, 16, 40, 14, 18, 9, 9, 8, 9, 8, 8, 9, 9, 9, 7, 12, 24, 11, 11, 11])
P['I1'].comment = Comment('Views, reach or impressions, from the platform. Pick one and use it every time. Leave blank until you have checked: blank is "not yet", never zero.', 'Scoreboard')
P['T1'].comment = Comment('Likes + comments + shares + saves. Blank until at least one of those is typed.', 'Scoreboard')
for r in range(2, POST_ROWS + 2):
    P.cell(row=r, column=1).number_format = 'yyyy-mm-dd'; P.cell(row=r, column=18).number_format = 'yyyy-mm-dd'
    calc(P, r, 20, f'=IF(COUNT(J{r}:M{r})=0,"",SUM(J{r}:M{r}))')
    calc(P, r, 21, f'=IF(OR(T{r}="",I{r}="",I{r}=0),"",T{r}/I{r})', '0.0%')
    calc(P, r, 22, f'=IF(A{r}="","",A{r}-WEEKDAY(A{r},2)+1)', 'yyyy-mm-dd')
last = POST_ROWS + 1
for col, rng in ((2, f'=Lists!$A$5:$A$40'), (3, f'=Lists!$B$5:$B$40'), (6, f'=Lists!$C$5:$C$40')):
    dv = DataValidation(type='list', formula1=rng, allow_blank=True); dv.error = 'Pick one from the list, or add it on the Lists tab.'; P.add_data_validation(dv)
    dv.add(f'{get_column_letter(col)}2:{get_column_letter(col)}{last}')
dvd = DataValidation(type='date', allow_blank=True); P.add_data_validation(dvd); dvd.add(f'A2:A{last}')
P.conditional_formatting.add(f'A2:S{last}', FormulaRule(formula=[f'AND($A2<>"",$I2="",$A2<=TODAY()-3)'], fill=PatternFill('solid', start_color='FFF2CC')))
P.freeze_panes = 'B2'; P.auto_filter.ref = f'A1:V{last}'

# ---------------------------------------------------------------- People
H = wb.create_sheet('People')
HCOLS = ['Name', 'Company', 'Role', 'Email', 'Phone', 'Platform', 'Handle or link', 'How we met', 'Stage', 'Became a lead on', 'Call booked on', 'Became a client on', 'Paid ($)', 'Days between touches', 'Contact them on', 'Tags', 'Came from this post', 'Notes', 'Added on',
         'Last touch', 'Rhythm (days)', 'Next touch', 'Status', 'Days over', 'How to reach out', 'sort key']
head(H, 1, HCOLS, [18, 16, 14, 22, 12, 11, 20, 18, 15, 12, 12, 12, 9, 10, 12, 16, 24, 24, 11, 11, 9, 11, 9, 8, 46, 8])
H['N1'].comment = Comment('Blank uses the stage\'s usual rhythm from the Lists tab. Type 0 to never be nudged about this person.', 'Scoreboard')
H['O1'].comment = Comment('A date here overrides the rhythm. Clear it to go back to the rhythm.', 'Scoreboard')
H['J1'].comment = Comment('The Scoreboard counts leads, calls and clients from these three dates, so type the day it happened.', 'Scoreboard')
T_DATES, T_NAMES = f'Touches!$A$2:$A${TOUCH_ROWS + 1}', f'Touches!$B$2:$B${TOUCH_ROWS + 1}'
for r in range(2, PEOPLE_ROWS + 2):
    for col in (10, 11, 12, 15, 19): H.cell(row=r, column=col).number_format = 'yyyy-mm-dd'
    H.cell(row=r, column=13).number_format = '$#,##0'
    calc(H, r, 20, f'=IF(A{r}="","",IF(COUNTIF({T_NAMES},A{r})=0,"",SUMPRODUCT(MAX(({T_NAMES}=A{r})*{T_DATES}))))', 'yyyy-mm-dd')
    calc(H, r, 21, f'=IF(A{r}="","",IF(N{r}<>"",N{r},IFERROR(INDEX(Lists!$H$5:$H$12,MATCH(I{r},Lists!$G$5:$G$12,0)),"")))')
    calc(H, r, 22, f'=IF(A{r}="","",IF(O{r}<>"",O{r},IF(OR(U{r}="",U{r}="never",U{r}=0),"",IF(T{r}="",TODAY(),T{r}+U{r}))))', 'yyyy-mm-dd')
    calc(H, r, 23, f'=IF(A{r}="","",IF(V{r}="","no rhythm",IF(V{r}<TODAY(),"Overdue",IF(V{r}=TODAY(),"Today",IF(V{r}<=TODAY()+7,"Soon","Later")))))')
    calc(H, r, 24, f'=IF(OR(A{r}="",V{r}=""),"",TODAY()-V{r})')
    calc(H, r, 25, f'=IF(A{r}="","",IFERROR(INDEX(Lists!$I$5:$I$12,MATCH(I{r},Lists!$G$5:$G$12,0)),""))')
    calc(H, r, 26, f'=IF(AND(X{r}<>"",X{r}>=0),X{r}+ROW()/100000,"")')
hl = PEOPLE_ROWS + 1
for col, rng in ((8, '=Lists!$J$5:$J$12'), (9, '=Lists!$G$5:$G$12')):
    dv = DataValidation(type='list', formula1=rng, allow_blank=True); H.add_data_validation(dv); dv.add(f'{get_column_letter(col)}2:{get_column_letter(col)}{hl}')
H.conditional_formatting.add(f'W2:W{hl}', CellIsRule(operator='equal', formula=['"Overdue"'], fill=BAD))
H.conditional_formatting.add(f'W2:W{hl}', CellIsRule(operator='equal', formula=['"Today"'], fill=WATCH))
H.column_dimensions['Z'].hidden = True
H.freeze_panes = 'B2'; H.auto_filter.ref = f'A1:Y{hl}'

# ---------------------------------------------------------------- Touches
T = wb.create_sheet('Touches')
head(T, 1, ['Date', 'Person', 'Kind', 'Lane', 'Who started it', 'What happened', 'Note'], [11, 20, 12, 20, 12, 14, 50])
T['B1'].comment = Comment('Pick the name exactly as it is on the People tab; the dropdown reads that tab.', 'Scoreboard')
for r in range(2, TOUCH_ROWS + 2):
    T.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
tl = TOUCH_ROWS + 1
for col, rng in ((2, f'=People!$A$2:$A${hl}'), (3, '=Lists!$K$5:$K$12'), (4, '=Lists!$J$5:$J$12'), (5, '=Lists!$M$5:$M$6'), (6, '=Lists!$L$5:$L$9')):
    dv = DataValidation(type='list', formula1=rng, allow_blank=True); T.add_data_validation(dv); dv.add(f'{get_column_letter(col)}2:{get_column_letter(col)}{tl}')
T.freeze_panes = 'A2'; T.auto_filter.ref = f'A1:G{tl}'

# ---------------------------------------------------------------- Scoreboard
S = wb.create_sheet('Scoreboard', 0)
S.column_dimensions['A'].width = 30
for c in 'BCDEFGH': S.column_dimensions[c].width = 16
title(S, 'Marketing Scoreboard', 'What gets measured gets managed. Every number here is a formula over the Posts, People and Touches tabs. Blank means not entered, never zero.')
label(S, 4, 1, 'Period from', True); c = S['B4']; c.value = '=TODAY()-29'; c.number_format = 'yyyy-mm-dd'; c.font = F(color='0000FF'); c.fill = INPUT_FILL
label(S, 5, 1, 'Period to', True); c = S['B5']; c.value = '=TODAY()'; c.number_format = 'yyyy-mm-dd'; c.font = F(color='0000FF'); c.fill = INPUT_FILL
label(S, 4, 3, 'Type dates over these two cells to read any period. The columns to the right compare it with the period of the same length just before.')
label(S, 6, 1, 'The period before, from'); calc(S, 6, 2, '=B4-(B5-B4+1)', 'yyyy-mm-dd'); label(S, 6, 3, 'to'); calc(S, 6, 4, '=B4-1', 'yyyy-mm-dd')
FROM, TO, PFROM, PTO = '$B$4', '$B$5', '$B$6', '$D$6'

# This week
section(S, 8, 'This week against the targets', 6)
label(S, 9, 1, 'Week starting', True); calc(S, 9, 2, '=TODAY()-WEEKDAY(TODAY(),2)+1', 'yyyy-mm-dd'); label(S, 9, 3, 'to'); calc(S, 9, 4, '=B9+6', 'yyyy-mm-dd')
label(S, 9, 5, 'Share of the week gone'); calc(S, 9, 6, '=(TODAY()-B9+1)/7', '0%')
WS, WE = '$B$9', '$D$9'
for i, h in enumerate(['', 'Target', 'So far', 'Should be by now', 'Still to do', 'Status'], 1):
    c = S.cell(row=10, column=i, value=h); c.font = F(bold=True); c.border = BORDER
week_rows = [
    ('Posts', 'Lists!$E$5', f'=COUNTIFS(Posts!$A$2:$A${last},">="&{WS},Posts!$A$2:$A${last},"<="&{WE})'),
    ('Reach-outs (touches I started)', 'Lists!$E$6', f'=COUNTIFS({T_DATES},">="&{WS},{T_DATES},"<="&{WE},Touches!$E$2:$E${tl},"Me")'),
    ('Conversations (they replied or wrote first)', 'Lists!$E$7', f'=COUNTIFS({T_DATES},">="&{WS},{T_DATES},"<="&{WE},Touches!$E$2:$E${tl},"Them")+COUNTIFS({T_DATES},">="&{WS},{T_DATES},"<="&{WE},Touches!$E$2:$E${tl},"Me",Touches!$F$2:$F${tl},"They replied")'),
    ('Leads (became a lead on)', 'Lists!$E$8', f'=COUNTIFS(People!$J$2:$J${hl},">="&{WS},People!$J$2:$J${hl},"<="&{WE})'),
    ('Calls booked (call booked on)', 'Lists!$E$9', f'=COUNTIFS(People!$K$2:$K${hl},">="&{WS},People!$K$2:$K${hl},"<="&{WE})'),
]
for i, (name, tgt, formula) in enumerate(week_rows):
    r = 11 + i
    label(S, r, 1, name); calc(S, r, 2, f'={tgt}'); calc(S, r, 3, formula); calc(S, r, 4, f'=ROUND(B{r}*$F$9,1)', '0.0'); calc(S, r, 5, f'=MAX(0,B{r}-C{r})')
    calc(S, r, 6, f'=IF(C{r}>=B{r},"Hit",IF(C{r}>=D{r},"On pace",IF(C{r}>=D{r}*0.6,"A little behind","Behind")))')
S.conditional_formatting.add('F11:F15', CellIsRule(operator='equal', formula=['"Hit"'], fill=GOOD))
S.conditional_formatting.add('F11:F15', CellIsRule(operator='equal', formula=['"On pace"'], fill=GOOD))
S.conditional_formatting.add('F11:F15', CellIsRule(operator='equal', formula=['"A little behind"'], fill=WATCH))
S.conditional_formatting.add('F11:F15', CellIsRule(operator='equal', formula=['"Behind"'], fill=BAD))

# The numbers, this period against the one before
section(S, 17, 'The numbers for the period, against the period before', 6)
for i, h in enumerate(['', 'This period', 'Before', 'Change', 'Change %', 'What it is'], 1):
    c = S.cell(row=18, column=i, value=h); c.font = F(bold=True); c.border = BORDER
def posts_in(fr, to): return f'Posts!$A$2:$A${last},">="&{fr},Posts!$A$2:$A${last},"<="&{to}'
def touches_in(fr, to): return f'{T_DATES},">="&{fr},{T_DATES},"<="&{to}'
def sum_or_blank(col, fr, to):   # blank when nobody typed that column in the period
    return f'=IF(COUNTIFS(Posts!${col}$2:${col}${last},"<>",{posts_in(fr, to)})=0,"",SUMIFS(Posts!${col}$2:${col}${last},{posts_in(fr, to)}))'
def ratio(a, b): return f'=IF(OR({a}="",{b}="",{b}=0),"",{a}/{b})'
NUM = [
    ('Posts', lambda fr, to: f'=COUNTIFS({posts_in(fr, to)})', '0', 'Rows on Posts with a date in the period.'),
    ('Posts with results typed', lambda fr, to: f'=COUNTIFS(Posts!$I$2:$I${last},"<>",{posts_in(fr, to)})', '0', 'Of those, how many have Reach typed.'),
    ('Reach', lambda fr, to: sum_or_blank('I', fr, to), '#,##0', 'Reach added up. Blank until a post in the period has it.'),
    ('Engagements', lambda fr, to: sum_or_blank('T', fr, to), '#,##0', 'Likes + comments + shares + saves.'),
    ('Engagement rate', None, '0.0%', 'Engagements over reach. Above 2% is fine, above 5% is good.'),
    ('Saves', lambda fr, to: sum_or_blank('M', fr, to), '#,##0', 'Saves are worth more than likes.'),
    ('New followers', lambda fr, to: sum_or_blank('P', fr, to), '#,##0', 'From what you typed on each post.'),
    ('Leads credited to posts', lambda fr, to: sum_or_blank('Q', fr, to), '0', 'The Leads column on Posts.'),
    ('Content hours', lambda fr, to: f'=IF(COUNTIFS(Posts!$H$2:$H${last},"<>",{posts_in(fr, to)})=0,"",SUMIFS(Posts!$H$2:$H${last},{posts_in(fr, to)})/60)', '0.0', 'Minutes to make, added up, over 60.'),
    ('Reach-outs', lambda fr, to: f'=COUNTIFS({touches_in(fr, to)},Touches!$E$2:$E${tl},"Me")', '0', 'Touches you started.'),
    ('Conversations', lambda fr, to: f'=COUNTIFS({touches_in(fr, to)},Touches!$E$2:$E${tl},"Them")+COUNTIFS({touches_in(fr, to)},Touches!$E$2:$E${tl},"Me",Touches!$F$2:$F${tl},"They replied")', '0', 'Touches they started, plus yours marked "They replied".'),
    ('Reply rate', None, '0.0%', 'Conversations over reach-outs. Above 30% warm, above 10% cold.'),
    ('Leads', lambda fr, to: f'=COUNTIFS(People!$J$2:$J${hl},">="&{fr},People!$J$2:$J${hl},"<="&{to})', '0', 'People whose "Became a lead on" is in the period.'),
    ('Calls booked', lambda fr, to: f'=COUNTIFS(People!$K$2:$K${hl},">="&{fr},People!$K$2:$K${hl},"<="&{to})', '0', 'People whose "Call booked on" is in the period.'),
    ('Clients won', lambda fr, to: f'=COUNTIFS(People!$L$2:$L${hl},">="&{fr},People!$L$2:$L${hl},"<="&{to})', '0', 'People whose "Became a client on" is in the period.'),
    ('Revenue agreed', lambda fr, to: f'=IF(COUNTIFS(People!$L$2:$L${hl},">="&{fr},People!$L$2:$L${hl},"<="&{to})=0,"",SUMIFS(People!$M$2:$M${hl},People!$L$2:$L${hl},">="&{fr},People!$L$2:$L${hl},"<="&{to}))', '$#,##0', 'What those clients paid.'),
    ('Hours a lead', None, '0.0', 'Content hours over leads. Lower is better.'),
]
NROW = {}
for i, (name, fn, fmt, what) in enumerate(NUM):
    r = 19 + i; NROW[name] = r
    label(S, r, 1, name)
    if fn:
        calc(S, r, 2, fn(FROM, TO), fmt); calc(S, r, 3, fn(PFROM, PTO), fmt)
    label(S, r, 6, what).font = F(color='555555', size=9)
    S.cell(row=r, column=2).number_format = fmt; S.cell(row=r, column=3).number_format = fmt
for col in ('B', 'C'):
    r = NROW['Engagement rate']; calc(S, r, 2 if col == 'B' else 3, ratio(f'{col}{NROW["Engagements"]}', f'{col}{NROW["Reach"]}'), '0.0%')
    r = NROW['Reply rate']; calc(S, r, 2 if col == 'B' else 3, ratio(f'{col}{NROW["Conversations"]}', f'{col}{NROW["Reach-outs"]}'), '0.0%')
    r = NROW['Hours a lead']; calc(S, r, 2 if col == 'B' else 3, f'=IF(OR({col}{NROW["Content hours"]}="",{col}{NROW["Leads"]}=0),"",{col}{NROW["Content hours"]}/{col}{NROW["Leads"]})', '0.0')
for name, fn, fmt, what in NUM:
    r = NROW[name]
    calc(S, r, 4, f'=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})', fmt if fmt != '0' else '+0;-0;0')
    calc(S, r, 5, f'=IF(OR(B{r}="",C{r}="",C{r}=0),"",(B{r}-C{r})/ABS(C{r}))', '+0%;-0%;0%')
NEND = 19 + len(NUM) - 1

# The funnel
r0 = NEND + 2
section(S, r0, 'The funnel, each step as a share of the one before', 6)
for i, h in enumerate(['', 'This period', 'Share of the step before'], 1):
    c = S.cell(row=r0 + 1, column=i, value=h); c.font = F(bold=True); c.border = BORDER
FUN = [('Reached', 'Reach'), ('Engaged', 'Engagements'), ('Followed', 'New followers'), ('Talked to me', 'Conversations'), ('Became a lead', 'Leads'), ('Booked a call', 'Calls booked'), ('Became a client', 'Clients won')]
for i, (name, src) in enumerate(FUN):
    r = r0 + 2 + i
    label(S, r, 1, name); calc(S, r, 2, f'=B{NROW[src]}', '#,##0')
    if i: calc(S, r, 3, f'=IF(OR(B{r}="",B{r-1}="",B{r-1}=0),"",B{r}/B{r-1})', '0.0%')
FEND = r0 + 2 + len(FUN) - 1

# By channel, by ask, by format
def group_block(start, heading, list_ref, col_letter, n):
    section(S, start, heading, 7)
    for i, h in enumerate(['', 'Posts', 'Reach', 'Engagement rate', 'New followers', 'Leads', 'Leads a post'], 1):
        c = S.cell(row=start + 1, column=i, value=h); c.font = F(bold=True); c.border = BORDER
    for i in range(n):
        r = start + 2 + i
        calc(S, r, 1, f'=IF({list_ref}{5 + i}="","",{list_ref}{5 + i})').fill = PatternFill()
        crit = f'Posts!${col_letter}$2:${col_letter}${last},$A{r},{posts_in(FROM, TO)}'
        calc(S, r, 2, f'=IF($A{r}="","",COUNTIFS({crit}))', '0')
        calc(S, r, 3, f'=IF(OR($A{r}="",COUNTIFS(Posts!$I$2:$I${last},"<>",{crit})=0),"",SUMIFS(Posts!$I$2:$I${last},{crit}))', '#,##0')
        calc(S, r, 4, f'=IF(OR(C{r}="",C{r}=0,COUNTIFS(Posts!$T$2:$T${last},"<>",{crit})=0),"",SUMIFS(Posts!$T$2:$T${last},{crit})/C{r})', '0.0%')
        calc(S, r, 5, f'=IF(OR($A{r}="",COUNTIFS(Posts!$P$2:$P${last},"<>",{crit})=0),"",SUMIFS(Posts!$P$2:$P${last},{crit}))', '#,##0')
        calc(S, r, 6, f'=IF(OR($A{r}="",COUNTIFS(Posts!$Q$2:$Q${last},"<>",{crit})=0),"",SUMIFS(Posts!$Q$2:$Q${last},{crit}))', '0')
        calc(S, r, 7, f'=IF(OR(F{r}="",B{r}=0),"",F{r}/B{r})', '0.00')
    return start + 2 + n
r1 = group_block(FEND + 2, 'By channel', 'Lists!$A$', 'B', len(CHANNELS))
r2 = group_block(r1 + 1, 'By ask (what you asked people to do)', 'Lists!$C$', 'F', len(ASKS))
r3 = group_block(r2 + 1, 'By format', 'Lists!$B$', 'C', len(FORMATS))

# By lane
section(S, r3 + 1, 'Where the money comes from, by how you met them (all people, dates in the period)', 6)
for i, h in enumerate(['', 'People', 'Leads', 'Clients', 'Revenue'], 1):
    c = S.cell(row=r3 + 2, column=i, value=h); c.font = F(bold=True); c.border = BORDER
for i in range(len(LANES)):
    r = r3 + 3 + i
    calc(S, r, 1, f'=Lists!$J${5 + i}').fill = PatternFill()
    calc(S, r, 2, f'=COUNTIFS(People!$H$2:$H${hl},$A{r},People!$A$2:$A${hl},"<>")', '0')
    calc(S, r, 3, f'=COUNTIFS(People!$H$2:$H${hl},$A{r},People!$J$2:$J${hl},">="&{FROM},People!$J$2:$J${hl},"<="&{TO})', '0')
    calc(S, r, 4, f'=COUNTIFS(People!$H$2:$H${hl},$A{r},People!$L$2:$L${hl},">="&{FROM},People!$L$2:$L${hl},"<="&{TO})', '0')
    calc(S, r, 5, f'=IF(D{r}=0,"",SUMIFS(People!$M$2:$M${hl},People!$H$2:$H${hl},$A{r},People!$L$2:$L${hl},">="&{FROM},People!$L$2:$L${hl},"<="&{TO}))', '$#,##0')
r4 = r3 + 3 + len(LANES)

# Conversion, all time
section(S, r4 + 1, 'Conversion, all time (from the three dates on People)', 6)
label(S, r4 + 2, 1, 'Leads, ever'); calc(S, r4 + 2, 2, f'=COUNTIF(People!$J$2:$J${hl},"<>")', '0')
label(S, r4 + 3, 1, 'Of those, booked a call'); calc(S, r4 + 3, 2, f'=COUNTIFS(People!$J$2:$J${hl},"<>",People!$K$2:$K${hl},"<>")', '0'); calc(S, r4 + 3, 3, f'=IF(B{r4 + 2}=0,"",B{r4 + 3}/B{r4 + 2})', '0%')
label(S, r4 + 4, 1, 'Of the calls, became clients'); calc(S, r4 + 4, 2, f'=COUNTIFS(People!$K$2:$K${hl},"<>",People!$L$2:$L${hl},"<>")', '0'); calc(S, r4 + 4, 3, f'=IF(B{r4 + 3}=0,"",B{r4 + 4}/B{r4 + 3})', '0%')
label(S, r4 + 5, 1, 'Average days from lead to client'); calc(S, r4 + 5, 2, f'=IF(COUNTIFS(People!$J$2:$J${hl},"<>",People!$L$2:$L${hl},"<>")=0,"",SUMPRODUCT((People!$J$2:$J${hl}<>"")*(People!$L$2:$L${hl}<>"")*(People!$L$2:$L${hl}-People!$J$2:$J${hl}))/COUNTIFS(People!$J$2:$J${hl},"<>",People!$L$2:$L${hl},"<>"))', '0')
r5 = r4 + 6

# Who to contact today
section(S, r5 + 1, 'Who to contact today (overdue first)', 6)
for i, h in enumerate(['Name', 'Stage', 'Days over', 'Next touch', 'How'], 1):
    c = S.cell(row=r5 + 2, column=i, value=h); c.font = F(bold=True); c.border = BORDER
for k in range(1, 13):
    r = r5 + 2 + k
    key = f'IFERROR(LARGE(People!$Z$2:$Z${hl},{k}),"")'
    calc(S, r, 1, f'=IF({key}="","",INDEX(People!$A$2:$A${hl},MATCH({key},People!$Z$2:$Z${hl},0)))').fill = PatternFill()
    calc(S, r, 2, f'=IF(A{r}="","",INDEX(People!$I$2:$I${hl},MATCH({key},People!$Z$2:$Z${hl},0)))')
    calc(S, r, 3, f'=IF(A{r}="","",INDEX(People!$X$2:$X${hl},MATCH({key},People!$Z$2:$Z${hl},0)))', '0')
    calc(S, r, 4, f'=IF(A{r}="","",INDEX(People!$V$2:$V${hl},MATCH({key},People!$Z$2:$Z${hl},0)))', 'yyyy-mm-dd')
    calc(S, r, 5, f'=IF(A{r}="","",INDEX(People!$Y$2:$Y${hl},MATCH({key},People!$Z$2:$Z${hl},0)))')
S.column_dimensions['E'].width = 16
S.freeze_panes = 'A4'
for row in S.iter_rows(min_row=1, max_row=r5 + 15):
    for c in row:
        if c.font.name != FONT: c.font = F(bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size or 10)

# ---------------------------------------------------------------- Weekly
W = wb.create_sheet('Weekly')
title(W, 'The last twelve weeks', 'One row a week, newest at the bottom. The streak counts weeks in a row on the posting target.')
head(W, 4, ['Week starting', 'Posts', 'Reach', 'Engagement rate', 'New followers', 'Reach-outs', 'Conversations', 'Leads', 'Calls booked', 'On target?', 'Streak'], [13, 8, 10, 12, 12, 11, 13, 8, 12, 10, 8])
for i in range(12):
    r = 5 + i
    calc(W, r, 1, f'=TODAY()-WEEKDAY(TODAY(),2)+1-7*{11 - i}', 'yyyy-mm-dd')
    pi = f'Posts!$A$2:$A${last},">="&$A{r},Posts!$A$2:$A${last},"<="&$A{r}+6'
    ti = f'{T_DATES},">="&$A{r},{T_DATES},"<="&$A{r}+6'
    calc(W, r, 2, f'=COUNTIFS({pi})', '0')
    calc(W, r, 3, f'=IF(COUNTIFS(Posts!$I$2:$I${last},"<>",{pi})=0,"",SUMIFS(Posts!$I$2:$I${last},{pi}))', '#,##0')
    calc(W, r, 4, f'=IF(OR(C{r}="",C{r}=0,COUNTIFS(Posts!$T$2:$T${last},"<>",{pi})=0),"",SUMIFS(Posts!$T$2:$T${last},{pi})/C{r})', '0.0%')
    calc(W, r, 5, f'=IF(COUNTIFS(Posts!$P$2:$P${last},"<>",{pi})=0,"",SUMIFS(Posts!$P$2:$P${last},{pi}))', '#,##0')
    calc(W, r, 6, f'=COUNTIFS({ti},Touches!$E$2:$E${tl},"Me")', '0')
    calc(W, r, 7, f'=COUNTIFS({ti},Touches!$E$2:$E${tl},"Them")+COUNTIFS({ti},Touches!$E$2:$E${tl},"Me",Touches!$F$2:$F${tl},"They replied")', '0')
    calc(W, r, 8, f'=COUNTIFS(People!$J$2:$J${hl},">="&$A{r},People!$J$2:$J${hl},"<="&$A{r}+6)', '0')
    calc(W, r, 9, f'=COUNTIFS(People!$K$2:$K${hl},">="&$A{r},People!$K$2:$K${hl},"<="&$A{r}+6)', '0')
    calc(W, r, 10, f'=IF(B{r}>=Lists!$E$5,"yes","no")')
    calc(W, r, 11, f'=IF(J{r}="yes",{"K" + str(r - 1) if i else "0"}+1,0)', '0')
label(W, 18, 1, 'Streak now', True); calc(W, 18, 2, '=IF(J16="yes",K16,IF(J16="no",K15,0))', '0'); label(W, 18, 3, 'weeks in a row (this week counts once it is on target)')
ch = BarChart(); ch.title = 'Posts a week'; ch.style = 10; ch.y_axis.title = 'Posts'; ch.legend = None
ch.add_data(Reference(W, min_col=2, min_row=4, max_row=16), titles_from_data=True); ch.set_categories(Reference(W, min_col=1, min_row=5, max_row=16)); ch.height, ch.width = 7, 16
W.add_chart(ch, 'A20')
lc = LineChart(); lc.title = 'Reach a week'; lc.style = 12; lc.legend = None
lc.add_data(Reference(W, min_col=3, min_row=4, max_row=16), titles_from_data=True); lc.set_categories(Reference(W, min_col=1, min_row=5, max_row=16)); lc.height, lc.width = 7, 16
W.add_chart(lc, 'E20')
bc = BarChart(); bc.title = 'Reach-outs, conversations, leads'; bc.style = 10
bc.add_data(Reference(W, min_col=6, max_col=8, min_row=4, max_row=16), titles_from_data=True); bc.set_categories(Reference(W, min_col=1, min_row=5, max_row=16)); bc.height, bc.width = 7, 16
W.add_chart(bc, 'A35')
W.freeze_panes = 'A5'

# ---------------------------------------------------------------- Start here
G = wb.create_sheet('Start here', 0)
G.column_dimensions['A'].width = 110
lines = [
    ('Marketing Scoreboard', 'h1'),
    ('What gets measured gets managed. Three logs, one scoreboard, nothing hidden: every number is a formula you can click on.', 'p'),
    ('', 'p'),
    ('The tabs', 'h2'),
    ('Posts: one row a post. Type the date, channel, format, topic, hook and ask when it goes out. Come back three days later and type the platform\'s numbers (reach, likes, comments, shares, saves, clicks, DMs, new followers, leads). Rows waiting for results turn yellow.', 'p'),
    ('People: one row a person. Stage, how you met them, and three dates that drive everything: became a lead on, call booked on, became a client on. The grey columns on the right work out their last touch, next touch and status.', 'p'),
    ('Touches: one row each time you and a person are in contact. Pick the person, the kind, the lane, who started it and what happened.', 'p'),
    ('Scoreboard: this week against the targets with pace; the period\'s numbers against the period before; the funnel; by channel, by ask, by format, by lane; conversion; who to contact today.', 'p'),
    ('Weekly: the last twelve weeks, one row a week, and the posting streak. Three charts.', 'p'),
    ('Lists: the dropdowns, the weekly targets, and each stage\'s rhythm (days between touches) and how to reach out.', 'p'),
    ('', 'p'),
    ('The rules', 'h2'),
    ('Blank is not zero. A result you have not checked stays blank; the Scoreboard says blank, never 0%. A typed 0 is a real zero.', 'p'),
    ('Cream cells with blue text are yours to change (targets, rhythms, the period). Grey cells are formulas; leave them alone.', 'p'),
    ('Type names on Touches exactly as they are on People; the dropdown does that for you.', 'p'),
    ('The rows marked EXAMPLE in the Notes column show the shape. Delete them when you start.', 'p'),
    ('', 'p'),
    ('The words', 'h2'),
    ('Reach: how many times a post was shown. Use the platform\'s number, the same one every time.', 'p'),
    ('Engagement rate: likes + comments + shares + saves, over reach. Above 2% is fine, above 5% is good.', 'p'),
    ('Reach-out: a touch you started. Conversation: a touch they started, or one of yours marked "They replied".', 'p'),
    ('Lane: where the person came from. Warm, cold, from content, paid, referral (Hormozi\'s Core Four, and referrals).', 'p'),
    ('Pace: where a weekly number should be by today if the week ends on target.', 'p'),
    ('Rhythm: days between touches for a stage. A person\'s own number or a "contact them on" date wins over it.', 'p'),
]
for i, (text, kind) in enumerate(lines, 1):
    c = G.cell(row=i, column=1, value=text); c.alignment = Alignment(wrap_text=True, vertical='top')
    c.font = F(bold=True, size=16) if kind == 'h1' else F(bold=True, size=12) if kind == 'h2' else F()

# ---------------------------------------------------------------- Example rows
d = lambda n: TODAY - dt.timedelta(days=n)
posts = [
    (d(20), 'Instagram', 'Carousel', 'money habits', 'The one number nobody checks', 'Comment a word', '', 60, 1840, 96, 14, 9, 31, None, 3, 12, 1, d(17)),
    (d(18), 'LinkedIn', 'Text post', 'client story', 'What a client asked me this week', 'DM me', '', 25, 2210, 118, 22, 11, 19, None, 6, 15, 2, d(15)),
    (d(15), 'YouTube', 'Long video', 'how to', 'The spreadsheet that runs my life', 'Click the link', 'https://example.com/video', 240, 960, 41, 8, 3, 12, 38, None, 9, 1, d(12)),
    (d(12), 'Instagram', 'Short video', 'a mistake I made', 'I did this wrong for ten years', 'Follow', '', 25, 3120, 155, 19, 24, 27, None, 2, 31, 0, d(9)),
    (d(9), 'Newsletter', 'Email', 'money habits', 'Why your emergency fund is the wrong size', 'Book a call', '', 90, 410, None, 4, None, None, 22, None, None, 2, d(6)),
    (d(6), 'LinkedIn', 'Carousel', 'behind the scenes', 'Three minutes on the thing you keep putting off', 'No ask', '', 60, 1530, 84, 11, 6, 22, None, 1, 8, 0, d(3)),
    (d(2), 'Instagram', 'Story', 'behind the scenes', 'Sunday planning, live', 'DM me', '', 15, None, None, None, None, None, None, None, None, None, None),
    (d(0), 'TikTok', 'Short video', 'how to', 'Stop budgeting. Do this instead.', 'Follow', '', 25, None, None, None, None, None, None, None, None, None, None),
]
for i, row in enumerate(posts):
    for j, v in enumerate(row, 1):
        if v is not None: P.cell(row=2 + i, column=j, value=v)
    P.cell(row=2 + i, column=19, value='EXAMPLE')
people = [
    ('Ada Okafor', 'Northwind Studio', 'Founder', 'ada@example.com', '', 'Instagram', '@ada.demo', 'From content', 'Lead', d(6), None, None, None, None, None, 'small business', 'What a client asked me this week', 'EXAMPLE', d(14)),
    ('Ben Sato', 'Bramble & Co', 'Owner', 'ben@example.com', '', 'LinkedIn', '', 'Referral', 'Call booked', d(10), d(4), None, None, None, None, '', '', 'EXAMPLE', d(20)),
    ('Cleo Marsh', 'Halcyon Physio', '', 'cleo@example.com', '555-0100', 'Email', '', 'Warm (they know me)', 'Client', d(40), d(33), d(21), 2500, None, None, 'parent', '', 'EXAMPLE', d(50)),
    ('Dev Patel', '', '', 'dev@example.com', '', 'Instagram', '@dev.demo', 'From content', 'In conversation', None, None, None, None, None, None, '', 'The one number nobody checks', 'EXAMPLE', d(12)),
    ('Esme Haddad', 'Ferro Fitness', 'Coach', 'esme@example.com', '', 'LinkedIn', '', 'Cold (they do not)', 'Follower', None, None, None, None, None, None, '', '', 'EXAMPLE', d(9)),
    ('Femi Lindqvist', 'Oak Lane Dental', '', 'femi@example.com', '', 'Email', '', 'Referral', 'Advocate', d(120), d(110), d(95), 4800, None, None, '', '', 'EXAMPLE', d(130)),
]
for i, row in enumerate(people):
    for j, v in enumerate(row, 1):
        if v is not None: H.cell(row=2 + i, column=j, value=v)
touches = [
    (d(13), 'Ada Okafor', 'DM', 'From content', 'Them', 'They replied', 'Asked about the carousel'),
    (d(11), 'Ada Okafor', 'DM', 'From content', 'Me', 'They replied', 'Sent the worksheet'),
    (d(6), 'Ada Okafor', 'Email', 'From content', 'Me', 'No reply yet', 'Offered a call'),
    (d(19), 'Ben Sato', 'Email', 'Referral', 'Me', 'They replied', ''),
    (d(4), 'Ben Sato', 'Call', 'Referral', 'Me', 'Call booked', 'Thursday 10am'),
    (d(21), 'Cleo Marsh', 'Call', 'Warm (they know me)', 'Me', 'They bought', 'Six sessions'),
    (d(12), 'Dev Patel', 'Comment', 'From content', 'Them', 'No reply yet', 'Commented "yes" on the post'),
    (d(9), 'Esme Haddad', 'DM', 'Cold (they do not)', 'Me', 'No reply yet', ''),
    (d(2), 'Esme Haddad', 'DM', 'Cold (they do not)', 'Me', 'No reply yet', 'Second try'),
]
for i, row in enumerate(touches):
    for j, v in enumerate(row, 1):
        if v is not None: T.cell(row=2 + i, column=j, value=v)

wb.save(OUT)
print('wrote', os.path.relpath(OUT))
