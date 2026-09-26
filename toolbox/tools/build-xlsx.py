#!/usr/bin/env python3
"""
toolbox/tools/build-xlsx.py, the Toolbox as one workbook (TB-014).

Ten tabs, one a tool, every answer a live formula on the inputs above it;
a cover with the legend; a Tax Tables tab that carries the 2026 figures
from data/ so no number is hidden inside a formula. Example numbers only.

  python3 toolbox/tools/build-xlsx.py            writes toolbox/SPARKS-Toolbox.xlsx
"""
import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
OUT = os.path.join(ROOT, 'toolbox', 'SPARKS-Toolbox.xlsx')
FED = json.load(open(os.path.join(ROOT, 'data', 'federal_brackets_2026.json')))
SE = json.load(open(os.path.join(ROOT, 'data', 'se_tax_2026.json')))
RULES = json.load(open(os.path.join(ROOT, 'toolbox', 'data', 'plan_loan_rules_2026.json')))

# ---- style ------------------------------------------------------------------
FONT = 'Arial'
NAVY = '1F3A5F'
f_title = Font(name=FONT, size=18, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color='555555')
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_label = Font(name=FONT, size=10, color='222222')
f_note = Font(name=FONT, size=9, italic=True, color='666666')
f_input = Font(name=FONT, size=10, color='0000FF')
f_calc = Font(name=FONT, size=10, color='000000')
f_link = Font(name=FONT, size=10, color='008000')
f_strong = Font(name=FONT, size=11, bold=True, color='000000')
f_verdict = Font(name=FONT, size=13, bold=True, color=NAVY)
fill_head = PatternFill('solid', fgColor=NAVY)
fill_input = PatternFill('solid', fgColor='FFFFCC')
fill_answer = PatternFill('solid', fgColor='EEF3FA')
fill_band = PatternFill('solid', fgColor='F4F6F9')
thin = Side(style='thin', color='C9D1DC')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '$#,##0;($#,##0);-'
MONEY2 = '$#,##0.00;($#,##0.00);-'
PCT = '0.00%'
PCT1 = '0.0%'
MONTHS = '0.0'
INT = '0'
DATE = 'mmm d, yyyy'

wb = Workbook()

def sheet(name, title, sub, widths=(2, 44, 18, 62)):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws['B1'] = title; ws['B1'].font = f_title
    ws['B2'] = sub; ws['B2'].font = f_sub
    ws['B2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('B2:D2'); ws.row_dimensions[2].height = 30
    ws.sheet_view.showGridLines = False
    return ws

def head(ws, row, text, cols=('B', 'C', 'D')):
    for c in cols:
        ws[f'{c}{row}'].fill = fill_head; ws[f'{c}{row}'].font = f_head
    ws[f'{cols[0]}{row}'] = text
    ws.row_dimensions[row].height = 18

def inp(ws, row, label, value, fmt=None, note=None, key=None, cells=None):
    ws[f'B{row}'] = label; ws[f'B{row}'].font = f_label
    c = ws[f'C{row}']; c.value = value; c.font = f_input; c.fill = fill_input; c.border = box
    c.alignment = Alignment(horizontal='right')
    if fmt: c.number_format = fmt
    if note: ws[f'D{row}'] = note; ws[f'D{row}'].font = f_note; ws[f'D{row}'].alignment = Alignment(wrap_text=True, vertical='top', indent=1)
    if key is not None and cells is not None: cells[key] = f'$C${row}'
    return f'C{row}'

def out(ws, row, label, formula, fmt=None, note=None, strong=False, key=None, cells=None):
    ws[f'B{row}'] = label; ws[f'B{row}'].font = f_strong if strong else f_label
    c = ws[f'C{row}']; c.value = formula; c.font = f_strong if strong else f_calc
    c.alignment = Alignment(horizontal='right')
    if strong: c.fill = fill_answer; ws[f'B{row}'].fill = fill_answer
    if fmt: c.number_format = fmt
    if note: ws[f'D{row}'] = note; ws[f'D{row}'].font = f_note; ws[f'D{row}'].alignment = Alignment(wrap_text=True, vertical='top', indent=1)
    if key is not None and cells is not None: cells[key] = f'$C${row}'
    return f'C{row}'

def verdict(ws, row, formula):
    ws[f'B{row}'] = 'THE ANSWER'; ws[f'B{row}'].font = Font(name=FONT, size=9, bold=True, color='888888')
    ws.merge_cells(f'B{row+1}:D{row+1}')
    c = ws[f'B{row+1}']; c.value = formula; c.font = f_verdict; c.fill = fill_answer
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row + 1].height = 44
    return row + 2

def dropdown(ws, cell, options):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(cell)

def small(ws, row, text):
    ws.merge_cells(f'B{row}:D{row}')
    c = ws[f'B{row}']; c.value = text; c.font = f_note; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[row].height = max(15, 13 * (1 + len(text) // 120))

def table_head(ws, row, col, labels, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab); c.font = f_head; c.fill = fill_head
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        if widths: ws.column_dimensions[get_column_letter(col + i)].width = widths[i]

def cell_fmt(c, fmt, kind='calc'):
    c.font = f_input if kind == 'input' else f_calc
    if kind == 'input': c.fill = fill_input
    c.border = box; c.number_format = fmt

STATUSES = ['Single', 'Married, joint', 'Married, separate', 'Head of household']
STATUS_KEY = {'Single': 'single', 'Married, joint': 'married_joint', 'Married, separate': 'married_separate', 'Head of household': 'head_of_household'}

# =============================================================================
# Cover
# =============================================================================
cv = wb.active; cv.title = 'Start Here'
cv.sheet_view.showGridLines = False
for col, w in zip('ABCD', (2, 30, 70, 8)): cv.column_dimensions[col].width = w
cv['B1'] = 'SPARKS Toolbox'; cv['B1'].font = Font(name=FONT, size=22, bold=True, color=NAVY)
cv['B2'] = 'Ten money calculators in one workbook. Each tab is one question with a number for an answer.'; cv['B2'].font = f_sub
cv['B4'] = 'HOW IT WORKS'; cv['B4'].font = f_head; cv['B4'].fill = fill_head; cv['C4'].fill = fill_head
rows = [
    ('Type only in the yellow boxes', 'Blue numbers on yellow are yours to change. Everything in black is worked out from them and should be left alone.'),
    ('The answer is at the top of each tab', 'A sentence in plain words, then the figures behind it. Change a yellow box and the sentence changes.'),
    ('Empty means not known', 'A yellow box marked optional can be left blank. Do not type 0 unless the true figure is zero.'),
    ('Every tab opens with example numbers', 'They are nobody\'s real numbers. Type over them with yours.'),
    ('Percentages are typed as percentages', 'Type 7 in a rate box and it reads 7.00%. Money is typed in dollars.'),
    ('Estimates, not advice', 'Taxes use the 2026 federal tables on the Tax Tables tab, marked unverified. No state schedules, no credits. A tool that says a bill is coming is a reason to check, not a return.'),
]
r = 5
for a, b in rows:
    cv[f'B{r}'] = a; cv[f'B{r}'].font = Font(name=FONT, size=10, bold=True)
    cv[f'C{r}'] = b; cv[f'C{r}'].font = f_label; cv[f'C{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    cv.row_dimensions[r].height = 30; r += 1
r += 1
cv[f'B{r}'] = 'LEGEND'; cv[f'B{r}'].font = f_head; cv[f'B{r}'].fill = fill_head; cv[f'C{r}'].fill = fill_head; r += 1
cv[f'B{r}'] = 12345; cv[f'B{r}'].font = f_input; cv[f'B{r}'].fill = fill_input; cv[f'B{r}'].number_format = MONEY; cv[f'B{r}'].border = box
cv[f'C{r}'] = 'A number you type. Blue, on yellow.'; cv[f'C{r}'].font = f_label; r += 1
cv[f'B{r}'] = 12345; cv[f'B{r}'].font = f_calc; cv[f'B{r}'].number_format = MONEY
cv[f'C{r}'] = 'A number the sheet works out. Black. Leave it.'; cv[f'C{r}'].font = f_label; r += 1
cv[f'B{r}'] = 12345; cv[f'B{r}'].font = f_strong; cv[f'B{r}'].fill = fill_answer; cv[f'B{r}'].number_format = MONEY
cv[f'C{r}'] = 'The figure that answers the question. Bold, shaded.'; cv[f'C{r}'].font = f_label; r += 2
cv[f'B{r}'] = 'THE TEN TOOLS'; cv[f'B{r}'].font = f_head; cv[f'B{r}'].fill = fill_head; cv[f'C{r}'].fill = fill_head; r += 1
TOOLS = [
    ('01 Paycheck', 'Paycheck Check', 'Is the withholding on this paystub right? Refund or bill in April, and the change per check.'),
    ('02 Refinance', 'Refinance', 'Is a new loan worth its closing costs? Three paths to the last payment.'),
    ('03 Move the Debt', 'Move the Debt', 'A balance transfer or a consolidation loan against staying put.'),
    ('04 Cash or Finance', 'Pay Cash or Finance', '0% for 24 months, or a discount for cash: which leaves you richer?'),
    ('05 Lump or Payments', 'Lump Sum or Payments', 'A pension buyout or severance: the return the payments promise, the age the lump runs out.'),
    ('06 Cash Ladder', 'Cash Ladder', 'CDs or Treasury bills in rungs against one savings account.'),
    ('07 Three Paychecks', 'Three Paychecks', 'Which months hold a third check, and the budget that never counts on it.'),
    ('08 401k Loan', 'Borrow From Yourself', 'What a 401(k) loan really costs, and the bill if you leave the job.'),
    ('09 Sinking Funds', 'Sinking Funds', 'The yearly bills, spread into a monthly set-aside.'),
    ('10 Stay or Move', 'Stay or Move', 'The rent went up: stay, or move, and the counter-offer.'),
]
tool_rows_start = r
for tab, name, blurb in TOOLS:
    c = cv[f'B{r}']; c.value = name; c.hyperlink = f"#'{tab}'!B1"; c.font = Font(name=FONT, size=10, color='0563C1', underline='single')
    cv[f'C{r}'] = blurb; cv[f'C{r}'].font = f_label; cv[f'C{r}'].alignment = Alignment(wrap_text=True)
    r += 1
r += 1
cv[f'B{r}'] = 'Tax Tables'; cv[f'B{r}'].hyperlink = "#'Tax Tables'!B1"; cv[f'B{r}'].font = Font(name=FONT, size=10, color='0563C1', underline='single')
cv[f'C{r}'] = 'The 2026 federal brackets, standard deductions, payroll tax rates and 401(k) loan limits the tools read. Change a figure here and every tool follows.'; cv[f'C{r}'].font = f_label; cv[f'C{r}'].alignment = Alignment(wrap_text=True)
cv.row_dimensions[r].height = 30
r += 2
cv[f'B{r}'] = 'Built from the SPARKS Toolbox (toolbox/), ' + datetime.date.today().isoformat() + '. The same maths as the web tools; the web tools are the reference.'
cv[f'B{r}'].font = f_note; cv.merge_cells(f'B{r}:C{r}')

# =============================================================================
# Tax Tables (built early so the others can reference it)
# =============================================================================
tt = sheet('Tax Tables', 'Tax Tables, 2026', 'Every figure the tools read, in one place. Change a number here and every tool follows. Source and confidence are stated beside each block.', widths=(2, 26, 16, 12, 12, 60))
tt['B4'] = 'FEDERAL BRACKETS (ordinary income)'; head(tt, 4, 'FEDERAL BRACKETS (ordinary income)', cols=('B', 'C', 'D', 'E', 'F'))
table_head(tt, 5, 2, ['Filing status', 'Taxable income from', 'Rate', 'Rate step', 'Note'])
tt['F5'] = 'Note'
r = 6
BR_FIRST = r
for label, key in STATUS_KEY.items():
    prev_rate = 0.0; lower = 0
    for b in FED['brackets'][key]:
        tt[f'B{r}'] = label; tt[f'B{r}'].font = f_label
        tt[f'C{r}'] = lower; cell_fmt(tt[f'C{r}'], MONEY, 'input')
        tt[f'D{r}'] = b['rate']; cell_fmt(tt[f'D{r}'], PCT1, 'input')
        tt[f'E{r}'] = f'=D{r}-' + (f'D{r-1}' if prev_rate else '0'); cell_fmt(tt[f'E{r}'], PCT1)
        prev_rate = b['rate']; lower = b['upToTaxableIncome'] if b['upToTaxableIncome'] is not None else lower
        r += 1
BR_LAST = r - 1
tt[f'F{BR_FIRST}'] = 'Source: ' + FED['source']; tt[f'F{BR_FIRST}'].font = f_note; tt[f'F{BR_FIRST}'].alignment = Alignment(wrap_text=True, vertical='top')
tt.merge_cells(f'F{BR_FIRST}:F{BR_FIRST+5}')
tt[f'F{BR_FIRST+6}'] = 'Confidence: ' + FED['confidence'] + '. ' + FED['confidenceNote']; tt[f'F{BR_FIRST+6}'].font = f_note; tt[f'F{BR_FIRST+6}'].alignment = Alignment(wrap_text=True, vertical='top')
tt.merge_cells(f'F{BR_FIRST+6}:F{BR_FIRST+12}')
tt[f'F{BR_FIRST+13}'] = 'Tax on a taxable income is the sum over rows of (income above the row\'s "from", if any) times the rate step. That is the bracket ladder in one formula, and it is what 01 Paycheck uses.'; tt[f'F{BR_FIRST+13}'].font = f_note; tt[f'F{BR_FIRST+13}'].alignment = Alignment(wrap_text=True, vertical='top')
tt.merge_cells(f'F{BR_FIRST+13}:F{BR_FIRST+18}')
BR_STATUS = f"'Tax Tables'!$B${BR_FIRST}:$B${BR_LAST}"
BR_LOWER = f"'Tax Tables'!$C${BR_FIRST}:$C${BR_LAST}"
BR_STEP = f"'Tax Tables'!$E${BR_FIRST}:$E${BR_LAST}"

r += 1
head(tt, r, 'STANDARD DEDUCTION AND ADDITIONAL MEDICARE THRESHOLD', cols=('B', 'C', 'D', 'E', 'F')); r += 1
table_head(tt, r, 2, ['Filing status', 'Standard deduction', 'Add. Medicare from', '', '']); r += 1
SD_FIRST = r
for label, key in STATUS_KEY.items():
    tt[f'B{r}'] = label; tt[f'B{r}'].font = f_label
    tt[f'C{r}'] = FED['standardDeduction'][key]; cell_fmt(tt[f'C{r}'], MONEY, 'input')
    tt[f'D{r}'] = SE['additionalMedicare']['thresholds'][key]; cell_fmt(tt[f'D{r}'], MONEY, 'input')
    r += 1
SD_LAST = r - 1
SD_RANGE = f"'Tax Tables'!$B${SD_FIRST}:$B${SD_LAST}"
SD_VALS = f"'Tax Tables'!$C${SD_FIRST}:$C${SD_LAST}"
AM_VALS = f"'Tax Tables'!$D${SD_FIRST}:$D${SD_LAST}"
tt[f'F{SD_FIRST}'] = 'Standard deduction: same source as the brackets. Additional Medicare threshold: ' + SE['additionalMedicare']['note']; tt[f'F{SD_FIRST}'].font = f_note; tt[f'F{SD_FIRST}'].alignment = Alignment(wrap_text=True, vertical='top'); tt.merge_cells(f'F{SD_FIRST}:F{SD_LAST}')

r += 1
head(tt, r, 'PAYROLL TAX (what an employee pays)', cols=('B', 'C', 'D', 'E', 'F')); r += 1
FICA = {}
for label, key, val, fmt, note in [
    ('Social Security rate, employee', 'ssRate', SE['socialSecurityRate'] / 2, PCT1, 'Half of the 12.4% combined rate; the employer pays the other half.'),
    ('Social Security wage base', 'ssBase', SE['socialSecurityWageBase'], MONEY, SE['precision']),
    ('Medicare rate, employee', 'medRate', SE['medicareRate'] / 2, PCT1, 'Half of 2.9%. No cap.'),
    ('Additional Medicare rate', 'addRate', SE['additionalMedicare']['rate'], PCT1, 'On wages above the threshold for the filing status.'),
]:
    tt[f'B{r}'] = label; tt[f'B{r}'].font = f_label
    tt[f'C{r}'] = val; cell_fmt(tt[f'C{r}'], fmt, 'input'); FICA[key] = f"'Tax Tables'!$C${r}"
    tt[f'F{r}'] = note; tt[f'F{r}'].font = f_note; tt[f'F{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
tt[f'F{r}'] = 'Source: ' + SE['source']; tt[f'F{r}'].font = f_note; tt[f'F{r}'].alignment = Alignment(wrap_text=True, vertical='top'); r += 1

r += 1
head(tt, r, '401(k) LOAN RULES', cols=('B', 'C', 'D', 'E', 'F')); r += 1
PL = {}
for label, key, val, fmt in [
    ('Most a plan can lend', 'maxLoan', RULES['maxLoanCents'] / 100, MONEY),
    ('Or this share of the vested balance', 'maxShare', RULES['maxShareOfVested'], PCT1),
    ('Small-balance floor', 'floor', RULES['smallBalanceFloorCents'] / 100, MONEY),
    ('Longest term, months', 'maxMonths', RULES['maxMonths'], INT),
    ('Early withdrawal penalty', 'penalty', RULES['penaltyRate'], PCT1),
    ('Penalty stops at age', 'penaltyAge', RULES['penaltyAge'], '0.0'),
]:
    tt[f'B{r}'] = label; tt[f'B{r}'].font = f_label
    tt[f'C{r}'] = val; cell_fmt(tt[f'C{r}'], fmt, 'input'); PL[key] = f"'Tax Tables'!$C${r}"
    r += 1
tt[f'F{r-6}'] = 'Source: ' + RULES['source'] + ' Confidence: ' + RULES['confidence'] + '. ' + RULES['confidenceNote']; tt[f'F{r-6}'].font = f_note; tt[f'F{r-6}'].alignment = Alignment(wrap_text=True, vertical='top'); tt.merge_cells(f'F{r-6}:F{r-1}')

r += 1
head(tt, r, 'PAY FREQUENCY', cols=('B', 'C', 'D', 'E', 'F')); r += 1
table_head(tt, r, 2, ['Paid', 'Checks a year', 'In a normal month', '', '']); r += 1
PF_FIRST = r
for label, n, base in [('Every week', 52, 4), ('Every two weeks', 26, 2), ('Twice a month', 24, 2), ('Once a month', 12, 1)]:
    tt[f'B{r}'] = label; tt[f'B{r}'].font = f_label
    tt[f'C{r}'] = n; cell_fmt(tt[f'C{r}'], INT, 'input')
    tt[f'D{r}'] = base; cell_fmt(tt[f'D{r}'], INT, 'input')
    r += 1
PF_LAST = r - 1
PF_LABELS = f"'Tax Tables'!$B${PF_FIRST}:$B${PF_LAST}"
PF_N = f"'Tax Tables'!$C${PF_FIRST}:$C${PF_LAST}"
PF_BASE = f"'Tax Tables'!$D${PF_FIRST}:$D${PF_LAST}"
FREQS = ['Every two weeks', 'Every week', 'Twice a month', 'Once a month']

# =============================================================================
# 01 Paycheck
# =============================================================================
ws = sheet('01 Paycheck', 'Paycheck Check', 'One paystub, read as a year. Is the withholding on track, or is April a refund or a bill? Type the lines off one check; every check is assumed to look like it.')
c = {}
head(ws, 4, 'ONE PAYSTUB'); r = 5
inp(ws, r, 'Gross pay, this check', 3000, MONEY2, 'Before anything comes out.', 'gross', c); r += 1
inp(ws, r, 'Paid', 'Every two weeks', None, 'Pick from the list.', 'freq', c); dropdown(ws, f'C{r}', FREQS); r += 1
inp(ws, r, 'Federal income tax withheld, this check', 280, MONEY2, 'The federal line on the stub.', 'fed', c); r += 1
inp(ws, r, '401(k) or other pre-tax, this check', 180, MONEY2, '0 if none. Comes off before federal tax, not before payroll tax.', 'pretax', c); r += 1
inp(ws, r, 'Health, HSA, FSA premiums, this check', 120, MONEY2, '0 if none. Comes off before federal and payroll tax.', 's125', c); r += 1
inp(ws, r, 'Social Security + Medicare withheld (optional)', None, MONEY2, 'Type it to check it. Leave blank to skip.', 'fica', c); r += 1
head(ws, r, 'YOU'); r += 1
inp(ws, r, 'Filing status', 'Single', None, 'Pick from the list.', 'status', c); dropdown(ws, f'C{r}', STATUSES); r += 1
inp(ws, r, 'State income tax rate, flat (optional)', 0.05, PCT, 'A single rate is a rough stand-in for a state schedule. 0 in a state with no income tax.', 'stRate', c); r += 1
inp(ws, r, 'State tax withheld, this check (optional)', 100, MONEY2, None, 'stW', c); r += 1
inp(ws, r, 'Checks already received this year', 18, INT, '0 in January.', 'soFar', c); r += 1
inp(ws, r, 'Other income this year with no withholding', 0, MONEY, 'Interest, a side gig.', 'other', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE YEAR, READ FROM THIS CHECK'); r += 1
out(ws, r, 'Checks a year', f"=INDEX({PF_N},MATCH({c['freq']},{PF_LABELS},0))", INT, None, False, 'n', c); r += 1
out(ws, r, 'Checks left this year', f"=MAX(0,{c['n']}-{c['soFar']})", INT, None, False, 'left', c); r += 1
out(ws, r, 'Gross for the year', f"={c['gross']}*{c['n']}", MONEY, None, False, 'annual', c); r += 1
out(ws, r, 'Wages for federal tax', f"=({c['gross']}-{c['pretax']}-{c['s125']})*{c['n']}+{c['other']}", MONEY, 'After both kinds of deduction, plus other income.', False, 'fedW', c); r += 1
out(ws, r, 'Wages for payroll tax', f"=({c['gross']}-{c['s125']})*{c['n']}", MONEY, 'Only the premiums come off.', False, 'ficaW', c); r += 1
out(ws, r, 'Standard deduction', f"=INDEX({SD_VALS},MATCH({c['status']},{SD_RANGE},0))", MONEY, 'From Tax Tables.', False, 'sd', c); r += 1
out(ws, r, 'Federal taxable income', f"=MAX(0,{c['fedW']}-{c['sd']})", MONEY, None, False, 'taxable', c); r += 1
out(ws, r, 'Federal tax for the year', f"=SUMPRODUCT(({BR_STATUS}={c['status']})*({c['taxable']}>{BR_LOWER})*({c['taxable']}-{BR_LOWER})*{BR_STEP})", MONEY, 'The bracket ladder on Tax Tables. No credits.', False, 'tax', c); r += 1
out(ws, r, 'Rate on the next dollar', f"=SUMPRODUCT(({BR_STATUS}={c['status']})*({c['taxable']}>{BR_LOWER})*{BR_STEP})", PCT1, None, False, 'marg', c); r += 1
out(ws, r, 'Withheld for the year, at this check\'s rate', f"={c['fed']}*{c['n']}", MONEY, None, False, 'withheld', c); r += 1
out(ws, r, 'April: tax less withholding', f"={c['tax']}-{c['withheld']}", MONEY, 'Positive: you will owe. Negative: a refund.', True, 'gap', c); r += 1
out(ws, r, 'Change per remaining check to land at zero', f"=IF({c['left']}>0,{c['gap']}/{c['left']},\"no checks left\")", MONEY2, 'Add this to line 4(c) of the W-4 (or reduce withholding by it if negative).', True, 'adj', c); r += 1
r += 1
head(ws, r, 'PAYROLL TAX, CHECKED'); r += 1
out(ws, r, 'Social Security + Medicare for the year', f"=MIN({c['ficaW']},{FICA['ssBase']})*{FICA['ssRate']}+{c['ficaW']}*{FICA['medRate']}+MAX(0,{c['ficaW']}-INDEX({AM_VALS},MATCH({c['status']},{SD_RANGE},0)))*{FICA['addRate']}", MONEY, None, False, 'ficaY', c); r += 1
out(ws, r, 'Per check, expected', f"={c['ficaY']}/{c['n']}", MONEY2, None, False, 'ficaPC', c); r += 1
out(ws, r, 'Yours against expected', f"=IF({c['fica']}=\"\",\"type the line to check it\",IF(ABS({c['fica']}-{c['ficaPC']})<5,\"matches\",IF({c['fica']}>{c['ficaPC']},\"yours is \"&TEXT({c['fica']}-{c['ficaPC']},\"$#,##0.00\")&\" more\",\"yours is \"&TEXT({c['ficaPC']}-{c['fica']},\"$#,##0.00\")&\" less; ask payroll\")))", None, 'Past the Social Security wage base late in the year, less is normal.'); r += 1
out(ws, r, 'Take-home per check', f"={c['gross']}-{c['pretax']}-{c['s125']}-{c['fed']}-{c['stW']}-IF({c['fica']}=\"\",{c['ficaPC']},{c['fica']})", MONEY2); r += 1
r += 1
head(ws, r, 'STATE, ROUGHLY'); r += 1
out(ws, r, 'State tax for the year at the flat rate', f"={c['taxable']}*{c['stRate']}", MONEY, 'Federal taxable income times the flat rate. A stand-in, not a schedule.', False, 'stTax', c); r += 1
out(ws, r, 'State withheld for the year', f"={c['stW']}*{c['n']}", MONEY, None, False, 'stWY', c); r += 1
out(ws, r, 'State, in April', f"={c['stTax']}-{c['stWY']}", MONEY, 'Positive: owe. Negative: refund.', True); r += 1
r += 1
small(ws, r, 'Every check is assumed to look like this one. The year\'s federal tax is the 2026 bracket ladder on the wages after the pre-tax deductions and the standard deduction; no credits, no itemising. Payroll tax is Social Security up to the wage base plus Medicare. An estimate, not a return; the W-4 is where the change is made. Under-withholding by more than 10% of the tax can carry a penalty: the safe harbours are 90% of this year\'s tax or 100% of last year\'s.')
verdict(ws, row_v, f"=IF(ABS({c['gap']})<200,\"On track. April comes within \"&TEXT(ABS({c['gap']}),\"$#,##0\")&\" either way. Leave the W-4 alone.\",IF({c['gap']}<0,\"A refund of about \"&TEXT(-{c['gap']},\"$#,##0\")&\" in April. You are lending the government \"&TEXT(-{c['adj']},\"$#,##0\")&\" a check, interest free.\",\"A bill of about \"&TEXT({c['gap']},\"$#,##0\")&\" in April. Add about \"&TEXT({c['adj']},\"$#,##0\")&\" of extra withholding to each of the \"&{c['left']}&\" checks left.\"))")

# =============================================================================
# 02 Refinance
# =============================================================================
ws = sheet('02 Refinance', 'Refinance', 'Is the new loan worth its closing costs? Any loan: a mortgage, a car, a student loan. Three paths costed to the last payment, and the month the saving has paid the costs back.')
c = {}
head(ws, 4, 'THE LOAN YOU HAVE'); r = 5
inp(ws, r, 'Balance left', 300000, MONEY, None, 'bal', c); r += 1
inp(ws, r, 'Rate now', 0.07, PCT, None, 'r0', c); r += 1
inp(ws, r, 'Months left', 300, INT, 'Or leave it and type the payment below.', 'm0', c); r += 1
inp(ws, r, 'Payment now (optional)', None, MONEY2, 'If typed, it is used instead of the months left.', 'p0in', c); r += 1
head(ws, r, 'THE LOAN ON OFFER'); r += 1
inp(ws, r, 'New rate', 0.055, PCT, None, 'r1', c); r += 1
inp(ws, r, 'New term, months', 360, INT, None, 'm1', c); r += 1
inp(ws, r, 'Closing costs', 6000, MONEY, None, 'close', c); r += 1
inp(ws, r, 'Roll the closing costs into the new balance?', 'No', None, 'Yes or No.', 'roll', c); dropdown(ws, f'C{r}', ['Yes', 'No']); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'KEEP THE LOAN YOU HAVE'); r += 1
out(ws, r, 'Payment now', f"=IF({c['p0in']}>0,{c['p0in']},PMT({c['r0']}/12,{c['m0']},-{c['bal']}))", MONEY2, 'The level payment for the rate and months, unless you typed one.', False, 'p0', c); r += 1
out(ws, r, 'Months to the last payment', f"=IF({c['p0in']}>0,IF({c['p0in']}<={c['bal']}*{c['r0']}/12,\"never: the payment does not cover the interest\",NPER({c['r0']}/12,-{c['p0in']},{c['bal']})),{c['m0']})", MONTHS, None, False, 'n0', c); r += 1
out(ws, r, 'Paid in all', f"=IFERROR({c['p0']}*{c['n0']},\"never\")", MONEY, None, False, 't0', c); r += 1
out(ws, r, 'Of which interest', f"=IFERROR({c['t0']}-{c['bal']},\"never\")", MONEY, None, False, 'i0', c); r += 1
r += 1
head(ws, r, 'TAKE THE NEW LOAN'); r += 1
out(ws, r, 'New balance', f"={c['bal']}+IF({c['roll']}=\"Yes\",{c['close']},0)", MONEY, None, False, 'bal1', c); r += 1
out(ws, r, 'New payment', f"=PMT({c['r1']}/12,{c['m1']},-{c['bal1']})", MONEY2, None, False, 'p1', c); r += 1
out(ws, r, 'Paid in all, closing costs included', f"={c['p1']}*{c['m1']}+IF({c['roll']}=\"Yes\",0,{c['close']})", MONEY, None, False, 't1', c); r += 1
out(ws, r, 'Of which interest and costs', f"={c['t1']}-{c['bal']}", MONEY, None, False, 'i1', c); r += 1
out(ws, r, 'Monthly saving', f"={c['p0']}-{c['p1']}", MONEY2, 'Negative: the new payment is higher.', False, 'save', c); r += 1
out(ws, r, 'Break-even', f"=IF({c['save']}<=0,\"never: the payment does not fall\",IF({c['close']}=0,0,{c['close']}/{c['save']}))", MONTHS, 'Months for the saving to pay back the closing costs.', True, 'be', c); r += 1
out(ws, r, 'Lifetime difference against keeping the loan', f"=IFERROR({c['t1']}-{c['t0']},\"n/a\")", MONEY, 'Negative: the new loan costs less over its life.', True, 'd1', c); r += 1
r += 1
head(ws, r, 'NEW LOAN, KEEP PAYING THE OLD AMOUNT'); r += 1
out(ws, r, 'Months to the last payment', f"=IF({c['p0']}>{c['p1']},NPER({c['r1']}/12,-{c['p0']},{c['bal1']}),\"n/a: the old payment is not higher\")", MONTHS, 'The lower rate with the same payment clears it sooner.', False, 'n2', c); r += 1
out(ws, r, 'Paid in all, closing costs included', f"=IFERROR({c['p0']}*{c['n2']}+IF({c['roll']}=\"Yes\",0,{c['close']}),\"n/a\")", MONEY, None, False, 't2', c); r += 1
out(ws, r, 'Lifetime difference against keeping the loan', f"=IFERROR({c['t2']}-{c['t0']},\"n/a\")", MONEY, None, True, 'd2', c); r += 1
r += 1
out(ws, r, 'Cheapest path', f"=IFERROR(IF(MIN({c['t0']},{c['t1']},IF(ISNUMBER({c['t2']}),{c['t2']},{c['t0']}))={c['t0']},\"Keep the loan you have\",IF(ISNUMBER({c['t2']}),IF({c['t2']}<={c['t1']},\"New loan, keep paying the old amount\",\"Take the new loan\"),\"Take the new loan\")),\"Take the new loan\")", None, None, True, 'best', c); r += 2
small(ws, r, 'The payment on each loan is the level payment for its rate and term. Closing costs count in the new loan\'s cost whether paid up front or rolled in. A lower payment on a longer term is the trap this tab exists to show: the payment falls and the lifetime cost rises. Taxes, escrow, insurance and points are not modelled.')
r += 2
# balance table + chart
head(ws, r, 'BALANCE LEFT, YEAR BY YEAR'); TB_ROW = r + 1; r += 1
table_head(ws, r, 2, ['Years from now', 'Keep the loan', 'Take the new loan']); r += 1
first = r
for k in range(0, 31):
    ws[f'B{r}'] = k; ws[f'B{r}'].font = f_calc; ws[f'B{r}'].alignment = Alignment(horizontal='center')
    ws[f'C{r}'] = f"=IFERROR(IF({k}*12>={c['n0']},0,MAX(0,-FV({c['r0']}/12,{k}*12,-{c['p0']},{c['bal']}))),0)"; cell_fmt(ws[f'C{r}'], MONEY)
    ws[f'D{r}'] = f"=IF({k}*12>={c['m1']},0,MAX(0,-FV({c['r1']}/12,{k}*12,-{c['p1']},{c['bal1']})))"; cell_fmt(ws[f'D{r}'], MONEY)
    ws[f'D{r}'].alignment = Alignment(horizontal='left')
    r += 1
ch = LineChart(); ch.title = 'What is still owed'; ch.height = 7.5; ch.width = 16; ch.y_axis.numFmt = '$#,##0'; ch.x_axis.title = 'years from now'
ch.add_data(Reference(ws, min_col=3, min_row=first - 1, max_col=4, max_row=r - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=first, max_row=r - 1))
ws.add_chart(ch, f'F{TB_ROW}')
verdict(ws, row_v, f"=IF({c['best']}=\"Keep the loan you have\",\"Keep the loan you have. Refinancing costs \"&TEXT({c['d1']},\"$#,##0\")&\" more over its life.\",IF({c['best']}=\"Take the new loan\",\"Take the new loan: \"&TEXT(-{c['d1']},\"$#,##0\")&\" less over its life. The payment falls by \"&TEXT({c['save']},\"$#,##0\")&\" a month and the closing costs are paid back in \"&IF(ISNUMBER({c['be']}),TEXT({c['be']},\"0\")&\" months.\",\"no time.\"),\"Take the new loan but keep paying the old amount: \"&TEXT(-{c['d2']},\"$#,##0\")&\" less over its life, gone in \"&TEXT({c['n2']},\"0\")&\" months. Take the lower rate, ignore the lower payment.\"))")

# =============================================================================
# 03 Move the Debt
# =============================================================================
ws = sheet('03 Move the Debt', 'Move the Debt', 'A balance transfer or a consolidation loan against leaving it where it is. Three paths, each run to the last payment. The cost of a path is everything paid minus the balance itself: interest plus fees.')
c = {}
head(ws, 4, 'THE DEBT'); r = 5
inp(ws, r, 'Balance', 8000, MONEY, None, 'bal', c); r += 1
inp(ws, r, 'APR now', 0.24, PCT, None, 'apr', c); r += 1
inp(ws, r, 'What you can pay a month', 400, MONEY2, None, 'pay', c); r += 1
head(ws, r, 'A BALANCE TRANSFER CARD (leave blank if none)'); r += 1
inp(ws, r, 'Transfer fee', 0.03, PCT, 'Goes onto the balance.', 'fee', c); r += 1
inp(ws, r, 'Promo, months at 0%', 18, INT, None, 'promo', c); r += 1
inp(ws, r, 'APR after the promo', 0.27, PCT, None, 'after', c); r += 1
head(ws, r, 'A CONSOLIDATION LOAN (leave blank if none)'); r += 1
inp(ws, r, 'Loan APR', 0.11, PCT, None, 'lapr', c); r += 1
inp(ws, r, 'Term, months', 24, INT, None, 'lm', c); r += 1
inp(ws, r, 'Origination fee', 0.02, PCT, 'Added to the loan.', 'lfee', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'STAY WHERE IT IS'); r += 1
out(ws, r, 'Months to clear', f"=IF({c['pay']}<={c['bal']}*{c['apr']}/12,\"never: the payment does not cover the interest\",NPER({c['apr']}/12,-{c['pay']},{c['bal']}))", MONTHS, None, False, 'n0', c); r += 1
out(ws, r, 'Cost: interest', f"=IFERROR({c['pay']}*{c['n0']}-{c['bal']},\"never\")", MONEY, None, True, 'c0', c); r += 1
r += 1
head(ws, r, 'BALANCE TRANSFER'); r += 1
out(ws, r, 'Balance after the fee', f"=IF({c['promo']}>0,{c['bal']}*(1+{c['fee']}),\"no offer\")", MONEY, None, False, 'tb', c); r += 1
out(ws, r, 'Payment that clears it inside the promo', f"=IFERROR({c['tb']}/{c['promo']},\"no offer\")", MONEY2, 'The number that matters.', True, 'tclear', c); r += 1
out(ws, r, 'Left when the promo ends, at your payment', f"=IFERROR(MAX(0,{c['tb']}-{c['pay']}*{c['promo']}),\"no offer\")", MONEY, None, False, 'tleft', c); r += 1
out(ws, r, 'Months to clear', f"=IFERROR(IF({c['tleft']}=0,{c['tb']}/{c['pay']},IF({c['pay']}<={c['tleft']}*{c['after']}/12,\"never\",{c['promo']}+NPER({c['after']}/12,-{c['pay']},{c['tleft']}))),\"no offer\")", MONTHS, None, False, 'n1', c); r += 1
out(ws, r, 'Cost: interest and fee', f"=IFERROR({c['pay']}*{c['n1']}-{c['bal']},\"n/a\")", MONEY, None, True, 'c1', c); r += 1
r += 1
head(ws, r, 'CONSOLIDATION LOAN'); r += 1
out(ws, r, 'Loan amount with the fee', f"=IF({c['lm']}>0,{c['bal']}*(1+{c['lfee']}),\"no offer\")", MONEY, None, False, 'lb', c); r += 1
out(ws, r, 'Its payment', f"=IFERROR(PMT({c['lapr']}/12,{c['lm']},-{c['lb']}),\"no offer\")", MONEY2, None, False, 'lp', c); r += 1
out(ws, r, 'Against what you can pay', f"=IFERROR(IF({c['lp']}>{c['pay']},TEXT({c['lp']}-{c['pay']},\"$#,##0\")&\" a month more than you said\",\"fits\"),\"no offer\")", None); r += 1
out(ws, r, 'Cost: interest and fee', f"=IFERROR({c['lp']}*{c['lm']}-{c['bal']},\"n/a\")", MONEY, None, True, 'c2', c); r += 1
r += 1
out(ws, r, 'Cheapest path', f"=IF(MIN(IF(ISNUMBER({c['c0']}),{c['c0']},9E+99),IF(ISNUMBER({c['c1']}),{c['c1']},9E+99),IF(ISNUMBER({c['c2']}),{c['c2']},9E+99))=9E+99,\"none clears it\",IF(MIN(IF(ISNUMBER({c['c0']}),{c['c0']},9E+99),IF(ISNUMBER({c['c1']}),{c['c1']},9E+99),IF(ISNUMBER({c['c2']}),{c['c2']},9E+99))=IF(ISNUMBER({c['c0']}),{c['c0']},9E+99),\"Stay where it is\",IF(MIN(IF(ISNUMBER({c['c1']}),{c['c1']},9E+99),IF(ISNUMBER({c['c2']}),{c['c2']},9E+99))=IF(ISNUMBER({c['c1']}),{c['c1']},9E+99),\"Balance transfer\",\"Consolidation loan\")))", None, None, True, 'best', c); r += 1
out(ws, r, 'Saves against staying', f"=IFERROR({c['c0']}-MIN(IF(ISNUMBER({c['c1']}),{c['c1']},9E+99),IF(ISNUMBER({c['c2']}),{c['c2']},9E+99)),\"n/a\")", MONEY, None, True, 'saves', c); r += 2
small(ws, r, 'The transfer fee goes onto the balance; the promo is taken at 0% and the after-rate applies to whatever is left. The trap is not in the arithmetic: a card moved to 0% is an empty card, and if it fills again you have two balances. Close it, freeze it, or cut it up. Minimum payments, late fees and credit score effects are not modelled.')
verdict(ws, row_v, f"=IF({c['best']}=\"none clears it\",\"None of these clears the balance at \"&TEXT({c['pay']},\"$#,##0\")&\" a month.\",IF({c['best']}=\"Stay where it is\",\"Stay where it is. The offers cost more than they save. It clears in \"&TEXT({c['n0']},\"0\")&\" months at \"&TEXT({c['c0']},\"$#,##0\")&\" of interest.\",{c['best']}&\": \"&TEXT({c['saves']},\"$#,##0\")&\" less than staying put.\"&IF({c['best']}=\"Balance transfer\",IF({c['tleft']}=0,\" Your payment clears it inside the promo.\",\" But \"&TEXT({c['tleft']},\"$#,##0\")&\" is left when the promo ends; \"&TEXT({c['tclear']},\"$#,##0\")&\" a month would clear it in time.\"),\"\")))")

# =============================================================================
# 04 Cash or Finance
# =============================================================================
ws = sheet('04 Cash or Finance', 'Pay Cash or Finance', 'The shop offers 0% for two years, or a little off for cash. Which leaves you with more after the last payment? You are assumed to hold the price in cash.')
c = {}
head(ws, 4, 'THE THING, AND THE TWO WAYS TO PAY'); r = 5
inp(ws, r, 'Price', 2400, MONEY, None, 'price', c); r += 1
inp(ws, r, 'Discount for paying cash', 0.05, PCT, '0 if none.', 'disc', c); r += 1
inp(ws, r, 'Financing APR', 0, PCT, '0 for a 0% offer.', 'apr', c); r += 1
inp(ws, r, 'Months of payments', 24, INT, None, 'm', c); r += 1
inp(ws, r, 'Any fee to finance', 0, MONEY, None, 'fee', c); r += 1
inp(ws, r, 'What your cash earns where it sits', 0.04, PCT, 'The savings account rate, usually.', 'park', c); r += 1
inp(ws, r, 'Is it a "no interest if paid in full by" plan?', 'Yes', None, 'Deferred interest. Yes or No.', 'def', c); dropdown(ws, f'C{r}', ['Yes', 'No']); r += 1
inp(ws, r, 'The rate on the fine print', 0.2999, PCT, 'Only used if the answer above is Yes.', 'dapr', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE TWO PATHS'); r += 1
out(ws, r, 'Cash price after the discount', f"={c['price']}*(1-{c['disc']})", MONEY, None, False, 'cp', c); r += 1
out(ws, r, 'Monthly payment', f"=IF({c['apr']}=0,{c['price']}/{c['m']},PMT({c['apr']}/12,{c['m']},-{c['price']}))", MONEY2, None, False, 'pmt', c); r += 1
out(ws, r, 'Interest and fees on the financing', f"={c['pmt']}*{c['m']}-{c['price']}+{c['fee']}", MONEY, None, False, 'fi', c); r += 1
out(ws, r, 'Pay cash: left in the account after the last month', f"=({c['price']}-{c['cp']})*(1+{c['park']}/12)^{c['m']}", MONEY, 'The discount, kept and growing.', False, 'cashEnd', c); r += 1
out(ws, r, 'Finance: left in the account after the last payment', f"=FV({c['park']}/12,{c['m']},{c['pmt']},-({c['price']}-{c['fee']}))", MONEY, 'The cash keeps earning while a payment leaves each month (a positive pmt in FV is money coming out).', False, 'finEnd', c); r += 1
out(ws, r, 'Earned by your cash while the payments run', f"={c['finEnd']}-({c['price']}-{c['fee']}-{c['pmt']}*{c['m']})", MONEY, None, False, 'earned', c); r += 1
out(ws, r, 'Financing leaves you ahead by', f"={c['finEnd']}-{c['cashEnd']}", MONEY, 'Negative: paying cash leaves you ahead.', True, 'edge', c); r += 1
out(ws, r, 'Back interest if a deferred payment slips', f"=IF({c['def']}=\"Yes\",PMT({c['dapr']}/12,{c['m']},-{c['price']})*{c['m']}-{c['price']},\"n/a\")", MONEY, 'Charged from day one at the fine-print rate if a single payment is late or any balance remains at the end.', True); r += 1
r += 1
small(ws, r, 'Paying cash: the discount stays in the account and earns the parked rate. Financing: the account keeps earning while a payment leaves it each month and any fee leaves at once. Whichever leaves more in the account after the last payment won. With a deferred-interest plan, set the payments to autopay and finish a month early.')
verdict(ws, row_v, f"=IF(ABS({c['edge']})<1,\"A wash: the two end within a dollar of each other.\",IF({c['edge']}>0,\"Take the payments. You end \"&TEXT({c['edge']},\"$#,##0\")&\" ahead: your cash earns \"&TEXT({c['earned']},\"$#,##0\")&\" while the payments run.\",\"Pay cash. You end \"&TEXT(-{c['edge']},\"$#,##0\")&\" ahead.\"))&IF({c['def']}=\"Yes\",\" Deferred interest: miss one payment and about \"&TEXT(PMT({c['dapr']}/12,{c['m']},-{c['price']})*{c['m']}-{c['price']},\"$#,##0\")&\" of back interest lands.\",\"\")")

# =============================================================================
# 05 Lump or Payments
# =============================================================================
ws = sheet('05 Lump or Payments', 'Lump Sum or Payments', 'One amount now, or so much a month for life. What return are the payments quietly promising, and how long would the lump sum last if you took it? Nothing here knows how long you will live; the age you plan to is the whole of that assumption.')
c = {}
head(ws, 4, 'THE TWO OFFERS'); r = 5
inp(ws, r, 'The lump sum, today', 200000, MONEY, None, 'lump', c); r += 1
inp(ws, r, 'Or, a month', 1200, MONEY2, None, 'mo', c); r += 1
inp(ws, r, 'Payments rise each year by', 0, PCT, '0 if fixed.', 'cola', c); r += 1
inp(ws, r, 'Payments start at age (optional)', None, INT, 'Blank: they start now.', 'startIn', c); r += 1
head(ws, r, 'YOU'); r += 1
inp(ws, r, 'Your age', 62, INT, None, 'age', c); r += 1
inp(ws, r, 'Plan to age', 90, INT, None, 'to', c); r += 1
inp(ws, r, 'Return you would earn on the lump', 0.05, PCT, None, 'ret', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE TWO READINGS'); r += 1
out(ws, r, 'Payments start at', f"=IF({c['startIn']}>0,{c['startIn']},{c['age']})", INT, None, False, 'start', c); r += 1
out(ws, r, 'Payments received by the planning age', f"=SUM(D50:D111)", MONEY, 'From the year table below.', False, 'total', c); r += 1
out(ws, r, 'Years of payments to equal the lump', f"={c['lump']}/({c['mo']}*12)", '0.0', 'Ignoring returns.', False); r += 1
out(ws, r, 'Implied return of the payments', f"=IFERROR(IRR(E50:E111),\"below zero\")", PCT, 'The yearly return the lump would have to earn to fund exactly these payments to the planning age.', True, 'irr', c); r += 1
out(ws, r, 'Lump sum runs out around age', f"=IF(ISNUMBER(MATCH(0,F51:F111,0)),{c['age']}+MATCH(0,F51:F111,0),\"not before \"&{c['to']})", INT, 'Invested at your rate, drawn at the monthly amount from the start age.', True, 'runout', c); r += 1
out(ws, r, 'Left at the planning age if you take the lump', f"=INDEX(F50:F111,{c['to']}-{c['age']}+1)", MONEY, None, True, 'left', c); r += 1
r += 1
small(ws, r, 'Payments for life are insurance against living long; a lump sum is a bet that you will not, or that you can earn more than the implied return. Someone in good health with long-lived parents should lean to the payments; someone who needs the money flexible, or has heirs, to the lump. Also ask who stands behind the payments. Tax is not modelled: a lump rolled to an IRA is not taxed on arrival; a payment is taxed as it comes.')
# year table at row 49
head(ws, 48, 'YEAR BY YEAR', cols=('B', 'C', 'D', 'E', 'F'))
table_head(ws, 49, 2, ['Year', 'Age', 'Payments that year', 'Cash flow for the return', 'Lump sum left, at your rate'])
ws.column_dimensions['E'].width = 22; ws.column_dimensions['F'].width = 24
for k in range(0, 62):
    rr = 50 + k
    ws[f'B{rr}'] = k; ws[f'B{rr}'].font = f_calc; ws[f'B{rr}'].alignment = Alignment(horizontal='center')
    ws[f'C{rr}'] = f"={c['age']}+{k}"; cell_fmt(ws[f'C{rr}'], INT)
    if k == 0:
        ws[f'D{rr}'] = 0; cell_fmt(ws[f'D{rr}'], MONEY)
        ws[f'E{rr}'] = f"=-{c['lump']}"; cell_fmt(ws[f'E{rr}'], MONEY)
        ws[f'F{rr}'] = f"={c['lump']}"; cell_fmt(ws[f'F{rr}'], MONEY)
    else:
        # payments during year k are for the age at the start of that year: age + k - 1
        ws[f'D{rr}'] = f"=IF(AND(C{rr-1}>={c['start']},C{rr-1}<{c['to']}),12*{c['mo']}*(1+{c['cola']})^(C{rr-1}-{c['start']}),0)"; cell_fmt(ws[f'D{rr}'], MONEY)
        ws[f'E{rr}'] = f"=D{rr}"; cell_fmt(ws[f'E{rr}'], MONEY)
        ws[f'F{rr}'] = f"=MAX(0,F{rr-1}*(1+{c['ret']})-D{rr})"; cell_fmt(ws[f'F{rr}'], MONEY)
ch = LineChart(); ch.title = 'The lump sum, invested and drawn'; ch.height = 7.5; ch.width = 16; ch.y_axis.numFmt = '$#,##0'; ch.x_axis.title = 'age'; ch.legend = None
ch.add_data(Reference(ws, min_col=6, min_row=49, max_row=111), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=3, min_row=50, max_row=111))
ws.add_chart(ch, 'H4')
verdict(ws, row_v, f"=IF(ISNUMBER({c['runout']}),\"The payments, if you expect to live past \"&{c['runout']}&\". At \"&TEXT({c['ret']},\"0.0%\")&\" the lump sum runs out at that age. The payments are quietly promising \"&IF(ISNUMBER({c['irr']}),TEXT({c['irr']},\"0.0%\"),\"less than nothing\")&\" a year.\",\"The lump sum, if you can earn \"&TEXT({c['ret']},\"0.0%\")&\". Drawn at \"&TEXT({c['mo']},\"$#,##0\")&\" a month it still holds \"&TEXT({c['left']},\"$#,##0\")&\" at \"&{c['to']}&\". The payments promise only \"&IF(ISNUMBER({c['irr']}),TEXT({c['irr']},\"0.0%\"),\"less than nothing\")&\" a year.\")")

# =============================================================================
# 06 Cash Ladder
# =============================================================================
ws = sheet('06 Cash Ladder', 'Cash Ladder', 'Cash you will not need this month, split across CDs or Treasury bills that mature at different times, against leaving it all in the savings account.', widths=(2, 30, 16, 14, 14, 16, 16, 16, 30))
c = {}
head(ws, 4, 'THE CASH, AND WHERE IT SITS NOW'); r = 5
inp(ws, r, 'Cash to place', 20000, MONEY, None, 'cash', c); r += 1
inp(ws, r, 'Savings account rate', 0.041, PCT, None, 'hysa', c); r += 1
inp(ws, r, 'Treasury bills? (no state income tax)', 'Yes', None, 'Yes or No.', 'tr', c); dropdown(ws, f'C{r}', ['Yes', 'No']); r += 1
inp(ws, r, 'Your state income tax rate', 0.05, PCT, 'Used only for Treasuries.', 'st', c); r += 1
inp(ws, r, 'Start date', datetime.date.today(), DATE, None, 'start', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE RUNGS (type a term and a rate; leave a row blank to skip it)', cols='BCDEFGHI'); r += 1
table_head(ws, r, 2, ['Rung', 'Months', 'Rate', 'Share (optional)', 'Amount', 'Matures on', 'Interest to maturity', 'Tax-equivalent rate'])
r += 1
R1 = r
for i, (m, rt) in enumerate([(3, 0.05), (6, 0.049), (9, 0.048), (12, 0.046)], start=1):
    ws[f'B{r}'] = f'Rung {i}'; ws[f'B{r}'].font = f_label
    ws[f'C{r}'] = m; cell_fmt(ws[f'C{r}'], INT, 'input')
    ws[f'D{r}'] = rt; cell_fmt(ws[f'D{r}'], PCT, 'input')
    ws[f'E{r}'] = None; cell_fmt(ws[f'E{r}'], PCT1, 'input')
    r += 1
R4 = r - 1
for rr in range(R1, R4 + 1):
    ws[f'F{rr}'] = f"=IF(C{rr}>0,{c['cash']}*IF(SUM($E${R1}:$E${R4})>0,E{rr}/SUM($E${R1}:$E${R4}),1/COUNTIF($C${R1}:$C${R4},\">0\")),0)"; cell_fmt(ws[f'F{rr}'], MONEY)
    ws[f'G{rr}'] = f"=IF(C{rr}>0,EDATE({c['start']},C{rr}),\"\")"; cell_fmt(ws[f'G{rr}'], DATE)
    ws[f'H{rr}'] = f"=F{rr}*D{rr}*C{rr}/12"; cell_fmt(ws[f'H{rr}'], MONEY)
    ws[f'I{rr}'] = f"=IF(C{rr}>0,IF({c['tr']}=\"Yes\",D{rr}/(1-{c['st']}),D{rr}),\"\")"; cell_fmt(ws[f'I{rr}'], PCT)
r += 1
head(ws, r, 'THE WHOLE'); r += 1
out(ws, r, 'Blended rate', f"=SUMPRODUCT(F{R1}:F{R4},D{R1}:D{R4})/{c['cash']}", PCT, None, False, 'bl', c); r += 1
out(ws, r, 'Blended tax-equivalent rate', f"=IF({c['tr']}=\"Yes\",{c['bl']}/(1-{c['st']}),{c['bl']})", PCT, 'What a taxable account would have to pay to match a Treasury, given state tax.', False, 'blt', c); r += 1
out(ws, r, 'Average months locked', f"=SUMPRODUCT(F{R1}:F{R4},C{R1}:C{R4})/{c['cash']}", MONTHS, None, False); r += 1
out(ws, r, 'First rung unlocks', f"=MIN(G{R1}:G{R4})", DATE, None, False, 'first', c); r += 1
out(ws, r, 'Interest by the last maturity, ladder', f"=SUM(H{R1}:H{R4})", MONEY, None, False); r += 1
out(ws, r, 'Same money in the account, same dates', f"=SUMPRODUCT(F{R1}:F{R4},C{R1}:C{R4})*{c['hysa']}/12", MONEY, None, False); r += 1
out(ws, r, 'A year of the ladder against the account', f"={c['cash']}*({c['blt']}-{c['hysa']})", MONEY, 'Positive: the ladder earns more.', True, 'edge', c); r += 2
small(ws, r, 'Each rung earns simple interest to its maturity. A Treasury\'s interest is exempt from state income tax and a savings account\'s is not, so the comparison uses the tax-equivalent rate when the rungs are Treasuries. A ladder is a promise not to touch the money: keep what you might need before the first rung matures in the account. When a rung matures, roll it to the longest term and the ladder keeps unlocking. Early withdrawal penalties, reinvestment and federal tax are not modelled.')
verdict(ws, row_v, f"=IF(ABS({c['edge']})<1,\"A wash: the ladder and the account earn within a dollar a year of each other.\",IF({c['edge']}>0,\"The ladder earns \"&TEXT({c['edge']},\"$#,##0\")&\" a year more than the account. Blended it pays \"&TEXT({c['bl']},\"0.00%\")&IF({c['tr']}=\"Yes\",\", worth \"&TEXT({c['blt']},\"0.00%\")&\" once state tax is counted\",\"\")&\"; the first rung unlocks \"&TEXT({c['first']},\"mmm d, yyyy\")&\".\",\"The account wins by \"&TEXT(-{c['edge']},\"$#,##0\")&\" a year. Leave it where it is.\"))")

# =============================================================================
# 07 Three Paychecks
# =============================================================================
ws = sheet('07 Three Paychecks', 'Three Paychecks', 'Paid every two weeks, two months a year hold a third check; paid weekly, four months hold a fifth. Here are the months, the amount, and the budget that never counts on it.', widths=(2, 34, 18, 18, 18, 14))
c = {}
head(ws, 4, 'HOW YOU ARE PAID'); r = 5
next_fri = datetime.date.today(); next_fri += datetime.timedelta(days=((4 - next_fri.weekday()) % 7) or 7)
inp(ws, r, 'How often', 'Every two weeks', None, 'Pick from the list.', 'freq', c); dropdown(ws, f'C{r}', FREQS); r += 1
inp(ws, r, 'Next payday', next_fri, DATE, None, 'next', c); r += 1
inp(ws, r, 'Net pay per check', 1800, MONEY2, None, 'net', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE NUMBERS', cols=('B', 'C', 'D')); r += 1
out(ws, r, 'Days between paydays', f"=IF({c['freq']}=\"Every week\",7,IF({c['freq']}=\"Every two weeks\",14,0))", INT, '0 means there is no extra month.', False, 'step', c); r += 1
out(ws, r, 'Checks a year', f"=INDEX({PF_N},MATCH({c['freq']},{PF_LABELS},0))", INT, None, False, 'n', c); r += 1
out(ws, r, 'Checks in a normal month', f"=INDEX({PF_BASE},MATCH({c['freq']},{PF_LABELS},0))", INT, None, False, 'base', c); r += 1
out(ws, r, 'Extra checks a year', f"={c['n']}-12*{c['base']}", INT, None, False, 'extra', c); r += 1
out(ws, r, 'Build the month on', f"={c['base']}*{c['net']}", MONEY, 'The normal count of checks.', True, 'budget', c); r += 1
out(ws, r, 'The extra checks are worth', f"={c['extra']}*{c['net']}", MONEY, 'A year, free to go somewhere on purpose.', True, 'extra$', c); r += 1
out(ws, r, 'Spread over the year, a month', f"={c['extra$']}/12", MONEY2, None, False); r += 1
out(ws, r, 'Take-home for the year', f"={c['n']}*{c['net']}", MONEY, None, False); r += 1
r += 1
head(ws, r, 'THE NEXT TWELVE MONTHS', cols=('B', 'C', 'D', 'E')); r += 1
table_head(ws, r, 2, ['Month', 'Paydays', 'Extra checks', 'Extra pay']); r += 1
M1 = r
# helper list of paydays in column H (hidden-ish)
ws['H4'] = 'Payday list (helper)'; ws['H4'].font = f_note
for k in range(0, 60):
    ws[f'H{5+k}'] = f"=IF({c['step']}>0,{c['next']}+{c['step']}*({k}-4),0)"; ws[f'H{5+k}'].number_format = DATE; ws[f'H{5+k}'].font = f_note
PAYDAYS = '$H$5:$H$64'
for k in range(12):
    ws[f'B{r}'] = f"=DATE(YEAR({c['next']}),MONTH({c['next']})+{k},1)"; cell_fmt(ws[f'B{r}'], 'mmmm yyyy')
    ws[f'C{r}'] = f"=IF({c['step']}>0,SUMPRODUCT(({PAYDAYS}>=B{r})*({PAYDAYS}<=EOMONTH(B{r},0))),{c['base']})"; cell_fmt(ws[f'C{r}'], INT)
    ws[f'D{r}'] = f"=MAX(0,C{r}-{c['base']})"; cell_fmt(ws[f'D{r}'], INT)
    ws[f'E{r}'] = f"=D{r}*{c['net']}"; cell_fmt(ws[f'E{r}'], MONEY)
    r += 1
M12 = r - 1
out(ws, r, 'Months with an extra check', f"=COUNTIF(D{M1}:D{M12},\">0\")", INT, None, True, 'xm', c); r += 2
small(ws, r, 'The months are counted from your next payday forward, in the helper list to the right. A payday that falls on a weekend or holiday and is paid early may shift a month by one check. Decide now where the extra goes: the whole check to the highest-rate debt, to the cushion, or to the sinking funds. A standing rule beats a good intention.')
verdict(ws, row_v, f"=IF({c['step']}=0,\"Paid \"&LOWER({c['freq']})&\", every month has the same number of checks. There is no extra one.\",{c['xm']}&\" of the next twelve months hold an extra check: \"&TEXT({c['extra$']},\"$#,##0\")&\" a year the budget does not need. Build the month on \"&TEXT({c['budget']},\"$#,##0\")&\" (\"&{c['base']}&\" checks).\")")

# =============================================================================
# 08 401k Loan
# =============================================================================
ws = sheet('08 401k Loan', 'Borrow From Yourself', 'A 401(k) loan pays its interest back to you, which makes it look free. It is not: its cost is what the money did not earn while it was out. Here is that number, and what it costs if you leave the job.')
c = {}
head(ws, 4, 'THE LOAN'); r = 5
inp(ws, r, 'Borrow', 20000, MONEY, None, 'loan', c); r += 1
inp(ws, r, 'Vested balance in the plan', 80000, MONEY, None, 'vested', c); r += 1
inp(ws, r, 'Loan rate', 0.085, PCT, 'Usually prime plus one or two.', 'lr', c); r += 1
inp(ws, r, 'Repay over, months', 60, INT, 'Up to 60 by law, except a first home.', 'm', c); r += 1
inp(ws, r, 'Return the plan would earn', 0.07, PCT, None, 'mk', c); r += 1
inp(ws, r, 'Contributions you would pause, a month', 500, MONEY, '0 if none.', 'paused', c); r += 1
head(ws, r, 'THE TWO RISKS'); r += 1
inp(ws, r, 'Your age', 40, INT, None, 'age', c); r += 1
inp(ws, r, 'Your marginal tax rate', 0.22, PCT, None, 'marg', c); r += 1
inp(ws, r, 'If you left the job after, months', 24, INT, 'Blank to skip.', 'leave', c); r += 1
head(ws, r, 'THE ALTERNATIVE'); r += 1
inp(ws, r, 'A personal loan at', 0.12, PCT, 'Blank to skip.', 'pr', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'WHAT IT COSTS'); r += 1
out(ws, r, 'Most the plan can lend', f"=MIN({PL['maxLoan']},MAX({c['vested']}*{PL['maxShare']},MIN({PL['floor']},{c['vested']})))", MONEY, 'The lesser of $50,000 and half the vested balance (Tax Tables).', False, 'allowed', c); r += 1
out(ws, r, 'Inside the limits?', f"=IF({c['loan']}>{c['allowed']},\"No: over the amount the plan can lend\",IF({c['m']}>{PL['maxMonths']},\"No: over \"&{PL['maxMonths']}&\" months\",\"Yes\"))", None); r += 1
out(ws, r, 'Monthly payment', f"=PMT({c['lr']}/12,{c['m']},-{c['loan']})", MONEY2, None, False, 'pmt', c); r += 1
out(ws, r, 'Interest, paid to yourself', f"={c['pmt']}*{c['m']}-{c['loan']}", MONEY, None, False); r += 1
out(ws, r, 'Left in the plan, the loan would have grown to', f"={c['loan']}*(1+{c['mk']}/12)^{c['m']}", MONEY, None, False, 'stayed', c); r += 1
out(ws, r, 'Paid back a month at a time, it grows to', f"=FV({c['mk']}/12,{c['m']},-{c['pmt']},0)", MONEY, 'Each payment lands and grows from then.', False, 'repaid', c); r += 1
out(ws, r, 'Growth the borrowed money misses', f"={c['stayed']}-{c['repaid']}", MONEY, 'Can be negative when the loan rate beats the plan return.', False, 'lost', c); r += 1
out(ws, r, 'Paused contributions, at what they would be worth', f"=FV({c['mk']}/12,{c['m']},-{c['paused']},0)", MONEY, None, False, 'pc', c); r += 1
out(ws, r, 'Cost of the plan loan', f"=MAX(0,{c['lost']}+{c['pc']})", MONEY, None, True, 'cost', c); r += 1
r += 1
head(ws, r, 'IF YOU LEAVE THE JOB'); r += 1
out(ws, r, 'Still owed at that point', f"=IF({c['leave']}>0,MAX(0,-FV({c['lr']}/12,{c['leave']},-{c['pmt']},{c['loan']})),\"n/a\")", MONEY, None, False, 'owed', c); r += 1
out(ws, r, 'Income tax on it', f"=IFERROR({c['owed']}*{c['marg']},\"n/a\")", MONEY, 'If not repaid by the tax deadline it becomes a withdrawal.', False, 'ltax', c); r += 1
out(ws, r, 'Penalty', f"=IFERROR(IF({c['age']}+{c['leave']}/12<{PL['penaltyAge']},{c['owed']}*{PL['penalty']},0),\"n/a\")", MONEY, 'Before 59 and a half.', False, 'lpen', c); r += 1
out(ws, r, 'The bill', f"=IFERROR({c['ltax']}+{c['lpen']},\"n/a\")", MONEY, 'On money that is also gone from the plan.', True, 'bill', c); r += 1
r += 1
head(ws, r, 'THE ALTERNATIVE'); r += 1
out(ws, r, 'Personal loan payment', f"=IF({c['pr']}>0,PMT({c['pr']}/12,{c['m']},-{c['loan']}),\"n/a\")", MONEY2, None, False, 'pp', c); r += 1
out(ws, r, 'Its interest, paid to a lender', f"=IFERROR({c['pp']}*{c['m']}-{c['loan']},\"n/a\")", MONEY, None, True, 'pi', c); r += 1
out(ws, r, 'Cheaper', f"=IFERROR(IF({c['pi']}<{c['cost']},\"The personal loan\",\"The plan loan\"),\"n/a\")", None, None, True, 'cheaper', c); r += 2
small(ws, r, 'A lost employer match is not counted and would make it worse. On leaving a job the unpaid balance is due; a plan may allow rolling the offset amount to an IRA by the tax deadline. The limits are on the Tax Tables tab; a plan may set tighter ones and the plan document wins.')
verdict(ws, row_v, f"=IF({c['cost']}>0,\"Borrowing from yourself costs about \"&TEXT({c['cost']},\"$#,##0\")&\" over \"&{c['m']}&\" months\",\"At a loan rate above the plan return, the loan costs nothing on the arithmetic\")&IF(ISNUMBER({c['bill']}),\". Leave the job after \"&{c['leave']}&\" months and \"&TEXT({c['owed']},\"$#,##0\")&\" is still owed: a bill of about \"&TEXT({c['bill']},\"$#,##0\")&\" if it is not repaid.\",\".\")&IF(ISNUMBER({c['pi']}),\" \"&{c['cheaper']}&\" is cheaper.\",\"\")")

# =============================================================================
# 09 Sinking Funds
# =============================================================================
ws = sheet('09 Sinking Funds', 'Sinking Funds', 'The bills that come once or twice a year and always feel like a surprise. Spread into a monthly amount, with what the fund should already hold today and what lands in each of the next twelve months.', widths=(2, 28, 14, 14, 14, 14, 14, 14, 16))
c = {}
head(ws, 4, 'TODAY'); r = 5
inp(ws, r, 'Today', '=TODAY()', DATE, 'Left as a formula, this is always today. Type a date to fix it.', 'today', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE BILLS (type up to fifteen; leave a row blank to skip it)', cols='BCDEFGHI'); r += 1
table_head(ws, r, 2, ['Bill', 'Amount', 'Every (months)', 'Next due', 'A month', 'Hold now', 'A year', 'Every, or 1 (helper)'])
r += 1
B1 = r
today = datetime.date.today(); y = today.year
demo = [('Car insurance', 900, 6, datetime.date(y, 12, 15)), ('Holidays and gifts', 1200, 12, datetime.date(y, 12, 1)), ('Property tax', 3200, 12, datetime.date(y + 1, 4, 10)), ('Tyres', 800, 36, datetime.date(y + 1, 8, 1)), ('Vet, the yearly visit', 350, 12, datetime.date(y + 1, 2, 20))]
for i in range(15):
    if i < len(demo):
        n, a, e, d = demo[i]
    else:
        n = a = e = d = None
    ws[f'B{r}'] = n; cell_fmt(ws[f'B{r}'], 'General', 'input')
    ws[f'C{r}'] = a; cell_fmt(ws[f'C{r}'], MONEY, 'input')
    ws[f'D{r}'] = e; cell_fmt(ws[f'D{r}'], INT, 'input')
    ws[f'E{r}'] = d; cell_fmt(ws[f'E{r}'], DATE, 'input')
    ws[f'F{r}'] = f"=IF(AND(C{r}>0,D{r}>0),C{r}/D{r},0)"; cell_fmt(ws[f'F{r}'], MONEY2)
    ws[f'G{r}'] = f"=IF(AND(C{r}>0,D{r}>0,E{r}>0),C{r}*MAX(0,MIN(1,1-MAX(0,(E{r}-{c['today']})/30.4375)/D{r})),0)"; cell_fmt(ws[f'G{r}'], MONEY)
    ws[f'H{r}'] = f"=IF(AND(C{r}>0,D{r}>0),C{r}*12/D{r},0)"; cell_fmt(ws[f'H{r}'], MONEY)
    ws[f'I{r}'] = f"=IF(D{r}>0,D{r},1)"; ws[f'I{r}'].font = f_note; ws[f'I{r}'].number_format = INT
    r += 1
B15 = r - 1
ws[f'B{r}'] = 'All bills'; ws[f'B{r}'].font = f_strong
for col in 'FGH':
    ws[f'{col}{r}'] = f"=SUM({col}{B1}:{col}{B15})"; ws[f'{col}{r}'].font = f_strong; ws[f'{col}{r}'].fill = fill_answer; ws[f'{col}{r}'].number_format = MONEY
c['monthly'] = f'$F${r}'; c['hold'] = f'$G${r}'; c['yearly'] = f'$H${r}'
r += 2
head(ws, r, 'THE NEXT TWELVE MONTHS', cols=('B', 'C', 'D', 'E')); r += 1
table_head(ws, r, 2, ['Month', 'Lands', 'Fund after', 'Set aside']); r += 1
L0 = r
ws[f'B{r}'] = 'Today'; ws[f'B{r}'].font = f_label
ws[f'D{r}'] = f"={c['hold']}"; cell_fmt(ws[f'D{r}'], MONEY)
r += 1
L1 = r
AMT = f'$C${B1}:$C${B15}'; DUE = f'$E${B1}:$E${B15}'; EV = f'$I${B1}:$I${B15}'; PRESENT = f'($C${B1}:$C${B15}>0)*($D${B1}:$D${B15}>0)*($E${B1}:$E${B15}>0)'
for k in range(12):
    ws[f'B{r}'] = f"=DATE(YEAR({c['today']}),MONTH({c['today']})+{k},1)"; cell_fmt(ws[f'B{r}'], 'mmm yyyy')
    diff = f"((YEAR(B{r})-YEAR({DUE}))*12+MONTH(B{r})-MONTH({DUE}))"
    ws[f'C{r}'] = f"=SUMPRODUCT({PRESENT}*{AMT}*({diff}>=0)*(MOD({diff},{EV})=0))"; cell_fmt(ws[f'C{r}'], MONEY)
    ws[f'D{r}'] = f"=D{r-1}+{c['monthly']}-C{r}"; cell_fmt(ws[f'D{r}'], MONEY)
    ws[f'E{r}'] = f"={c['monthly']}"; cell_fmt(ws[f'E{r}'], MONEY)
    r += 1
L12 = r - 1
out(ws, r, 'Lowest the fund gets', f"=MIN(D{L0}:D{L12})", MONEY, None, True, 'low', c); r += 1
out(ws, r, 'Start with this much if it runs short', f"=IF({c['low']}<0,{c['hold']}-{c['low']},{c['hold']})", MONEY, None, True, 'startWith', c); r += 1
ch = LineChart(); ch.title = 'The fund over the next year'; ch.height = 7; ch.width = 15; ch.y_axis.numFmt = '$#,##0'; ch.legend = None
ch.add_data(Reference(ws, min_col=4, min_row=L0 - 1, max_row=L12), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=L0, max_row=L12))
ws.add_chart(ch, f'G{L0 - 2}')
r += 1
small(ws, r, 'Each bill\'s set-aside is its amount over the months between one and the next. What the fund should hold today is the share of each bill that has accrued since the last one: a yearly bill due next month is nearly all owed already. The months ahead start from that amount, add the set-aside and take out what lands.')
verdict(ws, row_v, f"=IF({c['monthly']}=0,\"Add a bill to see the set-aside.\",\"Set aside \"&TEXT({c['monthly']},\"$#,##0\")&\" a month. The fund should hold \"&TEXT({c['hold']},\"$#,##0\")&\" today.\"&IF({c['low']}<0,\" Starting there it runs \"&TEXT(-{c['low']},\"$#,##0\")&\" short in the year ahead: start with \"&TEXT({c['startWith']},\"$#,##0\")&\" instead.\",\"\"))")

# =============================================================================
# 10 Stay or Move
# =============================================================================
ws = sheet('10 Stay or Move', 'Stay or Move', 'The renewal letter came. Over the lease you would sign, does moving pay for itself? And the one number to take back to the landlord.')
c = {}
head(ws, 4, 'THE RENTS'); r = 5
inp(ws, r, 'Rent now', 1800, MONEY, None, 'now', c); r += 1
inp(ws, r, 'Rent if you stay', 1950, MONEY, 'The renewal letter.', 'stay', c); r += 1
inp(ws, r, 'Rent at the new place', 1700, MONEY, None, 'move', c); r += 1
head(ws, r, 'WHAT MOVING COSTS, ONCE'); r += 1
inp(ws, r, 'Movers, truck, boxes', 1200, MONEY, None, 'movers', c); r += 1
inp(ws, r, 'Fees and deposit you will not see again', 500, MONEY, 'A returned deposit is not a cost.', 'fees', c); r += 1
inp(ws, r, 'Overlap rent, time off, the rest', 800, MONEY, None, 'once', c); r += 1
head(ws, r, 'WHAT CHANGES EVERY MONTH'); r += 1
inp(ws, r, 'Commute, more (+) or less (-) a month', 50, MONEY, None, 'commute', c); r += 1
inp(ws, r, 'Utilities, parking, the rest, + or -', 0, MONEY, None, 'otherM', c); r += 1
inp(ws, r, 'Lease you would sign, months', 12, INT, None, 'term', c); r += 1
r += 1
row_v = r; r += 3
head(ws, r, 'THE NUMBERS'); r += 1
out(ws, r, 'The rise', f"=IF({c['now']}>0,({c['stay']}-{c['now']})/{c['now']},\"n/a\")", PCT1, None, False, 'rise', c); r += 1
out(ws, r, 'Moving, once', f"={c['movers']}+{c['fees']}+{c['once']}", MONEY, None, False, 'oneoff', c); r += 1
out(ws, r, 'The new place, a month, all in', f"={c['move']}+{c['commute']}+{c['otherM']}", MONEY, None, False, 'mm', c); r += 1
out(ws, r, 'Moving saves, a month', f"={c['stay']}-{c['mm']}", MONEY, 'Negative: the new place costs more once everything is counted.', False, 'saving', c); r += 1
out(ws, r, 'Break-even', f"=IF({c['saving']}>0,{c['oneoff']}/{c['saving']},\"never\")", MONTHS, 'Months for the saving to pay back the moving costs.', True, 'be', c); r += 1
out(ws, r, 'Stay: cost over the lease', f"={c['stay']}*{c['term']}", MONEY, None, False, 'ts', c); r += 1
out(ws, r, 'Move: cost over the lease, moving included', f"={c['mm']}*{c['term']}+{c['oneoff']}", MONEY, None, False, 'tm', c); r += 1
out(ws, r, 'Moving saves over the lease', f"={c['ts']}-{c['tm']}", MONEY, 'Negative: staying is cheaper.', True, 'delta', c); r += 1
out(ws, r, 'The counter-offer', f"={c['tm']}/{c['term']}", MONEY, 'The rent at which staying costs exactly what moving does. Anything under it and you are ahead staying.', True, 'counter', c); r += 2
small(ws, r, 'Both paths are summed over the lease term. Staying costs the new rent each month; moving costs the new place\'s rent plus the monthly changes, plus the one-off costs on day one. An empty month costs the landlord a month of rent, which is why the counter-offer is worth making. The hassle of moving is real and has no box.')
r += 2
head(ws, r, 'MONTH BY MONTH'); T0 = r; r += 1
table_head(ws, r, 2, ['Month', 'Stay, so far', 'Move, so far']); r += 1
first = r
for k in range(0, 37):
    ws[f'B{r}'] = k; ws[f'B{r}'].font = f_calc; ws[f'B{r}'].alignment = Alignment(horizontal='center')
    ws[f'C{r}'] = f"={c['stay']}*{k}"; cell_fmt(ws[f'C{r}'], MONEY)
    ws[f'D{r}'] = f"={c['oneoff']}+{c['mm']}*{k}"; cell_fmt(ws[f'D{r}'], MONEY); ws[f'D{r}'].alignment = Alignment(horizontal='left')
    r += 1
ch = LineChart(); ch.title = 'What each path has cost'; ch.height = 7.5; ch.width = 16; ch.y_axis.numFmt = '$#,##0'; ch.x_axis.title = 'months'
ch.add_data(Reference(ws, min_col=3, min_row=first - 1, max_col=4, max_row=r - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=first, max_row=r - 1))
ws.add_chart(ch, f'F{T0}')
verdict(ws, row_v, f"=IF(ABS({c['delta']})<50,\"A wash over \"&{c['term']}&\" months: the two come within \"&TEXT(ABS({c['delta']}),\"$#,##0\")&\". Stay, unless you want to move.\",IF({c['delta']}>0,\"Move. Over \"&{c['term']}&\" months it is \"&TEXT({c['delta']},\"$#,##0\")&\" cheaper, moving costs included; they are paid back in \"&TEXT({c['be']},\"0\")&\" months.\",\"Stay. Over \"&{c['term']}&\" months moving would cost \"&TEXT(-{c['delta']},\"$#,##0\")&\" more.\"))&\" The counter-offer is \"&TEXT({c['counter']},\"$#,##0\")&\".\"")

# ---- order and finish -----------------------------------------------------------
order = ['Start Here'] + [t[0] for t in TOOLS] + ['Tax Tables']
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = NAVY if ws.title in ('Start Here', 'Tax Tables') else '3987E5'
    ws.page_setup.fitToWidth = 1; ws.page_setup.orientation = 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print('wrote', OUT)
